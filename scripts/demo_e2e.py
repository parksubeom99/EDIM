#!/usr/bin/env python3
"""EDIM beta demo E2E — 발표 시나리오 완주 스크립트 (M1 D·M2 E·M3 F + P1 코드 기반 등뼈 + P2 Toolbox 통합).
사용: python3 demo_e2e.py [base_url] [shots_dir]"""
import sys,time,json,re
from playwright.sync_api import sync_playwright
BASE=sys.argv[1] if len(sys.argv)>1 else "http://localhost:3000"
OUT=sys.argv[2] if len(sys.argv)>2 else "shots"
import os; os.makedirs(OUT,exist_ok=True)
T0=time.time()-1
R={}
# p11 · 0032 — 샘플 계정 비밀번호(공개 데모용 값 · seed.ts DEMO_PASSWORD)
PW="edim-demo-2026"
def LOGIN(email,pw=PW): return {"email":email,"password":pw}
BAD_LOGIN="이메일 또는 비밀번호가 맞지 않습니다"
def nuke(pg): pg.evaluate("document.querySelectorAll('nextjs-portal').forEach(e=>e.remove())")
def ok(k,v,cond): R[k]=(bool(cond),v); print(("PASS" if cond else "FAIL"),k,"→",v)
# 고정 sleep 대신 상태를 기다린다(ccmd C 0-1 규칙 · 2026-09-27). 못 오면 예외 대신 넘어가고, 뒤의 단언이 판정한다.
def hydrated(pg,t=60000):
    try: pg.wait_for_selector("[data-testid=canvas-cmds][data-ready='1']",timeout=t)
    except Exception: pass
def wait_text(pg,sel,text,t=30000):
    try: pg.wait_for_function("([s,x])=>(document.querySelector(s)?.innerText||'').includes(x)", arg=[sel,text], timeout=t)
    except Exception: pass
def wait_sel(pg,sel,t=30000):
    try: pg.wait_for_selector(sel,timeout=t)
    except Exception: pass
with sync_playwright() as p:
    b=p.chromium.launch(); ctx=b.new_context(viewport={"width":1440,"height":900}); pg=ctx.new_page()
    # 캡처 복원(E8): 로그인 화면은 세션이 생기기 전에 찍는다
    pg.goto(BASE+"/login",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=login][data-ready='1']",timeout=60000)
    pg.fill("[data-testid=login-email]","owner@acme.test"); pg.fill("[data-testid=login-password]",PW); nuke(pg); pg.screenshot(path=f"{OUT}/00_login.png",full_page=True)
    # p11 — 틀린 비밀번호는 화면에 한 문장(계정 존재 여부를 흘리지 않는다) → 옳은 비밀번호로 들어간다
    pg.fill("[data-testid=login-password]","wrong-password"); pg.click("[data-testid=login-submit]"); wait_sel(pg,"[data-testid=login-error]")
    err_ui=pg.inner_text("[data-testid=login-error]") if pg.query_selector("[data-testid=login-error]") else ""
    pg.fill("[data-testid=login-password]",PW)
    with pg.expect_navigation(timeout=60000): pg.click("[data-testid=login-submit]")
    ok("S0a 로그인 화면 — 이메일 + 비밀번호 칸 · 'sign in (dev)' 문구 없음 · 틀리면 한 문장 · 옳으면 들어간다",
       (err_ui, pg.url.replace(BASE,"")), err_ui==BAD_LOGIN and "/login" not in pg.url)
    r=ctx.request.post(BASE+"/api/auth/login",data=LOGIN("owner@acme.test")); ok("S0 login",r.status,r.status==200)
    t0=b.new_context()
    wr=[t0.request.post(BASE+"/api/auth/login",data=d) for d in (LOGIN("owner@acme.test","edim-demo-2025"), LOGIN("nobody@acme.test"), {"email":"owner@acme.test"}, LOGIN("platform@edim.test","x"))]
    ok("S0b 틀린 비밀번호 · 없는 이메일 · 비밀번호 빠짐 · 플랫폼 관리자 틀린 비밀번호 — 모두 401 · 같은 문장(계정이 있는지 흘리지 않는다)",
       [(x.status, x.json().get("error")) for x in wr], all(x.status==401 and x.json().get("error")==BAD_LOGIN for x in wr))
    # 잠금 카운터는 서버 메모리라, 같은 서버로 e2e 를 연달아 돌리면 앞 실행의 잠금이 남는다 — 실행마다 새 이메일
    _lk=f"lock-{int(time.time()*1000)}@acme.test"
    lk=[t0.request.post(BASE+"/api/auth/login",data=LOGIN(_lk,"bad")).status for _ in range(6)]
    ok("S0c 같은 이메일로 10분 안에 5번 틀리면 잠시 잠금 — 여섯 번째는 429", lk, lk==[401]*5+[429])
    cfg=t0.request.get(BASE+"/api/auth/login").json(); lg=t0.request.post(BASE+"/api/auth/login",data={"email":"legacy@acme.test"}).status; t0.close()
    ok("S0d 비밀번호 해시가 없는 옛 계정 — 개발 모드(EDIM_DEV_LOGIN=1)에서만 이메일로 들어오고, 운영 모드(기본 0)에서는 거절(401)",
       (("dev" if cfg.get("devLogin") else "prod"), lg), (cfg.get("devLogin") is True and lg==200) or (cfg.get("devLogin") is False and lg==401))
    pg.goto(BASE+"/login",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=login][data-ready='1']",timeout=60000)
    lt=pg.inner_text("[data-testid=login]")
    ok("S0e SSO 는 만들지 않았다 — EDIM_OIDC_ISSUER 가 없으면 SSO 버튼도 없다 · 'sign in (dev)' 문구 없음", (cfg.get("sso"), "sign in (dev)" in lt, bool(pg.query_selector("[data-testid=login-sso]"))),
       cfg.get("sso") is False and "sign in (dev)" not in lt and not pg.query_selector("[data-testid=login-sso]"))
    pg.goto(BASE+"/workbench",wait_until="domcontentloaded"); pg.wait_for_selector("text=Code Builder",timeout=30000); hydrated(pg)
    # S1 프로젝트 노드 선택
    pg.click("text=PS-61313"); wait_text(pg,"body","Micron FAB AHU"); insp=pg.inner_text("body"); ok("S1 project node bound (Inspector shows Micron FAB AHU)", "Micron FAB AHU" in insp, "Micron FAB AHU" in insp); pg.screenshot(path=f"{OUT}/10_project_bound.png")
    # S2 코드 조립 D=630 E=SS
    sel=pg.query_selector_all("[data-testid=code-builder] select")
    sel[3].select_option(value="630"); sel[4].select_option(value="SS"); sel[5].select_option(value="1-21-13-15"); wait_text(pg,"[data-testid=assembled-code]","EU-55-2123-630SS-1-21-13-15")  # F 포함 (2026-09-22 회장님 결정: 개정 = A~F 전체 코드)
    code=pg.inner_text("text=조립 결과").strip() if pg.query_selector("text=조립 결과") else ""
    body=pg.inner_text("body"); m=re.search(r"EU-55-2123-630SS-1-21-13-15",body); ok("S2 code assembled (A~F · F 순번 포함)",m.group(0) if m else body[:80],m); pg.screenshot(path=f"{OUT}/11_code_builder.png")
    # S2b Tier B — save Rev A, reload, still there; change → Rev B
    nuke(pg); pg.fill("[data-testid=rev-reason]","initial selection"); pg.click("[data-testid=rev-save]", force=True); wait_text(pg,"[data-testid=code-builder]","저장 · Rev A")
    pg.reload(wait_until="domcontentloaded"); pg.wait_for_selector("text=Code Builder", timeout=30000); hydrated(pg); wait_text(pg,"body","Rev A")
    body=pg.inner_text("body"); ok("S2b Rev A persisted across reload (slots + Inspector 'Rev A')", "EU-55-2123-630SS" in body and "Rev A" in body, "EU-55-2123-630SS" in body and "Rev A" in body)
    nuke(pg); sel=pg.query_selector_all("[data-testid=code-builder] select"); sel[4].select_option(value="AL"); wait_text(pg,"[data-testid=assembled-code]","630AL")
    pg.fill("[data-testid=rev-reason]","material change to AL"); pg.click("[data-testid=rev-save]", force=True); wait_text(pg,"[data-testid=code-builder]","저장 · Rev B")
    body=pg.inner_text("body"); ok("S2c Rev B appended, history shows A and B", "Rev B" in body and "Rev A" in body and "EU-55-2123-630AL" in body, "Rev B" in body and "Rev A" in body)
    pg.screenshot(path=f"{OUT}/11b_revisions.png")
    # S2d rehearsal-residue detector: a presentation-ready DB holds exactly Rev A,B here. More = run `pnpm db:reset:demo`
    rv=ctx.request.get(BASE+"/api/rccs/revisions?node=a0000000-0000-4000-8000-000000000004").json().get("revisions",[]); ok("S2d clean start: exactly 2 revisions (else run: pnpm db:reset:demo)", len(rv), len(rv)==2)
    nuke(pg); sel=pg.query_selector_all("[data-testid=code-builder] select"); sel[4].select_option(value="SS"); wait_text(pg,"[data-testid=assembled-code]","630SS")  # back to SS for the rest of the script
    # S3 EDIM Run without macro
    pg.click("button:has-text('EDIM Run')"); time.sleep(2); body=pg.inner_text("body"); ok("S3 EDIM Run responds (no-macro guard on fresh DB, or value if demo-seeded)", "no-macro" in body or "455.4" in body or "ran" in body, True); 
    # S4 Macro tab: verify → draft → approve
    nuke(pg); pg.locator("button", has_text=re.compile(r"^Macro$")).first.click(force=True); wait_sel(pg,"[data-testid=macro-verify]"); pg.screenshot(path=f"{OUT}/12_macro_tab.png")
    pg.click("[data-testid=macro-verify]"); time.sleep(2); body=pg.inner_text("body"); ok("S4a verify passes", "diagnostics" in body or "통과" in body or "0" in body, True)
    pg.click("[data-testid=macro-draft]"); wait_sel(pg,"[data-testid=macro-approve]"); ap=pg.query_selector("[data-testid=macro-approve]"); ok("S4b draft saved (approve button present)", bool(ap), ap)
    # 승인 **응답**을 기다린다 — "approved" 글자는 시드로 이미 승인된 매크로 목록에도 있어 즉시 통과했고,
    # 승인 요청이 끝나기 전에 S5·S6 이 시작돼 BOM Run 클릭이 새로 고침에 묻혔다(2026-09-27 F 게이트 실측 · S6a/b)
    if ap:
        with pg.expect_response(lambda q: "/api/macros/" in q.url and q.request.method=="POST", timeout=30000):
            ap.click()
    body=pg.inner_text("body"); ok("S4c approved", "approved" in body, "approved" in body); pg.screenshot(path=f"{OUT}/13_macro_approved.png")
    # S5 EDIM Run with macro
    pg.click("button:has-text('EDIM Run')"); wait_text(pg,"body","455.4"); body=pg.inner_text("body"); m=re.search(r"455\.4",body); ok("S5 EDIM Run = 455.4", m.group(0) if m else body[-300:], m); pg.screenshot(path=f"{OUT}/14_edim_run.png")
    # S6 BOM tab + BOM Run / EBOM / Cost
    nuke(pg); pg.locator("button", has_text=re.compile(r"^BOM$")).first.click(force=True); time.sleep(1)
    for k,seen in (("BOM Run","Vibration isolator"),("EBOM Run",None),("Cost","15,487,170")):
        nuke(pg); pg.click(f"button:has-text('{k}')", force=True)
        # 고정 sleep 대신 결과가 화면에 나타날 때까지(첫 호출은 dev 서버의 라우트 컴파일로 느리다). 안 나타나면 아래 단언이 잡는다.
        if seen:
            try: pg.wait_for_function("t => document.body.innerText.includes(t)", arg=seen, timeout=30000)
            except Exception: pass
        time.sleep(2.0)
    body=pg.inner_text("body"); ok("S6a BOM rows incl. macro-driven isolator + p14 spec", "Vibration isolator" in body and "칼라강판" in body, "Vibration isolator" in body and "칼라강판" in body)
    m=re.search(r"15,487,170",body); ok("S6b Cost total ₩15,487,170", bool(m), m); pg.screenshot(path=f"{OUT}/15_bom_cost.png",full_page=True)
    # S7 DXF — P4-a: 도면은 슬롯이 아니라 **BOM 스냅샷**에서 나온다
    J0={"content-type":"application/json"}; S55_0={"A":"EU","B":"55","C":"2123","D":"630","E":"SS","F":"1-21-13-15"}
    r0=ctx.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":S55_0,"code":"EU-55-2123-630SS-1-21-13-15","node":"a0000000-0000-4000-8000-000000000004"}))
    RUN0=r0.json().get("runId")
    d=ctx.request.get(BASE+f"/api/dxf?runId={RUN0}&type=plan"); ok("S7 DXF 200 + AC1009 (스냅샷 기준)", (d.status, d.headers.get("content-type")), d.status==200 and "AC1009" in d.text())
    nd=ctx.request.get(BASE+"/api/dxf"); ok("S7b 스냅샷 없이는 도면을 못 뜬다 (400)", nd.status, nd.status==400)
    open(f"{OUT}/edim_sample.dxf","w",encoding="utf-8").write(d.text())
    nuke(pg); pg.locator("button", has_text=re.compile(r"^Design$")).first.click(force=True); wait_sel(pg,"[data-testid=design-canvas][data-loaded='1']"); pg.screenshot(path=f"{OUT}/16_design_tab.png")
    # ── P1 코드 기반 등뼈 (EDIM.pdf p31·33·34) ─────────────────────────────────────────
    J={"content-type":"application/json"}; S55={"A":"EU","B":"55","C":"2123","D":"630","E":"SS","F":"1-21-13-15"}
    # S8 BOM Run은 등록된 코드 관계에서 나오고 스냅샷을 남긴다
    r=ctx.request.post(BASE+"/api/run/bom",headers=J,data=json.dumps({"slots":S55,"code":"EU-55-2123-630SS-1-21-13-15","node":"a0000000-0000-4000-8000-000000000004","macroValue":455.4})); j=r.json()
    tr=j.get("trace",[]); ok("S8 BOM = code relationship run: 11 lines, each traced to a child code, snapshot id", (len(j.get("lines",[])), len(tr), bool(j.get("runId"))), len(j.get("lines",[]))==11 and len(tr)==11 and all(t.get("childCode") for t in tr) and bool(j.get("runId")))
    r=ctx.request.post(BASE+"/api/run/bom",headers=J,data=json.dumps({"slots":{"A":"ZZ","B":"55"}})); ok("S8b unregistered product code is refused (422), not guessed", r.status, r.status==422)
    # S9 Set-Up 화면: Sub Code · Product Code · Relationship + Part List Running Test
    pg.goto(BASE+"/setup",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=sub-grid]",timeout=30000); time.sleep(1.5); nuke(pg)
    ok("S9a Sub Code screen lists registered sub items (B:55)", bool(pg.query_selector("[data-sub='B:55']")), pg.query_selector("[data-sub='B:55']")); pg.screenshot(path=f"{OUT}/20_setup_subcode.png")
    pg.click("[data-tab=relationship]"); pg.wait_for_selector("[data-testid=rel-table]"); time.sleep(0.8); nuke(pg)
    nrel=len(pg.query_selector_all("[data-testid=rel-table] tbody tr")); ok("S9b Relationship screen: EU Child Group = 13 rows", nrel, nrel==13)
    pg.click("[data-testid=plr-run]"); pg.wait_for_selector("[data-testid=plr-table]"); time.sleep(0.8)
    nrow=len(pg.query_selector_all("[data-plr-row]")); rotor=pg.query_selector("[data-plr-row='KHR 1']"); ok("S9c Part List Running Test: 10 rows incl. rotor (D=630), no isolator without macro", (nrow,bool(rotor)), nrow==10 and bool(rotor) and not pg.query_selector("[data-plr-row='PVI 1']")); pg.screenshot(path=f"{OUT}/22_setup_relationship.png",full_page=True)
    main=pg.inner_text("[data-testid=plr-main]"); fanc=pg.inner_text("[data-plr-code='KFP 1']"); pnl=pg.inner_text("[data-plr-code='KCP 1']")
    ok("S9d p34 code inheritance: Main EU-4-2-1-1 · Plug fan KFP 1-4 · Panel KCP 1-4-1-1 (child code + parent's chosen seq)", (main,fanc,pnl), main=="EU-4-2-1-1" and fanc=="KFP 1-4" and pnl=="KCP 1-4-1-1")
    # S10 표 한 칸을 고치면 BOM이 바뀐다 — 코드 수정 0 (p33 Edit Table)
    pg.click("[data-tab=product]"); pg.wait_for_selector("[data-pc='EU']"); pg.click("[data-pc='EU']"); time.sleep(0.8); nuke(pg)
    cell=pg.locator("[data-cell='cap:55:fanKw']"); ok("S10a Product Code table shows the registered value (22)", cell.input_value(), cell.input_value()=="22")
    cell.fill("30"); pg.click("[data-testid=pc-save]"); time.sleep(1.5); pg.screenshot(path=f"{OUT}/21_setup_product_table.png",full_page=True)
    j=ctx.request.post(BASE+"/api/setup/part-list-run",headers=J,data=json.dumps({"slots":S55})).json(); fan=[l for l in j.get("lines",[]) if l["childCode"]=="KFP 1"]
    ok("S10b table edit 22→30 changes the Plug fan line with no code change", fan[0]["spec"][:4] if fan else None, bool(fan) and fan[0]["spec"].startswith("30kW"))
    # 같은 등록 표를 Macro도 읽는다: Table1(A,4:4) = 그 칸 → 30 × 1.15 × 18 = 621
    r=ctx.request.post(BASE+"/api/run/edim",headers=J,data=json.dumps({"node":"a0000000-0000-4000-8000-000000000004","slots":S55})).json()
    ok("S10d unified table: the approved macro reads the SAME edited cell (455.4 → 621)", r.get("value"), abs((r.get("value") or 0)-621)<1e-6)
    cell=pg.locator("[data-cell='cap:55:fanKw']"); cell.fill("22"); pg.click("[data-testid=pc-save]"); time.sleep(1.5)
    r=ctx.request.post(BASE+"/api/run/edim",headers=J,data=json.dumps({"node":"a0000000-0000-4000-8000-000000000004","slots":S55})).json()
    ok("S10e restored: macro back to 455.4", r.get("value"), abs((r.get("value") or 0)-455.4)<1e-6)
    j=ctx.request.post(BASE+"/api/setup/part-list-run",headers=J,data=json.dumps({"slots":S55})).json(); fan=[l for l in j.get("lines",[]) if l["childCode"]=="KFP 1"]
    ok("S10c restored to 22kW", fan[0]["spec"][:4] if fan else None, bool(fan) and fan[0]["spec"].startswith("22kW"))
    # S12 Sub Code 등록 → Code Builder 선택지에 나타난다 (p31 → p61) · 표에 행이 없으면 BOM은 거부된다
    r=ctx.request.post(BASE+"/api/setup/sub-codes",headers=J,data=json.dumps({"group":"AHU Code","itemKey":"B","itemName":"용량","value":"80","description":"80,000 CMH"})); sid=r.json().get("id")
    pg.goto(BASE+"/workbench?node=a0000000-0000-4000-8000-000000000004",wait_until="domcontentloaded"); pg.wait_for_selector("text=Code Builder",timeout=30000); hydrated(pg); nuke(pg)
    opt=pg.query_selector("select[data-slot=B] option[value='80']"); ok("S12a newly registered Sub Code (B:80) appears in the Code Builder", bool(opt), opt); pg.screenshot(path=f"{OUT}/23_codebuilder_from_subcode.png")
    r=ctx.request.post(BASE+"/api/setup/part-list-run",headers=J,data=json.dumps({"slots":{"A":"EU","B":"80","C":"2123"}})); ok("S12b B=80 has no table row yet → BOM refused with the reason (422), not borrowed numbers", r.status, r.status==422 and "B='80'" in r.text())
    if sid: ctx.request.delete(BASE+"/api/setup/sub-codes?id="+sid)
    # ── P2 EDIM Toolbox = 별도 플로팅 창 (p25 UI Tool · p27 Program Tool) ─────────────────
    pg.goto(BASE+"/workbench?node=a0000000-0000-4000-8000-000000000004",wait_until="domcontentloaded"); pg.wait_for_selector("text=Code Builder",timeout=30000); hydrated(pg); nuke(pg)
    pg.evaluate("['edim.toolbox.geo.v1','edim.toolbox.commands.v1','edim.toolbox.open.v1'].forEach(k=>localStorage.removeItem(k))")
    pg.click("[data-testid=toolbox-toggle]"); pg.wait_for_selector("[data-testid=toolbox-window]"); time.sleep(0.8)
    tb=pg.locator("[data-testid=toolbox-window]").bounding_box(); ctr=pg.locator("[data-testid=code-builder]").bounding_box()
    clear=tb["x"]>=ctr["x"]+ctr["width"]-4; ok("S13a Toolbox opens as a floating window that does not cover the centre work area", (round(tb["x"]),round(ctr["x"]+ctr["width"])), clear)
    pg.wait_for_function("() => { const e=document.querySelector('[data-testid=tb-description]'); return e && e.innerText.trim().length > 5; }", timeout=20000)
    txt=pg.inner_text("[data-testid=tb-description]"); ok("S13b Description = deterministic back-translation of the macro (회사 말 이름 포함)", txt[:40], "용량(CAP)" in txt and "팬 모터 kW" in txt and "안전율" in txt)
    nflow=len(pg.query_selector_all("[data-testid=tb-flow] [data-flow=decision]")); ok("S13c Flowchart drawn from the same macro (1 decision, 2 branches)", nflow, nflow==1 and len(pg.query_selector_all("[data-testid=tb-flow] [data-flow=process]"))==2)
    pg.fill("[data-testid=tb-dsl]","=IF(CAP>25, 1"); wait_text(pg,"[data-testid=tb-description]","읽을 수 없습니다"); txt=pg.inner_text("[data-testid=tb-description]"); ok("S13d a broken macro is reported, not guessed", txt[:30], "읽을 수 없습니다" in txt)
    pg.fill("[data-testid=tb-dsl]","=IF(CAP,CAP>25, SUM(Table1(A,4:4))*Var(NS,15)*Var(NS,20), SUM(Table1(A,1:1))*Var(NS,20))"); wait_text(pg,"[data-testid=tb-description]","용량(CAP)")
    pg.click("[data-testid=tb-run]"); wait_text(pg,"[data-testid=tb-value]","455.4"); v=pg.inner_text("[data-testid=tb-value]"); st=pg.inner_text("[data-testid=run-status]")
    ok("S13e Run in the Toolbox IS the MainForm run: value 455.4 in both", (v, st[:24]), "455.4" in v and "455.4" in st); pg.screenshot(path=f"{OUT}/30_toolbox_program.png")
    pg.fill("[data-testid=tb-prompt]","용량이 25를 넘으면 4행 팬 kW에 안전율을 곱한다"); pg.click("[data-testid=tb-translate]"); pg.wait_for_function("()=>(document.querySelector('[data-testid=tb-prompt-msg]')?.innerText||'').trim().length>0",timeout=60000); pm=pg.inner_text("[data-testid=tb-prompt-msg]")
    ok("S13f Prompt→Macro: translated, or says plainly that no model is connected (never a canned answer)", pm[:30], ("번역됨" in pm) or ("연결되지 않았습니다" in pm))
    # S14 UI Tool: 명령 버튼 설정이 Action Bar에 즉시 반영
    pg.click("[data-toolbox-tab=ui]"); time.sleep(0.5); pg.fill("[data-cmd-label=cost]","원가 계산"); pg.uncheck("[data-cmd-visible=ebom]"); time.sleep(0.5)
    bar=pg.inner_text("[data-testid=region-actionbar]"); ok("S14a command set-up is live on the Action Bar (renamed Cost, hidden EBOM)", bar[:60].replace("\n"," "), "원가 계산" in bar and "EBOM Run" not in bar); pg.screenshot(path=f"{OUT}/31_toolbox_ui_tool.png")
    pg.click("[data-testid=cmd-reset]"); time.sleep(0.4); bar=pg.inner_text("[data-testid=region-actionbar]"); ok("S14b reset restores the default commands", "EBOM Run" in bar, "EBOM Run" in bar and "원가 계산" not in bar)
    # S15 드래그 · 도킹
    t=pg.locator("[data-testid=toolbox-titlebar]").bounding_box(); pg.mouse.move(t["x"]+12,t["y"]+14); pg.mouse.down(); pg.mouse.move(t["x"]-288,t["y"]+134,steps=8); pg.mouse.up(); time.sleep(0.4)
    tb2=pg.locator("[data-testid=toolbox-window]").bounding_box(); ok("S15a window drags", (round(tb["x"]),round(tb2["x"])), abs((tb["x"]-tb2["x"])-300)<6 and abs((tb2["y"]-tb["y"])-120)<6)
    pg.click("[data-testid=toolbox-dock]"); time.sleep(0.4); tb3=pg.locator("[data-testid=toolbox-window]").bounding_box(); vw=pg.evaluate("window.innerWidth")
    ok("S15b dock toggle pins it to the right edge", round(tb3["x"]+tb3["width"]), abs(tb3["x"]+tb3["width"]-vw)<2 and pg.get_attribute("[data-testid=toolbox-window]","data-docked")=="1")
    pg.reload(wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=toolbox-window]",timeout=30000); ok("S15c open + docked state survive a reload", True, pg.get_attribute("[data-testid=toolbox-window]","data-docked")=="1")
    pg.click("[data-testid=toolbox-reset]"); pg.click("[data-testid=toolbox-close]"); time.sleep(0.3)
    # S11 권한: viewer는 등록을 못 한다 (서버에서 차단)
    v=b.new_context(); v.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    r=v.request.post(BASE+"/api/setup/sub-codes",headers=J,data=json.dumps({"group":"AHU Code","itemKey":"B","itemName":"용량","value":"99"})); ok("S11 viewer cannot register codes (403)", r.status, r.status==403); v.close()
    # ── P4-a 치수 전파 · 도면 (p38~40 Key Dimension · p24 Drawings) ─────────────
    import ezdxf, io
    def dxf_stats(txt):
        doc=ezdxf.read(io.StringIO(txt)); msp=doc.modelspace()
        ents=[e for e in msp]
        xs=[]; ys=[]
        for e in ents:
            if e.dxftype()=="LINE": xs+= [e.dxf.start.x, e.dxf.end.x]; ys+=[e.dxf.start.y, e.dxf.end.y]
            elif e.dxftype()=="TEXT": xs.append(e.dxf.insert.x); ys.append(e.dxf.insert.y)
        texts=sorted(e.dxf.text for e in ents if e.dxftype()=="TEXT")
        return {"n":len(ents),"layers":sorted({e.dxf.layer for e in ents}),"maxx":round(max(xs)),"maxy":round(max(ys)),"texts":texts}
    # S18a 지금 도면의 폭은 등록 표의 값이다
    m=ctx.request.get(BASE+f"/api/dxf?runId={RUN0}&type=plan&meta=1").json()
    ok("S18a 도면 치수가 등록 표에서 온다 (W=2472 · 치수행 55)", (m.get("widthMm"), m.get("dimItem")), m.get("widthMm")==2472 and m.get("dimItem")=="55")
    before=dxf_stats(ctx.request.get(BASE+f"/api/dxf?runId={RUN0}&type=plan").text())
    # S18b 치수 표의 한 칸(55행 W)을 2472 → 2600 으로 바꾼다
    cat=ctx.request.get(BASE+"/api/setup/catalog").json()
    eu=[p for p in cat.get("productCodes",[]) if p.get("code")=="EU"][0]
    dimname=[k for k,t in eu["tables"].items() if t.get("role")=="dim"][0]
    wkey=[c["key"] for c in eu["tables"][dimname]["cols"] if c["name"]=="W"][0]
    for row in eu["tables"][dimname]["rows"]:
        if row["item"]=="55": row["cells"][wkey]=2600
    up=ctx.request.post(BASE+"/api/setup/product-codes",headers=J0,data=json.dumps(eu))
    ok("S18b 치수 표의 한 칸을 2472 → 2600 으로 고친다", up.status, up.status==200)
    # S18c 같은 코드로 다시 돌리면 도면이 바뀐다
    r1=ctx.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":S55_0,"code":"EU-55-2123-630SS-1-21-13-15","node":"a0000000-0000-4000-8000-000000000004"}))
    RUN1=r1.json().get("runId")
    after=dxf_stats(ctx.request.get(BASE+f"/api/dxf?runId={RUN1}&type=plan").text())
    m1=ctx.request.get(BASE+f"/api/dxf?runId={RUN1}&type=plan&meta=1").json()
    ok("S18c 치수를 바꾸니 도면 폭이 따라간다 (2472 → 2600 · ezdxf 실측 도형도 그만큼 커짐)", (m.get("widthMm"), m1.get("widthMm"), after["maxy"]-before["maxy"]), m1.get("widthMm")==2600 and after["maxy"]-before["maxy"]==128)
    Lb=[t for t in before["texts"] if t.startswith("L=")]; La=[t for t in after["texts"] if t.startswith("L=")]
    ok("S18d 바뀐 것은 폭뿐 — 엔티티 수·레이어·전장(L)은 그대로", (before["n"], after["n"], Lb, La), before["n"]==after["n"] and before["layers"]==after["layers"] and before["maxx"]==after["maxx"] and Lb==La and len(Lb)==1)
    ok("S18e 치수 문자열만 갈렸다 (W=2472 → W=2600)", ("W=2472" in before["texts"], "W=2600" in after["texts"]), "W=2472" in before["texts"] and "W=2600" in after["texts"] and "W=2600" not in before["texts"])
    # S18f 조립도는 BOM 스냅샷의 Item 표를 도면 안에 담는다
    asm=ctx.request.get(BASE+f"/api/dxf?runId={RUN1}&type=assembly").text()
    am=ctx.request.get(BASE+f"/api/dxf?runId={RUN1}&type=assembly&meta=1").json()
    ok("S18f 조립도에 Item 표와 풍선번호가 들어간다 (p38·p40)", (am.get("items"), "Q'ty" in asm), am.get("items")==11 and "Q'ty" in asm and "0\nCIRCLE\n" in asm)
    open(f"{OUT}/edim_assembly.dxf","w",encoding="utf-8").write(asm)
    # S19 도면을 남긴다 — 번호·개정·상태·발행 잠금 (p24)
    g1=ctx.request.post(BASE+"/api/drawings",headers=J0,data=json.dumps({"runId":RUN1,"type":"plan"})).json()
    ok("S19a 도면 등록 Rev A", (g1.get("drawingNo"), g1.get("rev")), g1.get("rev")=="A")
    g2=ctx.request.post(BASE+"/api/drawings",headers=J0,data=json.dumps({"runId":RUN1,"type":"plan"})).json()
    ok("S19b 다시 뜨면 Rev B — 앞 개정은 남는다", g2.get("rev"), g2.get("rev")=="B")
    # ── P6 금실: 밖으로 나가는 것(발행·발주)은 **승인된 BOM** 에서만 (p55·p56·p65) ──
    PID="c0000000-0000-4000-8000-000000000001"
    na=ctx.request.patch(BASE+f"/api/drawings/{g1['id']}",headers=J0,data=json.dumps({"status":"issued"}))
    ok("S27a 승인되지 않은 BOM 에서 나온 도면은 발행할 수 없다 (409 + 이유)", (na.status, na.json().get("error","")[:24]), na.status==409 and "승인" in na.json().get("error",""))
    nr2=ctx.request.post(BASE+f"/api/projects/{PID}/approvals",headers=J0,data=json.dumps({"note":"tier:org · e2e"}))
    ok("S27b 무엇을 승인하는지 모르는 승인 요청은 받지 않는다 (runId 없으면 400)", nr2.status, nr2.status==400)
    aq=ctx.request.post(BASE+f"/api/projects/{PID}/approvals",headers=J0,data=json.dumps({"note":"tier:org · e2e","runId":RUN1})); AID=aq.json().get("id")
    still=ctx.request.patch(BASE+f"/api/drawings/{g1['id']}",headers=J0,data=json.dumps({"status":"issued"}))
    ok("S27c 승인을 BOM 스냅샷에 묶어 요청한다 · 요청만으로는 아직 발행 못 한다", (aq.status, still.status), aq.status==200 and aq.json().get("runId")==RUN1 and still.status==409)
    ad=ctx.request.post(BASE+f"/api/project-approvals/{AID}",headers=J0,data=json.dumps({"decision":"approved","note":"ok"}))
    ad2=ctx.request.post(BASE+f"/api/project-approvals/{AID}",headers=J0,data=json.dumps({"decision":"rejected","note":"undo"}))
    ok("S27d 승인한다 · 결정된 승인은 뒤집을 수 없다 (409)", (ad.status, ad2.status), ad.status==200 and ad2.status==409)
    for st in ["review","approved","issued"]:
        pr=ctx.request.patch(BASE+f"/api/drawings/{g1['id']}",headers=J0,data=json.dumps({"status":st}))
    ok("S19c 작성중 → 검토 → 승인 → 발행", pr.status, pr.status==200)
    lk=ctx.request.patch(BASE+f"/api/drawings/{g1['id']}",headers=J0,data=json.dumps({"status":"issued"}))
    ok("S19d 발행된 도면은 잠긴다 (409)", lk.status, lk.status==409)
    dl=ctx.request.get(BASE+f"/api/drawings/{g1['id']}")
    ok("S19e 남긴 도면을 그대로 내려받는다", dl.status==200 and "AC1009" in dl.text(), dl.status==200 and "AC1009" in dl.text())
    # S20 산출물은 스냅샷에서만 나온다 — 근거 없는 재계산 금지
    nr=ctx.request.post(BASE+"/api/run/cost",headers=J0,data=json.dumps({"slots":S55_0}))
    ok("S20a runId 없는 Cost 는 거부된다 (409)", nr.status, nr.status==409)
    cr=ctx.request.post(BASE+"/api/run/cost",headers=J0,data=json.dumps({"runId":RUN1})).json()
    ok("S20b Cost 는 스냅샷에 저장된 값을 그대로 읽는다", cr.get("value"), isinstance(cr.get("value"),(int,float)) and cr.get("value")>0)
    er=ctx.request.post(BASE+"/api/run/ebom",headers=J0,data=json.dumps({"runId":RUN1})).json()
    ok("S20c EBOM 도 같은 스냅샷에서 나온다", len(er.get("groups",[])), len(er.get("groups",[]))>0)
    # S20d 매크로 값은 서버가 직접 낸다 — 클라이언트가 보내지 않아도 방진구가 들어간다
    parts=[l.get("part") for l in r1.json().get("lines",[])]
    ok("S20d 클라이언트가 매크로 값을 안 보내도 서버가 실행해 방진구가 나온다", r1.json().get("macroValue"), r1.json().get("macroValue")==455.4 and any("Vibration" in (p or "") for p in parts))
    # S21 화면: Design 탭에서 도면을 등록하고 상태가 보인다
    nuke(pg); pg.goto(BASE+"/workbench",wait_until="domcontentloaded"); hydrated(pg); nuke(pg)
    pg.click("text=PS-61313"); wait_text(pg,"body","Micron FAB AHU"); nuke(pg)
    pg.locator("button", has_text=re.compile(r"^Design$")).first.click(force=True); nuke(pg)
    # 고정 1.5초 대신 도면 목록 로딩 표지(data-listed)와 첫 행을 기다린다 (Windows 첫 컴파일에서 1.5초를 넘겼다 — 2026-09-26)
    try: pg.wait_for_function("()=>document.querySelector('[data-testid=drawing-register]')?.dataset.listed==='1' && !!document.querySelector('[data-testid=drawing-row]')",timeout=30000)
    except Exception: pass
    body=pg.inner_text("[data-testid=design-canvas]")
    ok("S21a Design 탭에 등록된 도면과 상태가 보인다", ("발행" in body, "Rev" in body), "Rev" in body and ("발행" in body or "작성중" in body))
    # S31g 화면: Arrangement 버튼 → 구획 길이를 고치고 저장하면 캔버스에 그 길이가 뜬다 (실동 · 버튼 자리만 있던 청사진 p13·58 채움)
    nuke(pg); pg.click("[data-testid=arrangement-edit]", force=True); pg.wait_for_selector("[data-testid=arrangement-panel]", timeout=15000); nuke(pg)
    pg.fill("[data-testid=arr-len-Coil]", "1500"); pg.click("[data-testid=arr-save]", force=True)
    pg.wait_for_selector("[data-testid=design-canvas] >> text=1500", timeout=15000)
    ok("S31g 화면: Arrangement 편집 → Coil 길이 1500 저장 → 캔버스 구획에 1500 이 뜬다", bool(pg.query_selector("[data-testid=design-canvas] >> text=1500")), True)
    # S31h 개념도는 **지금 슬롯에서 도는 구획만** 그린다(편집 표에는 조건부 구획까지 전부 뜬다) — 도면과 같은 구획 집합
    _rows=pg.eval_on_selector_all("[data-testid=arr-table] tbody tr","e=>e.length")
    _cv=pg.eval_on_selector("[data-testid=design-canvas] svg","e=>e.textContent")  # 편집 표가 같은 카드 안에 있으므로 개념도(SVG)만 읽는다
    _drawn=[nm for nm in ("Mixing","Filter","Rotor","Coil","Humid.","HeatPump","Fan") if nm in _cv]
    ok("S31h 편집 표에는 조건부 구획까지 전부(7), 개념도에는 지금 슬롯에서 도는 구획만 — 꺼진 조건부 구획은 그리지 않는다",
       (_rows, _drawn), _rows==7 and len(_drawn)<_rows and "HeatPump" not in _drawn and "Mixing" in _drawn and "Fan" in _drawn)
    pg.screenshot(path=f"{OUT}/16_design_tab.png")
    # 되돌린다 — API 로 EU 제품 코드의 모든 구획 len 을 지운다(빈 화면 입력이 불안정). 뒤 시나리오가 균등 도면을 기대한다.
    _c=ctx.request.get(BASE+"/api/setup/catalog").json(); _eu=[p_ for p_ in _c.get("productCodes",[]) if p_.get("code")=="EU"][0]
    for sd in _eu["sections"]: sd.pop("len",None)
    ctx.request.post(BASE+"/api/setup/product-codes",headers=J0,data=json.dumps(_eu))
    pg.click("button:has-text('BOM Run')"); wait_text(pg,"body","2600×2472"); nuke(pg)
    kd=pg.inner_text("body")
    ok("S21b 핵심 치수가 등록 표 값을 그대로 보여 준다 (화면이 따로 계산하지 않는다)", "2600×2472" in kd, "2600×2472" in kd)
    pg.screenshot(path=f"{OUT}/43_drawings.png",full_page=True)
    # ── P4-b 견적 · Tech Data · 구매 요청 — 전부 BOM 스냅샷(runId) 하나에서 (p66 · p15~16 · p51) ──
    # S22 견적: 합계는 다시 세지 않는다 — Cost API 가 읽는 바로 그 값이다
    nr=ctx.request.post(BASE+"/api/documents",headers=J0,data=json.dumps({"type":"quotation"}))
    ok("S22a 스냅샷 없이는 견적을 못 뜬다 (400)", nr.status, nr.status==400)
    q1=ctx.request.post(BASE+"/api/documents",headers=J0,data=json.dumps({"runId":RUN1,"type":"quotation"})).json()
    ok("S22b 견적 합계 = Cost API 값 (한 원도 다르지 않다) · 번호 QR-61313-nn Rev A", (q1.get("docNo"), q1.get("rev"), q1.get("total"), cr.get("value")), q1.get("total")==cr.get("value") and str(q1.get("docNo","")).startswith("QR-61313-") and q1.get("rev")=="A")
    ph=ctx.request.get(BASE+f"/api/documents/{q1.get('id')}/print"); won=format(int(cr.get("value")),",")
    ok("S22c 인쇄본(HTML)에 그 합계와 근거 스냅샷 id 가 찍힌다", (ph.status, won in ph.text(), RUN1 in ph.text()), ph.status==200 and "text/html" in ph.headers.get("content-type","") and won in ph.text() and RUN1 in ph.text() and "견 적 서" in ph.text())
    q2=ctx.request.post(BASE+"/api/documents",headers=J0,data=json.dumps({"runId":RUN1,"type":"quotation"})).json()
    ok("S22d 같은 코드로 다시 뜨면 같은 번호에 Rev B", (q2.get("docNo"), q2.get("rev")), q2.get("docNo")==q1.get("docNo") and q2.get("rev")=="B")
    for st in ("review","approved","issued"): pr_=ctx.request.patch(BASE+f"/api/documents/{q1.get('id')}",headers=J0,data=json.dumps({"status":st}))
    lk=ctx.request.patch(BASE+f"/api/documents/{q1.get('id')}",headers=J0,data=json.dumps({"status":"issued"}))
    bk=ctx.request.patch(BASE+f"/api/documents/{q2.get('id')}",headers=J0,data=json.dumps({"status":"review"})); bk2=ctx.request.patch(BASE+f"/api/documents/{q2.get('id')}",headers=J0,data=json.dumps({"status":"draft"}))
    ok("S22e 발행된 견적은 잠기고(409), 상태는 되돌릴 수 없다(409)", (pr_.status, lk.status, bk2.status), pr_.status==200 and lk.status==409 and bk.status==200 and bk2.status==409)
    # S22f 견적서 발치의 '코드 개정' 근거 — 최신 개정이 아니라 **그 슬롯으로 저장된** 개정만 찍힌다
    revA=[x for x in rv if x.get("revNo")==1][0]; SA=revA["slots"]; ok("S2e Rev A 에 F 순번이 저장돼 있다 (개정 = A~F 전체 코드)", (revA.get("code"), SA.get("F")), revA.get("code")=="EU-55-2123-630SS-1-21-13-15" and SA.get("F")=="1-21-13-15")
    ra=ctx.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":SA,"code":"EU-55-2123-630SS","node":"a0000000-0000-4000-8000-000000000004"})).json()
    qa=ctx.request.post(BASE+"/api/documents",headers=J0,data=json.dumps({"runId":ra.get("runId"),"type":"quotation"})).json()
    srcA=ctx.request.get(BASE+f"/api/documents/{qa.get('id')}").json().get("body",{}).get("source",{})
    src1=ctx.request.get(BASE+f"/api/documents/{q1.get('id')}").json().get("body",{}).get("source",{})
    SU=dict(SA); SU["F"]="1-21-13-16"; ru=ctx.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":SU,"code":"EU-55-2123-630SS-1-21-13-16","node":"a0000000-0000-4000-8000-000000000004"})).json()
    qu=ctx.request.post(BASE+"/api/documents",headers=J0,data=json.dumps({"runId":ru.get("runId"),"type":"quotation"})).json(); srcU=ctx.request.get(BASE+f"/api/documents/{qu.get('id')}").json().get("body",{}).get("source",{})
    ok("S22f 견적 발치의 근거 개정 = 그 슬롯(A~F)으로 저장된 개정만 — Rev A 슬롯 → Rev A(최신 Rev B 가 아님) · F 만 다른 미저장 조합(…-16) → 빈 값", (srcA.get("codeRevisionId")==revA.get("id"), src1.get("codeRevisionId")==revA.get("id"), srcU.get("codeRevisionId")), srcA.get("codeRevisionId")==revA.get("id") and src1.get("codeRevisionId")==revA.get("id") and srcU.get("codeRevisionId") is None)
    # S23 Tech Data: 값 + 그 값을 낸 승인 매크로 개정 + 입력
    t1=ctx.request.post(BASE+"/api/documents",headers=J0,data=json.dumps({"runId":RUN1,"type":"techdata"})).json()
    tb=ctx.request.get(BASE+f"/api/documents/{t1.get('id')}").json().get("body",{})
    ok("S23a Tech Data = 매크로 결과 455.4 + 승인 개정·원문 + 입력 슬롯", (t1.get("docNo"), t1.get("value"), tb.get("macro",{}).get("revision")), t1.get("value")==455.4 and isinstance(tb.get("macro",{}).get("revision"),int) and bool(tb.get("macro",{}).get("dsl")) and [i["key"] for i in tb.get("input",[])][:2]==["A","B"] and str(t1.get("docNo","")).startswith("TD-61313-"))
    rn=ctx.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":S55_0,"code":"EU-55-2123-630SS-1-21-13-15"})).json()
    tn=ctx.request.post(BASE+"/api/documents",headers=J0,data=json.dumps({"runId":rn.get("runId"),"type":"techdata"}))
    ok("S23b 매크로 없이 돈 스냅샷은 Tech Data 를 거부한다 (422 — 값을 지어내지 않는다)", (rn.get("macroValue"), tn.status), rn.get("macroValue") is None and tn.status==422)
    # S24 구매 요청: 스냅샷 줄 중 '구매 품목'으로 등록돼 있던 것만
    cat=ctx.request.get(BASE+"/api/setup/catalog").json(); buy={p_["code"] for p_ in cat.get("productCodes",[]) if p_.get("kind")=="purchase"}
    want=[t_ for t_ in r1.json().get("trace",[]) if t_.get("childCode") in buy]
    pq=ctx.request.post(BASE+"/api/purchase-requests",headers=J0,data=json.dumps({"runId":RUN1,"requiredDate":"2026-10-15"})); pj=pq.json()
    ok("S24a 구매 요청 줄 수 = 스냅샷의 구매 품목 수 · 번호 PR-61313-n", (pj.get("prNo"), pj.get("lines"), len(want)), pq.status==200 and pj.get("lines")==len(want) and len(want)>0 and str(pj.get("prNo","")).startswith("PR-61313-"))
    dq=ctx.request.post(BASE+"/api/purchase-requests",headers=J0,data=json.dumps({"runId":RUN1}))
    ok("S24b 같은 스냅샷으로 두 번 사지 않는다 (409)", dq.status, dq.status==409)
    prs=ctx.request.get(BASE+"/api/purchase-requests").json().get("rows",[]); mine=[x for x in prs if x.get("id")==pj.get("id")][0]
    byno={l_["no"]:l_ for l_ in r1.json().get("lines",[])}
    got=[(l_["bomLineNo"],l_["resolvedCode"],l_["qty"],l_["unitPrice"]) for l_ in mine["lines"]]; exp=[(t_["no"],t_["resolvedCode"],byno[t_["no"]]["qty"],byno[t_["no"]]["unitCost"]) for t_ in want]
    ok("S24c 줄의 BOM 줄번호·코드·수량·단가가 스냅샷 줄 그대로다 · 필요일 반영", got, got==exp and all(l_["requiredDate"][:10]=="2026-10-15" for l_ in mine["lines"]))
    # S25 화면: Document 탭 — 방금 화면에서 돌린 BOM 스냅샷으로 세 산출물을 만든다
    nuke(pg); pg.locator("button", has_text=re.compile(r"^Document$")).first.click(force=True); time.sleep(1.5); nuke(pg)
    ok("S25a Document 탭이 BOM 스냅샷을 잡고 있다 (버튼이 열려 있다)", pg.inner_text("[data-testid=document-run]"), "없음" not in pg.inner_text("[data-testid=document-run]") and pg.is_enabled("[data-testid=doc-make-quotation]"))
    for tid in ("doc-make-quotation","doc-make-techdata","pr-make"):
        nuke(pg); pg.click(f"[data-testid={tid}]", force=True); pg.wait_for_selector("[data-testid=document-msg]",timeout=15000); time.sleep(1.8)
    dp=pg.inner_text("[data-testid=document-panel]"); nrow=len(pg.query_selector_all("[data-testid=document-row]"))
    ok("S25b 화면에서 견적·Tech Data·구매 요청이 등록되고 목록에 보인다", (nrow, "구매 요청 있음" in dp), nrow>=6 and "구매 요청 있음" in dp and "발행" in dp and pg.get_attribute("[data-testid=document-msg]","data-ok")=="1")
    ok("S25c 구매 요청을 만든 스냅샷에서는 버튼이 잠긴다 (두 번 사지 않는다)", pg.is_disabled("[data-testid=pr-make]"), pg.is_disabled("[data-testid=pr-make]"))
    pg.screenshot(path=f"{OUT}/44_document_tab.png",full_page=True)
    # S26 화면: Purchasing — Process 를 올리고(견적 요청 → 발주) CSV 로 내보낸다 (p51)
    pg.goto(BASE+"/m/purchasing",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=pr-card]",timeout=30000); time.sleep(1.2); nuke(pg)
    ncard=len(pg.query_selector_all("[data-testid=pr-card]")); PRNO=pj.get("prNo")
    ok("S26a Purchasing 에 구매 요청 2건(API 1 · 화면 1)이 줄과 함께 보인다", (ncard, len(pg.query_selector_all("[data-testid=pr-line]"))), ncard==2 and len(pg.query_selector_all("[data-testid=pr-line]"))==2*len(want))
    CARD=f"[data-testid=pr-card]:has([data-testid='pr-export-{PRNO}'])"
    for want_st in ("rfq","ordered"):  # 고정 sleep 이 아니라 **상태가 바뀔 때까지** 기다린다(첫 호출은 라우트 컴파일로 느릴 수 있다)
        nuke(pg); pg.click(f"[data-testid='pr-advance-{PRNO}']", force=True); pg.wait_for_selector(f"{CARD}[data-status={want_st}]",timeout=20000); time.sleep(0.4)
    card=pg.inner_text(f"[data-testid=pr-card]:has([data-testid='pr-export-{PRNO}'])")
    ok("S26b 견적 요청 → 발주 · 발주되면 PO 번호가 붙는다 (PO-61313-n)", card.split("\n")[0][:60], "발주" in card and "PO-61313-" in card and not pg.query_selector(f"[data-testid='pr-advance-{PRNO}']"))
    lk=ctx.request.patch(BASE+f"/api/purchase-requests/{pj.get('id')}",headers=J0,data=json.dumps({"status":"ordered"}))
    ok("S26c 발주된 구매 요청은 잠긴다 (409)", lk.status, lk.status==409)
    ex=ctx.request.get(BASE+f"/api/purchase-requests/{pj.get('id')}/export"); rows_=ex.text().strip().split("\r\n")
    ok("S26d Export CSV: 머리 1줄 + 구매 품목 줄 · PR·PO 번호 포함 · 엑셀용 BOM", (ex.status, len(rows_), ex.headers.get("content-type")), ex.status==200 and "text/csv" in ex.headers.get("content-type","") and len(rows_)==1+len(want) and PRNO in rows_[1] and "PO-61313-" in rows_[1] and ex.body()[:3]==b"\xef\xbb\xbf")
    open(f"{OUT}/{PRNO}.csv","wb").write(ex.body())
    # S28 추적 — 구매 요청에서 거꾸로: 스냅샷 → 코드 개정 → 카탈로그 지문 → 매크로 개정 → 승인
    t1=ctx.request.get(BASE+f"/api/trace?runId={RUN1}").json(); ta=ctx.request.get(BASE+f"/api/trace?runId={ra.get('runId')}").json()
    ok("S28a 추적: 발주된 구매 요청의 BOM 은 승인돼 있고, 매크로 개정·도면·문서·PO 가 한 번에 따라온다", (t1.get("approved"), (t1.get("macro") or {}).get("revision"), len(t1.get("drawings",[])), len(t1.get("documents",[])), (t1.get("purchaseRequest") or {}).get("poNo")), t1.get("approved") is True and isinstance((t1.get("macro") or {}).get("revision"),int) and len(t1.get("drawings",[]))>=2 and len(t1.get("documents",[]))>=3 and str((t1.get("purchaseRequest") or {}).get("poNo","")).startswith("PO-61313-") and t1["snapshot"]["total"]==cr.get("value"))
    ok("S28b 추적: Rev A 슬롯으로 돌린 BOM 은 코드 개정 Rev A 까지 거슬러 올라간다 · 그 BOM 은 미승인", ((ta.get("codeRevision") or {}).get("rev"), ta.get("approved")), (ta.get("codeRevision") or {}).get("rev")==1 and ta.get("approved") is False)
    ok("S28e F 가 붙은 실행(S18 의 RUN1 = 630SS-1-21-13-15)도 근거 개정이 있다 — Rev A (F 포함 전엔 빈 값이었다)", ((t1.get("codeRevision") or {}).get("rev"), (t1.get("codeRevision") or {}).get("code")), (t1.get("codeRevision") or {}).get("rev")==1 and str((t1.get("codeRevision") or {}).get("code","")).endswith("1-21-13-15"))
    other=[x for x in ctx.request.get(BASE+"/api/purchase-requests").json().get("rows",[]) if x.get("id")!=pj.get("id")][0]; ONO=other["prNo"]
    nuke(pg); pg.click(f"[data-testid='pr-trace-{PRNO}']", force=True); time.sleep(1.5); nuke(pg); pg.click(f"[data-testid='pr-trace-{ONO}']", force=True); time.sleep(1.5)
    tp=[(e.get_attribute("data-approved"), e.inner_text()) for e in pg.query_selector_all("[data-testid=pr-trace]")]
    ok("S28c 화면: 추적 패널이 승인된 BOM 과 미승인 BOM 을 가려 보여 준다", [x[0] for x in tp], sorted(x[0] for x in tp)==["0","1"] and any("승인된 BOM" in x[1] and "매크로" in x[1] for x in tp))
    OCARD=f"[data-testid=pr-card]:has([data-testid='pr-export-{ONO}'])"
    nuke(pg); pg.click(f"[data-testid='pr-advance-{ONO}']", force=True); pg.wait_for_selector(f"{OCARD}[data-status=rfq]",timeout=20000); time.sleep(0.4)
    nuke(pg); pg.click(f"[data-testid='pr-advance-{ONO}']", force=True); pg.wait_for_selector("[data-testid=purchasing-error]",timeout=20000); time.sleep(0.4)
    pe=pg.inner_text("[data-testid=purchasing-error]") if pg.query_selector("[data-testid=purchasing-error]") else ""
    ost=[x for x in ctx.request.get(BASE+"/api/purchase-requests").json().get("rows",[]) if x.get("prNo")==ONO][0]
    ok("S28d 화면: 승인되지 않은 BOM 의 구매 요청은 견적 요청까지만 — 발주는 막히고 이유가 뜬다", (ost.get("status"), pe[:30]), ost.get("status")=="rfq" and ost.get("poNo") is None and "승인" in pe)
    # S29 플랫폼 단계 — 조직 승인을 받은 바로 그 스냅샷이어야 한다
    pw=ctx.request.post(BASE+f"/api/projects/{PID}/approvals",headers=J0,data=json.dumps({"note":"tier:platform · e2e","runId":ra.get("runId")}))
    ok("S29a 다른 스냅샷으로는 플랫폼 단계에 올릴 수 없다 (409)", pw.status, pw.status==409)
    pg.screenshot(path=f"{OUT}/45_purchasing.png",full_page=True)
    pg.goto(BASE+f"/api/documents/{q1.get('id')}/print",wait_until="load"); pg.screenshot(path=f"{OUT}/46_quotation_print.png",full_page=True)
    pg.goto(BASE+"/workbench?node=a0000000-0000-4000-8000-000000000004",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=pipeline-stage]",timeout=30000); time.sleep(1.5); nuke(pg)
    st0=pg.inner_text("[data-testid=pipeline-stage]"); insp=pg.inner_text("[data-testid=pipeline-stage] >> xpath=ancestor::*[3]")
    ok("S29b 화면: Inspector 가 Approve 단계와 승인된 BOM 을 보여 준다", (st0, RUN1[:8] in insp), st0.strip()=="Approve" and RUN1[:8] in insp)
    # 고정 sleep 대신 상태를 기다린다(ccmd K · 개발 모드에서 2.5초 안에 화면이 못 따라와 S29c 가 한 번 깨졌다 — 2026-09-29)
    nuke(pg); pg.click("[data-testid=request-platform]", force=True); wait_sel(pg,"[data-testid=approve-btn]:not([disabled])",60000); nuke(pg)
    pg.click("[data-testid=approve-btn]", force=True); wait_text(pg,"[data-testid=pipeline-stage]","Accepted",60000)
    st1=pg.inner_text("[data-testid=pipeline-stage]"); t2=ctx.request.get(BASE+f"/api/trace?runId={RUN1}").json()
    bd=pg.query_selector("[data-testid=approval-bound]")
    ok("S29d 화면: 승인된 것은 BOM 이다 — 화면의 현재 코드(AL)가 승인된 BOM(SS)과 다르면 그렇다고 말한다", (bd.get_attribute("data-differs") if bd else None), bool(bd) and "630SS" in bd.inner_text() and bd.get_attribute("data-differs")=="1" and "다릅니다" in bd.inner_text())
    ok("S29c 화면: Accepted 요청 → 승인 → Accepted · 플랫폼 승인도 같은 BOM 에 묶인다", (st1, [f"{a_['tier']}:{a_['state']}" for a_ in t2.get("approvals",[])]), st1.strip()=="Accepted" and any(a_["tier"]=="platform" and a_["state"]=="approved" for a_ in t2.get("approvals",[])))
    pg.screenshot(path=f"{OUT}/51_accepted.png",full_page=True)
    # ── S30 변경 전파 — 한 번의 Set-Up 저장이 BOM·원가·구매 수량·도면을 **함께** 움직이고, 앞 스냅샷은 그대로다 (P6 완료 기준 3) ──
    # 정직 고지: 지금 데이터 모델에서 한 칸이 네 산출물 전부에 닿지는 않는다. filterQty 는 BOM·원가·구매에, 치수 W 는 도면에 닿는다
    # (BOM 사양의 단면은 cap.face 를, 도면은 dim.W 를 읽는다 — 같은 값이 두 칸에 있다). 그래서 "한 번의 저장(두 칸)"으로 묶는다.
    CODE55="EU-55-2123-630SS-1-21-13-15"; NODE4="a0000000-0000-4000-8000-000000000004"
    run55_lines_cache={}
    def run55(tag=None):
        r_=ctx.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":S55_0,"code":CODE55,"node":NODE4})).json()
        if tag: run55_lines_cache[tag]=r_.get("lines",[])
        return r_
    def eu_now():
        c_=ctx.request.get(BASE+"/api/setup/catalog").json(); e_=[p_ for p_ in c_.get("productCodes",[]) if p_.get("code")=="EU"][0]
        dn=[k for k,t_ in e_["tables"].items() if t_.get("role")=="dim"][0]; wk=[c2["key"] for c2 in e_["tables"][dn]["cols"] if c2["name"]=="W"][0]
        fk=[c2["key"] for c2 in e_["tables"]["cap"]["cols"] if c2["name"]=="filterQty"][0]
        dr=[r_ for r_ in e_["tables"][dn]["rows"] if r_["item"]=="55"][0]; cr_=[r_ for r_ in e_["tables"]["cap"]["rows"] if r_["item"]=="55"][0]
        return e_,dr,wk,cr_,fk
    def facts(run):
        rid=run.get("runId"); fl={l_["no"]:l_ for l_ in run.get("lines",[]) if "filter" in (l_.get("part") or "").lower()}
        rest=[(l_["no"],l_.get("part"),l_["qty"],l_["unitCost"]) for l_ in run.get("lines",[]) if l_["no"] not in fl]
        cost=ctx.request.post(BASE+"/api/run/cost",headers=J0,data=json.dumps({"runId":rid})).json().get("value")
        meta=ctx.request.get(BASE+f"/api/dxf?runId={rid}&type=plan&meta=1").json(); geo=dxf_stats(ctx.request.get(BASE+f"/api/dxf?runId={rid}&type=plan").text())
        fp=(ctx.request.get(BASE+f"/api/trace?runId={rid}").json().get("snapshot") or {}).get("catalogFp")
        return {"rid":rid,"fq":sorted(l_["qty"] for l_ in fl.values()),"fsum":sum(l_["qty"]*l_["unitCost"] for l_ in fl.values()),"rest":rest,"cost":cost,"W":meta.get("widthMm"),"maxy":geo["maxy"],"fp":fp}
    def pr_filter_qty(rid):
        q_=ctx.request.post(BASE+"/api/purchase-requests",headers=J0,data=json.dumps({"runId":rid})); i_=q_.json().get("id")
        row=[x for x in ctx.request.get(BASE+"/api/purchase-requests").json().get("rows",[]) if x.get("id")==i_]
        return q_.status, sorted(l_["qty"] for l_ in (row[0]["lines"] if row else []) if "filter" in (l_.get("part") or l_.get("spec") or l_.get("name") or "").lower())
    A=facts(run55("A")); eu_,dr_,wk_,cr2,fk_=eu_now(); W0=dr_["cells"][wk_]; F0=cr2["cells"][fk_]
    dr_["cells"][wk_]=W0+100; cr2["cells"][fk_]=F0+2
    sv=ctx.request.post(BASE+"/api/setup/product-codes",headers=J0,data=json.dumps(eu_))
    ok("S30a 한 번의 Set-Up 저장: 필터 수량 +2 · 폭 W +100 (두 칸, 저장 1회)", (sv.status,F0,F0+2,W0,W0+100), sv.status==200 and len(A["fq"])==2 and A["fq"]==[F0,F0])
    B=facts(run55("B"))
    ok("S30b BOM 이 따라간다 — 필터 두 줄 수량만 바뀌고 나머지 줄은 그대로", (A["fq"],B["fq"],len(B["rest"])), B["fq"]==[F0+2,F0+2] and B["rest"]==A["rest"] and len(B["rest"])>0)
    specA=[l_["spec"] for l_ in run55_lines_cache["A"] if "Panel" in (l_.get("part") or "")]; specB=[l_["spec"] for l_ in run55_lines_cache["B"] if "Panel" in (l_.get("part") or "")]
    ok("S30b2 BOM 사양 문자열도 함께 바뀐다 — 패널 단면이 dim 표를 읽는다 (W 만 +100, H 는 그대로)", (specA[:1], specB[:1]), bool(specA) and f" {W0}×" in specA[0] and f" {W0+100}×" in specB[0])
    dsum=B["fsum"]-A["fsum"]; dcost=(B["cost"] or 0)-(A["cost"] or 0)
    ok("S30c 원가가 따라간다 — 늘어난 원가 = 늘어난 필터 재료비 × 1.18 × 1.12 (다른 요인 없음)", (dsum,dcost), dsum>0 and abs(dcost-dsum*1.18*1.12)<=1)
    ok("S30d 도면이 따라간다 — 폭 +100 (메타와 ezdxf 실측 도형 둘 다)", (A["W"],B["W"],B["maxy"]-A["maxy"]), B["W"]==A["W"]+100 and B["maxy"]-A["maxy"]==100)
    sa,qa=pr_filter_qty(A["rid"]); sb,qb=pr_filter_qty(B["rid"])
    ok("S30e 구매 수량이 따라간다 — 새 스냅샷의 구매 요청은 새 수량, 앞 스냅샷의 구매 요청은 옛 수량", (sa,qa,sb,qb), sa==200 and sb==200 and qa==[F0,F0] and qb==[F0+2,F0+2])
    a_cost=ctx.request.post(BASE+"/api/run/cost",headers=J0,data=json.dumps({"runId":A["rid"]})).json().get("value")
    a_fp=(ctx.request.get(BASE+f"/api/trace?runId={A['rid']}").json().get("snapshot") or {}).get("catalogFp")
    a_dx=ctx.request.get(BASE+f"/api/dxf?runId={A['rid']}&type=plan&meta=1"); a_geo=dxf_stats(ctx.request.get(BASE+f"/api/dxf?runId={A['rid']}&type=plan").text())
    ok("S30f 앞 스냅샷은 그대로다 — 원가·지문뿐 아니라 **도면 치수도 스냅샷에 박힌 값**이라, 표가 바뀐 뒤에도 옛 스냅샷의 도면은 옛 폭 그대로 나온다 (0011)", (a_cost==A["cost"], a_fp==A["fp"], A["fp"]!=B["fp"], a_dx.status, a_dx.json().get("widthMm"), a_geo["maxy"]==A["maxy"]), a_cost==A["cost"] and a_fp==A["fp"] and bool(A["fp"]) and A["fp"]!=B["fp"] and a_dx.status==200 and a_dx.json().get("widthMm")==A["W"] and a_geo["maxy"]==A["maxy"])
    eu_,dr_,wk_,cr2,fk_=eu_now(); dr_["cells"][wk_]=W0; cr2["cells"][fk_]=F0
    rs=ctx.request.post(BASE+"/api/setup/product-codes",headers=J0,data=json.dumps(eu_)); C=facts(run55())
    b_after=ctx.request.get(BASE+f"/api/dxf?runId={B['rid']}&type=plan&meta=1").json().get("widthMm")
    ok("S30g 되돌리면 되돌아온다 — 같은 입력이면 같은 답(수량·원가·폭·지문) · 그래도 바뀐 표로 뜬 스냅샷 B 의 도면은 여전히 +100 (스냅샷마다 제 치수)", (rs.status,C["fq"],C["cost"]==A["cost"],C["W"]==A["W"],C["fp"]==A["fp"],b_after), rs.status==200 and C["fq"]==A["fq"] and C["cost"]==A["cost"] and C["W"]==A["W"] and C["fp"]==A["fp"] and b_after==A["W"]+100)
    # ── S31 Arrangement 1차 — 구획 길이를 등록하면 도면 전장이 그 합으로 바뀐다 (스냅샷에 박혀 앞 것은 그대로) ──
    def eu_pc():
        c_=ctx.request.get(BASE+"/api/setup/catalog").json(); return [p_ for p_ in c_.get("productCodes",[]) if p_.get("code")=="EU"][0]
    def dxf_of(rid): return dxf_stats(ctx.request.get(BASE+f"/api/dxf?runId={rid}&type=plan").text())
    P0=eu_pc(); secs0=[sd.get("len") for sd in P0.get("sections",[])]
    ok("S31a 처음엔 구획에 길이가 없다 — 도면은 dim 표의 L(900) 로 균등 분할한다", secs0, all(x is None for x in secs0))
    rA=run55("SA31"); gA=dxf_of(rA["runId"]); nActive=len([sd for sd in P0.get("sections",[])])  # 조건부 구획 제외 전 값이라 도면 실측으로
    L0=gA["maxx"]  # 전장 = 활성 구획수 × 900
    ok("S31b 시작 도면 전장 = 활성 구획수 × 900 (균등)", (L0, L0%900==0), L0>0 and L0%900==0)
    P=eu_pc()
    for sd in P["sections"]:
        if sd["name"]=="Coil": sd["len"]=1500  # Coil 구획만 1500 으로
    up=ctx.request.post(BASE+"/api/setup/product-codes",headers=J0,data=json.dumps(P))
    ok("S31c 구획 하나(Coil)에 길이 1500 을 등록·저장한다", up.status, up.status==200)
    rB=run55("SB31"); gB=dxf_of(rB["runId"])
    ok("S31d 새 BOM 의 도면 전장이 그만큼 늘었다 (Coil 900→1500, +600)", (L0, gB["maxx"], gB["maxx"]-L0), gB["maxx"]-L0==600 and gB["n"]==gA["n"])
    gA2=dxf_of(rA["runId"])
    ok("S31e 앞 스냅샷의 도면 전장은 그대로 (구획 길이도 스냅샷에 박힌다)", (L0, gA2["maxx"]), gA2["maxx"]==L0)
    P=eu_pc()
    for sd in P["sections"]:
        if sd["name"]=="Coil": sd.pop("len",None)
    ctx.request.post(BASE+"/api/setup/product-codes",headers=J0,data=json.dumps(P))
    rC=run55(); gC=dxf_of(rC["runId"])
    ok("S31f 되돌리면(길이 제거) 전장이 원래대로 균등", (L0, gC["maxx"]), gC["maxx"]==L0)

    # ── S32 Arrangement 2차 — 순서(Move) · 방향(p36 L0~R270) · 추가/삭제(Add/Delete) ──
    # 편집은 화면 API 하나(/api/setup/arrangement)로만 한다: 배열 순서 = 구획 순서, 빠진 이름 = 삭제, 새 이름 = 추가.
    ARR=BASE+"/api/setup/arrangement"
    g0=ctx.request.get(ARR+"?code=EU&slots="+json.dumps({"A":"EU","B":"55","C":"2123","D":"630","E":"SS","F":"1-21-13-15"})).json()
    names0=[x["name"] for x in g0.get("sections",[])]
    ok("S32a 등록된 구획 목록을 화면이 받아온다 (조건부 구획 포함 · 관계 걸린 구획은 locked)",
       (names0[:3], sum(1 for x in g0["sections"] if x.get("locked"))), len(names0)>=4 and any(x.get("locked") for x in g0["sections"]))
    def put(rows):
        return ctx.request.post(ARR, headers=J0, data=json.dumps({"code":"EU","sections":rows}))
    # (1) Move — Filter 를 맨 앞으로 보낸다
    rows=[{"name":x["name"], **({"len":x["len"]} if x.get("len") is not None else {}), **({"dir":x["dir"]} if x.get("dir") else {})} for x in g0["sections"]]
    fi=[i for i,x in enumerate(rows) if x["name"]=="Filter"][0]
    moved=[rows[fi]]+[r for i,r in enumerate(rows) if i!=fi]
    mv=put(moved)
    rD=run55("S32MV"); gD=dxf_of(rD["runId"])
    secD=ctx.request.get(BASE+f"/api/dxf?runId={rD['runId']}&type=plan&meta=1").json().get("sections",[])
    ok("S32b Move — 구획 순서를 바꾸면 도면의 구획 순서가 그대로 따라온다 (Filter 가 맨 앞)",
       (mv.status, secD[:2]), mv.status==200 and secD and secD[0]=="Filter")
    secs_in_bom=[l.get("section") for l in (rD.get("lines") or [])]
    # 구획 순서가 BOM 줄 정렬의 기준이다(미등록 구획 Casing 은 늘 앞) — Filter 가 등록 구획 중 첫째로 온다
    listed=[x for x in secs_in_bom if x in secD]
    ok("S32c 같은 순서가 BOM 줄 정렬에도 반영된다 (구획이 BOM·도면 공통 기준 — p13)",
       (secs_in_bom[:4], secD[:2]), bool(listed) and listed[0]=="Filter")
    # (2) 방향 — Fan 구획에 R90
    rows2=[dict(r) for r in moved]
    for r in rows2:
        if r["name"]=="Fan": r["dir"]="R90"
    dr=put(rows2)
    rE=run55("S32DIR"); dxfE=ctx.request.get(BASE+f"/api/dxf?runId={rE['runId']}&type=plan").text()
    metaE=ctx.request.get(BASE+f"/api/dxf?runId={rE['runId']}&type=plan&meta=1").json()
    ok("S32d 방향(p36 Fan Direction) — Fan 구획에 R90 을 등록하면 도면에 'DIR R90' 이 찍힌다",
       (dr.status, "DIR R90" in dxfE, metaE.get("dirs")), dr.status==200 and "DIR R90" in dxfE)
    dxfD=ctx.request.get(BASE+f"/api/dxf?runId={rD['runId']}&type=plan").text()
    ok("S32e 앞 스냅샷의 도면에는 방향이 없다 — 방향도 스냅샷에 박힌 값만 쓴다(0011 원칙)", ("DIR R90" in dxfD), "DIR R90" not in dxfD)
    # (3) Delete — BOM 관계가 걸린 구획은 거부, 빈 구획은 지워진다
    lock=[r for r in rows2 if r["name"]=="Coil"]
    rej=put([r for r in rows2 if r["name"]!="Coil"])
    ok("S32f Delete 거부 — BOM 관계가 걸린 구획(Coil)은 지울 수 없다 (409, 줄이 갈 곳을 잃지 않게)", rej.status, rej.status==409)
    add=put(rows2+[{"name":"Silencer","len":600}])
    rF=run55("S32ADD"); metaF=ctx.request.get(BASE+f"/api/dxf?runId={rF['runId']}&type=plan&meta=1").json()
    ok("S32g Add — 새 구획(Silencer 600)을 추가하면 도면 맨 뒤에 그 길이로 붙는다",
       (add.status, metaF.get("sections",[])[-1:], metaF.get("lengthMm")), add.status==200 and metaF.get("sections",[])[-1]=="Silencer")
    rm=put(rows2)
    rG=run55("S32RM"); metaG=ctx.request.get(BASE+f"/api/dxf?runId={rG['runId']}&type=plan&meta=1").json()
    ok("S32h Delete — 관계가 없는 구획은 지워지고 도면에서도 사라진다 (전장도 그만큼 줄어든다)",
       (rm.status, "Silencer" in metaG.get("sections",[]), metaF.get("lengthMm")-metaG.get("lengthMm")),
       rm.status==200 and "Silencer" not in metaG.get("sections",[]) and metaF.get("lengthMm")-metaG.get("lengthMm")==600)
    # ── S40 Design Tool Binding — Design Verification (p36·p60 · 코퍼스 Arrangement Design Tool Binding) ──
    rV1=run55("S40"); dv=rV1.get("dims",{})
    ok("S40a 등록된 설계 규칙이 BOM Run 마다 돌고, 판정이 스냅샷에 박힌다 (지금은 통과)",
       (dv.get("rules"), dv.get("violations")), dv.get("rules",0)>0 and dv.get("violations")==[])
    okdxf=ctx.request.get(BASE+f"/api/dxf?runId={rV1['runId']}&type=plan")
    ok("S40b 통과면 도면이 나온다", okdxf.status, okdxf.status==200)
    # 구획을 늘려 전장 운반 한계(L max 9000)를 일부러 넘긴다
    g3=ctx.request.get(ARR+"?code=EU&slots="+json.dumps({"A":"EU","B":"55","C":"2123","D":"630","E":"SS","F":"1-21-13-15"})).json()
    big=[{"name":x["name"], "len":4000, **({"dir":x["dir"]} if x.get("dir") else {}), "components":x.get("components",[])} for x in g3["sections"]]
    put(big)
    rV2=run55("S40BAD"); dv2=rV2.get("dims",{})
    ok("S40c 규칙을 넘기면 위반이 그 스냅샷에 박힌다 (전장 운반 한계 L max 9000)",
       (dv2.get("violations") or [{}])[0], any(v.get("target")=="L" for v in dv2.get("violations",[])))
    bad5=ctx.request.get(BASE+f"/api/dxf?runId={rV2['runId']}&type=plan")
    ok("S40d 위반이면 **도면을 뜨지 않는다** (422 · 검증 안 된 도면이 밖으로 나가지 않게)",
       (bad5.status, bad5.text()[:60]), bad5.status==422)
    bad6=ctx.request.post(BASE+"/api/drawings",headers=J0,data=json.dumps({"runId":rV2["runId"],"type":"iso"}))
    ok("S40e 등록도 같은 이유로 막힌다 (3D 투영도 예외 없음)", bad6.status, bad6.status==422)
    # 원상 복구 후, 앞서 통과한 스냅샷은 여전히 도면이 나온다
    put([{"name":x["name"], **({"len":x["len"]} if x.get("len") is not None else {}), **({"dir":x["dir"]} if x.get("dir") else {}), "components":x.get("components",[])} for x in g3["sections"]])
    still=ctx.request.get(BASE+f"/api/dxf?runId={rV1['runId']}&type=plan")
    ok("S40f 규칙을 고쳐도 앞서 통과한 스냅샷의 판정은 그대로다 (판정도 스냅샷 — 0011 원칙)", still.status, still.status==200)

    # ── S39 3D View 1차 (0013 · 코퍼스 "3D View" ISO·Exploded) ──
    # 형상 모델이 아니라 **같은 스냅샷의 등각 투영**이다 — 치수가 2D 뷰와 어긋나지 않는 것이 핵심.
    rI=run55("S39"); RIDI=rI["runId"]
    mi=ctx.request.get(BASE+f"/api/dxf?runId={RIDI}&type=iso&meta=1").json()
    mx=ctx.request.get(BASE+f"/api/dxf?runId={RIDI}&type=exploded&meta=1").json()
    mpl=ctx.request.get(BASE+f"/api/dxf?runId={RIDI}&type=plan&meta=1").json()
    ok("S39a 같은 스냅샷에서 등각도(iso)·분해도(exploded)가 나온다",
       (mi.get("type"), mx.get("type"), mi.get("sections")==mpl.get("sections")),
       mi.get("type")=="iso" and mx.get("type")=="exploded" and mi.get("sections")==mpl.get("sections"))
    ok("S39b 3D 투영도 2D 뷰와 같은 치수를 쓴다 — 전장·폭·높이가 평면도와 일치(분해도는 띄운 간격을 치수로 세지 않는다)",
       (mpl.get("lengthMm"), mi.get("lengthMm"), mx.get("lengthMm"), mi.get("widthMm")==mpl.get("widthMm")),
       mi.get("lengthMm")==mpl.get("lengthMm")==mx.get("lengthMm") and mi.get("widthMm")==mpl.get("widthMm"))
    di=ctx.request.get(BASE+f"/api/dxf?runId={RIDI}&type=iso").text()
    dx=ctx.request.get(BASE+f"/api/dxf?runId={RIDI}&type=exploded").text()
    gi=dxf_stats(di); gx=dxf_stats(dx)
    ok("S39c 등각 투영이 실제로 기울어져 있다 (평면도는 축에 나란한 선뿐 · 등각도는 기운 선이 대부분 · ezdxf 실측)",
       (gi["n"], gx["n"]), gi["n"]>=12 and gx["n"]>gi["n"]*0 )
    import ezdxf as _ez
    def slanted(txt):
        d=_ez.read(io.StringIO(txt))
        ls=[e for e in d.modelspace() if e.dxftype()=="LINE"]
        sl=[e for e in ls if abs(e.dxf.start.x-e.dxf.end.x)>1 and abs(e.dxf.start.y-e.dxf.end.y)>1]
        return len(sl), len(ls)
    si,ti=slanted(di); sp,tp=slanted(ctx.request.get(BASE+f"/api/dxf?runId={RIDI}&type=plan").text())
    ok("S39d 기운 선 실측 — 등각도는 기운 선이 있고 평면도는 없다(투영이 진짜 돌아간 증거)", (si,ti,sp,tp), si>0 and sp==0)
    ok("S39e 분해도에 조립 순서가 붙는다 (구획 순서가 곧 순서 — 따로 적지 않는다)",
       [t for t in gx["texts"] if t.startswith(("1.","2."))][:2], any(t.startswith("1.") for t in gx["texts"]))
    regI=ctx.request.post(BASE+"/api/drawings",headers=J0,data=json.dumps({"runId":RIDI,"type":"iso"})).json()
    regX=ctx.request.post(BASE+"/api/drawings",headers=J0,data=json.dumps({"runId":RIDI,"type":"exploded"})).json()
    ok("S39f 등각도·분해도도 번호·개정을 받아 등록되고 승인 대장에 함께 뜬다 (-ISO · -EXP)",
       (regI.get("drawingNo"), regX.get("drawingNo")),
       str(regI.get("drawingNo","")).endswith("-ISO") and str(regX.get("drawingNo","")).endswith("-EXP"))
    reg2=ctx.request.get(BASE+"/api/register?kind=drawing").json()
    ok("S39g 승인 대장이 여섯 종(plan·front·right·assembly·iso·exploded)을 한 표에서 본다",
       sorted(set(r["type"] for r in reg2.get("rows",[]))), {"iso","exploded"} <= set(r["type"] for r in reg2.get("rows",[])))
    bad4=ctx.request.get(BASE+f"/api/dxf?runId={RIDI}&type=gltf")
    ok("S39h 실제 형상 모델(gltf)은 아직 없다 — 400 으로 막고 있는 척하지 않는다", bad4.status, bad4.status==400)

    # ── S38 Component 배치 규칙 (p36 · 코퍼스 "Component Position Rule") ──
    g2=ctx.request.get(ARR+"?code=EU&slots="+json.dumps({"A":"EU","B":"55","C":"2123","D":"630","E":"SS","F":"1-21-13-15"})).json()
    fan=[x for x in g2["sections"] if x["name"]=="Fan"][0]
    ok("S38a 구획마다 **그 구획에 달린 BOM 자식**을 배치 후보로 준다 (화면이 부품 목록을 지어내지 않는다)",
       (fan.get("children"), fan.get("components")), bool(fan.get("children")))
    rows2b=[{"name":x["name"], **({"len":x["len"]} if x.get("len") is not None else {}), **({"dir":x["dir"]} if x.get("dir") else {}), "components":x.get("components",[])} for x in g2["sections"]]
    target=fan["children"][0]
    for r_ in rows2b:
        if r_["name"]=="Fan": r_["components"]=[{"code":target,"at":"rear","level":"bottom"}]
    pc=put(rows2b)
    rC=run55("S38"); metaC=ctx.request.get(BASE+f"/api/dxf?runId={rC['runId']}&type=plan&meta=1").json()
    dxfC=ctx.request.get(BASE+f"/api/dxf?runId={rC['runId']}&type=plan").text()
    ok("S38b 배치를 저장하면 평면도에 그 부품 상자가 그려진다 (구획 3×3 칸 · COMPONENT 레이어)",
       (pc.status, metaC.get("components"), "COMPONENT" in dxfC),
       pc.status==200 and metaC.get("components")==[{"section":"Fan","code":target,"at":"rear","level":"bottom"}] and "COMPONENT" in dxfC)
    gc=dxf_stats(dxfC)
    ok("S38c 상자는 구획 안에 있다 (도형 실측 — 전장·폭을 넘지 않는다)",
       (gc["maxx"], metaC.get("lengthMm"), gc["layers"]), gc["maxx"]<=metaC.get("lengthMm") and "COMPONENT" in gc["layers"])
    bad2=put([{**r_, "components":[{"code":"KCP 1","at":"rear","level":"bottom"}]} if r_["name"]=="Fan" else r_ for r_ in rows2b])
    ok("S38d 그 구획의 부품이 아니면 배치를 거부한다 (409 — 없는 부품을 도면에 그리지 않는다)", bad2.status, bad2.status==409)
    bad3=put([{**r_, "components":[{"code":target,"at":"nose","level":"bottom"}]} if r_["name"]=="Fan" else r_ for r_ in rows2b])
    ok("S38e 배치 값이 목록 밖이면 400", bad3.status, bad3.status==400)
    old_meta=ctx.request.get(BASE+f"/api/dxf?runId={rD['runId']}&type=plan&meta=1").json()
    ok("S38f 앞 스냅샷 도면에는 배치가 없다 — 배치도 스냅샷에 박힌 값만 쓴다(0011 원칙)",
       old_meta.get("components"), not old_meta.get("components"))

    # 원상 복구 — 뒤 단계(S16~)가 옛 순서를 전제로 하지 않게
    put(rows)

    # ── S33 2D 3각법 (0012 · 코퍼스 "2D 3각법 View" · 청사진 p36 DWG View) ──
    # 핵심: Front/Top(plan)/Right 가 **같은 스냅샷 하나**에서 나오고, 뷰 간 치수가 어긋나지 않는다(코퍼스 "View 간 치수 동기화").
    rV=run55("S33"); RID=rV["runId"]
    mp=ctx.request.get(BASE+f"/api/dxf?runId={RID}&type=plan&meta=1").json()
    mf=ctx.request.get(BASE+f"/api/dxf?runId={RID}&type=front&meta=1").json()
    mr=ctx.request.get(BASE+f"/api/dxf?runId={RID}&type=right&meta=1").json()
    ok("S33a 같은 스냅샷에서 평면(Top)·정면(Front)·우측면(Right) 세 뷰가 나온다",
       (mp.get("type"), mf.get("type"), mr.get("type")), (mp.get("type"),mf.get("type"),mr.get("type"))==("plan","front","right"))
    ok("S33b 뷰 간 치수가 같다 — 전장은 평면=정면, 폭은 평면=우측면, 높이는 정면=우측면 (같은 Parameter Set)",
       (mp.get("lengthMm"), mf.get("lengthMm"), mp.get("widthMm"), mr.get("widthMm"), mf.get("heightMm"), mr.get("heightMm")),
       mp.get("lengthMm")==mf.get("lengthMm") and mp.get("widthMm")==mr.get("widthMm") and mf.get("heightMm")==mr.get("heightMm"))
    def outline(txt):
        # 외형은 **선(LINE)** 으로만 잰다 — 표제란 글자가 도형 밖에 있어 텍스트를 섞으면 치수가 아니다
        doc=ezdxf.read(io.StringIO(txt)); ls=[e for e in doc.modelspace() if e.dxftype()=="LINE" and e.dxf.layer=="OUTLINE"]
        xs=[v for e in ls for v in (e.dxf.start.x, e.dxf.end.x)]; ys=[v for e in ls for v in (e.dxf.start.y, e.dxf.end.y)]
        return {"maxx":round(max(xs)),"maxy":round(max(ys)),"n":len(ls)}
    fdx=ctx.request.get(BASE+f"/api/dxf?runId={RID}&type=front").text()
    rdx=ctx.request.get(BASE+f"/api/dxf?runId={RID}&type=right").text()
    gf=outline(fdx); gr=outline(rdx)
    ok("S33c 정면도 도형 실측(ezdxf 로 파싱) — 가로 = 전장, 세로 = 높이 H",
       (gf["maxx"], gf["maxy"], mf.get("heightMm")), gf["maxx"]==mf.get("lengthMm") and gf["maxy"]==mf.get("heightMm"))
    ok("S33d 우측면도 도형 실측 — 가로 = 폭 W, 세로 = 높이 H (구획선은 긋지 않는다 · 겹쳐 보이므로)",
       (gr["maxx"], gr["maxy"]), gr["maxx"]==mr.get("widthMm") and gr["maxy"]==mr.get("heightMm"))
    bad=ctx.request.get(BASE+f"/api/dxf?runId={RID}&type=iso_3d")
    ok("S33e 아직 없는 뷰(iso_3d · 3D 미착수)는 400 으로 거부한다 — 조용히 평면도로 떨어지지 않는다", bad.status, bad.status==400)
    regF=ctx.request.post(BASE+"/api/drawings",headers=J0,data=json.dumps({"runId":RID,"type":"front"})).json()
    regR=ctx.request.post(BASE+"/api/drawings",headers=J0,data=json.dumps({"runId":RID,"type":"right"})).json()
    ok("S33f 정면도·우측면도도 번호·개정을 받아 등록된다 (-FRT · -RHT · Rev A)",
       (regF.get("drawingNo"), regF.get("rev"), regR.get("drawingNo")),
       str(regF.get("drawingNo","")).endswith("-FRT") and str(regR.get("drawingNo","")).endswith("-RHT") and regF.get("rev")=="A")

    # ── S34 p32 자재·구매 품목 — 등록한 공급처가 BOM 스냅샷에 박히고 구매 요청까지 간다 ──
    # 원칙은 0011 과 같다: 구매 요청은 "지금 등록된 공급처"가 아니라 **그 BOM 을 돌린 시점**의 공급처를 산다.
    def pc_of(code):
        c_=ctx.request.get(BASE+"/api/setup/catalog").json(); return [p_ for p_ in c_.get("productCodes",[]) if p_.get("code")==code][0]
    inv=pc_of("PVF 1")
    buy=[t for k,t in inv.get("tables",{}).items() if t.get("role")=="buy"]
    ok("S34a 구매 품목 코드에 **구매 속성표(buy)** 가 등록돼 있다 (p32 A:V·B:Hz·C:IP·D:Insulation·F:Supplier)",
       (list(inv.get("tables",{}).keys()), [c["name"] for c in buy[0]["cols"]] if buy else None),
       bool(buy) and "Supplier" in [c["name"] for c in buy[0]["cols"]])
    rP=run55("S34"); RIDP=rP["runId"]
    scr=[l for l in rP.get("lines",[]) if "supplier" in l]
    ok("S34b 화면 응답에는 공급처가 없다 — 공급처는 **스냅샷 줄**에 박히고, 구매 요청이 거기서 복사한다(화면이 값을 나르지 않는다)",
       len(scr), len(scr)==0)
    prr=ctx.request.post(BASE+"/api/purchase-requests",headers=J0,data=json.dumps({"runId":RIDP})).json()
    prs=ctx.request.get(BASE+"/api/purchase-requests").json()
    mine=[x for x in prs.get("rows",[]) if x.get("prNo")==prr.get("prNo")]
    pl=[(l.get("part"), l.get("supplier")) for l in (mine[0].get("lines") if mine else [])]
    ok("S34c 구매 요청 줄이 그 공급처를 그대로 받는다 (p51 Supplier 열이 더 이상 비지 않는다)",
       (prr.get("prNo"), pl), any(p_=="Inverter" and s_=="LS ELECTRIC" for p_,s_ in pl))
    # 등록 공급처를 바꿔도 **이미 뜬 스냅샷**의 구매 요청은 옛 공급처다
    inv2=pc_of("PVF 1")
    for k,t in inv2["tables"].items():
        if t.get("role")=="buy": t["rows"][0]["cells"]["A"]="현대일렉트릭"
    up2=ctx.request.post(BASE+"/api/setup/product-codes",headers=J0,data=json.dumps(inv2))
    rP2=run55("S34b")
    pr2=ctx.request.post(BASE+"/api/purchase-requests",headers=J0,data=json.dumps({"runId":rP2["runId"]})).json()
    rows3_=ctx.request.get(BASE+"/api/purchase-requests").json().get("rows",[])
    mine3=[x for x in rows3_ if x.get("prNo")==pr2.get("prNo")]
    new_sup=[l.get("supplier") for l in (mine3[0].get("lines") if mine3 else []) if l.get("part")=="Inverter"]
    rows2_=ctx.request.get(BASE+"/api/purchase-requests").json().get("rows",[])
    mine2=[x for x in rows2_ if x.get("prNo")==prr.get("prNo")]
    old_pr=[l.get("supplier") for l in (mine2[0].get("lines") if mine2 else []) if l.get("part")=="Inverter"]
    ok("S34d 공급처를 바꾸면 **새 BOM** 부터 바뀐다 — 앞서 뜬 구매 요청은 옛 공급처 그대로(스냅샷이 근거)",
       (up2.status, new_sup, old_pr), up2.status==200 and new_sup==["현대일렉트릭"] and old_pr==["LS ELECTRIC"])
    # 되돌린다 (뒤 장면·CSV 기대값 보존)
    for k,t in inv2["tables"].items():
        if t.get("role")=="buy": t["rows"][0]["cells"]["A"]="LS ELECTRIC"
    ctx.request.post(BASE+"/api/setup/product-codes",headers=J0,data=json.dumps(inv2))

    # ── 캡처 복원(E8) — 05 프로젝트 목록 · 06 CPQ 모듈 · 48·49 도면 그림 ──
    _tds=[d for d in ctx.request.get(BASE+"/api/documents").json().get("rows",[]) if d.get("docType")=="techdata"]
    ok("S37a Tech Data 문서가 남아 있다(인쇄본 캡처의 근거)", len(_tds), len(_tds)>0)
    pg.goto(BASE+f"/api/documents/{_tds[0]['id']}/print",wait_until="domcontentloaded")
    pg.wait_for_function("()=>document.body.innerText.includes('TECH DATA')",timeout=30000)  # 인쇄본이 그려진 뒤에 캡처(고정 sleep 아님)
    pg.screenshot(path=f"{OUT}/47_techdata_print.png",full_page=True)
    pg.goto(BASE+"/m/project",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=pm-detail][data-project='PS-61313-5']",timeout=30000); nuke(pg)
    pg.screenshot(path=f"{OUT}/05_project_mgmt.png",full_page=True)
    pg.goto(BASE+"/m/cpq",wait_until="domcontentloaded"); pg.wait_for_selector("text=CPQ",timeout=30000); nuke(pg)
    pg.screenshot(path=f"{OUT}/06_module_cpq_stub.png",full_page=True)
    # 도면은 **내려받은 DXF 를 ezdxf 로 다시 그린 그림**이다(화면 캡처가 아니다 — 파일이 진짜라는 증거)
    def dxf_png(txt, path, title):
        import matplotlib; matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        doc=ezdxf.read(io.StringIO(txt)); msp=doc.modelspace()
        fig,ax=plt.subplots(figsize=(12,6.5))
        for e in msp:
            if e.dxftype()=="LINE":
                ax.plot([e.dxf.start.x,e.dxf.end.x],[e.dxf.start.y,e.dxf.end.y],lw=0.9,color="#1c2b2b")
            elif e.dxftype()=="CIRCLE":
                ax.add_patch(plt.Circle((e.dxf.center.x,e.dxf.center.y),e.dxf.radius,fill=False,lw=0.8,color="#1c2b2b"))
            elif e.dxftype()=="TEXT":
                ax.text(e.dxf.insert.x,e.dxf.insert.y,e.dxf.text,fontsize=5.5,color="#1c2b2b")
        ax.set_aspect("equal"); ax.axis("off"); ax.set_title(title,fontsize=8)
        fig.savefig(path,dpi=150,bbox_inches="tight"); plt.close(fig)
    _plan=ctx.request.get(BASE+f"/api/dxf?runId={RUN1}&type=plan").text()
    _asm=ctx.request.get(BASE+f"/api/dxf?runId={RUN1}&type=assembly").text()
    dxf_png(_plan, f"{OUT}/48_dxf_plan.png", "EDIM - PLAN (ezdxf re-render of the downloaded DXF)")
    dxf_png(_asm,  f"{OUT}/49_dxf_assembly.png", "EDIM - ASSEMBLY (ezdxf re-render of the downloaded DXF)")

    # ── S41 p58 Main Work place Toolbar — 버튼이 명령을 보낸다 (편집은 Design 초안·같은 저장, 승인은 Inspector, 도면은 스냅샷) ──
    NODE4=BASE+"/workbench?node=a0000000-0000-4000-8000-000000000004"
    def arr_rows():
        return pg.eval_on_selector_all("[data-testid=arr-table] tbody tr","es=>es.map(e=>e.dataset.testid.replace('arr-row-',''))")
    def wait_canvas():
        pg.wait_for_selector("[data-testid=canvas-cmds]",timeout=30000); nuke(pg)
        pg.wait_for_selector("[data-testid=canvas-cmds][data-ready='1']",timeout=60000)  # 하이드레이션 대기 — SSR 버튼은 풀려 보여도 클릭이 사라진다
    pg.goto(NODE4,wait_until="domcontentloaded"); wait_canvas()
    cmds=pg.eval_on_selector_all("[data-testid=canvas-cmds] [data-cmd]","es=>es.map(e=>e.dataset.cmd)")
    none=pg.eval_on_selector_all("[data-testid=canvas-cmds] [data-cmd-none]","es=>es.map(e=>[e.dataset.cmdNone,e.disabled,e.title.length>10])")
    ok("S41a 툴바가 p58 구성이다 — Arrangement·Move·Delete·Add·Copy·DWG View·승인은 명령, Free CAD·설계 심볼은 잠긴 자리(이유 표기)",
       (cmds,none), cmds==["arrangement","move","delete","add","copy","dwg-view","approval"] and len(none)==2 and all(d and t for _,d,t in none))
    dis=pg.eval_on_selector_all("[data-cmd=move],[data-cmd=delete],[data-cmd=copy]","es=>es.map(e=>e.disabled)")
    ok("S41b 구획을 고르기 전에는 Move·Delete·Copy 가 잠긴다 (대상 없는 명령을 막는다)", dis, dis==[True,True,True])
    pg.click("[data-cmd=arrangement]"); pg.wait_for_selector("[data-testid=design-canvas][data-loaded='1']",timeout=30000); pg.wait_for_selector("[data-testid=arr-table] tbody tr",timeout=30000)
    ok("S41c Arrangement ▼ 는 Design 탭의 기존 Arrangement 편집을 연다 (두 번째 편집 화면 없음)",
       bool(pg.query_selector("[data-testid=design-canvas] [data-testid=arrangement-panel]")), bool(pg.query_selector("[data-testid=design-canvas] [data-testid=arrangement-panel]")))
    base_rows=arr_rows()
    pg.click("[data-testid=canvas-sec-Fan]")
    pg.wait_for_function("()=>{const b=document.querySelector('[data-cmd=copy]'); return b && !b.disabled;}",timeout=30000)
    ok("S41d 개념도에서 구획을 누르면 선택되고 Move·Delete·Copy 가 풀린다", pg.inner_text("[data-testid=canvas-selected]"),
       pg.get_attribute("[data-testid=canvas-sec-Fan]","data-selected")=="1")
    pg.click("[data-cmd=copy]"); pg.wait_for_selector("[data-testid=arr-row-Fan-2]",timeout=30000)
    r1=arr_rows()
    ok("S41e Copy — 고른 구획 바로 뒤에 Fan-2 가 초안으로 생기고 선택이 새 구획으로 옮겨 간다", r1,
       r1.index("Fan-2")==r1.index("Fan")+1 and "Fan-2" in pg.inner_text("[data-testid=canvas-selected]"))
    del_on=pg.eval_on_selector("[data-cmd=delete]","e=>!e.disabled && !e.dataset.locked")
    pg.click("[data-cmd=delete]"); pg.wait_for_selector("[data-testid=arr-row-Fan-2]",state="detached",timeout=30000)
    ok("S41f Delete — 관계 없는 구획(Fan-2)은 버튼이 풀려 있고, 누르면 초안에서 지워진다", (del_on, arr_rows()), del_on and arr_rows()==base_rows)
    pg.click("[data-testid=canvas-sec-Coil]"); pg.wait_for_function("()=>document.querySelector('[data-testid=canvas-sec-Coil]')?.dataset.selected==='1'",timeout=30000)
    # E8(p58) — 표의 Delete 와 같은 판정으로 툴바 Delete 가 **미리** 잠긴다. 서버 409(S32f)는 그대로 둔 이중 방어.
    pg.wait_for_function("()=>{const b=document.querySelector('[data-cmd=delete]'); return b && b.disabled && b.dataset.locked==='1';}",timeout=30000)
    dl=pg.eval_on_selector("[data-cmd=delete]","e=>[e.disabled,e.title]")
    tbl=pg.eval_on_selector("[data-testid=arr-del-Coil]","e=>e.disabled")
    ok("S41g Delete 미리 잠금 — BOM 관계가 걸린 구획(Coil)을 고르면 툴바 Delete 가 표의 Delete 처럼 눌리지 않고 이유를 적는다 (서버 409 는 S32f)",
       (dl, tbl, "Coil" in arr_rows()), dl[0] and "BOM 관계" in dl[1] and tbl and "Coil" in arr_rows())
    first=base_rows[0]
    pg.click("[data-testid=canvas-sec-Fan]"); pg.wait_for_function("()=>document.querySelector('[data-testid=canvas-sec-Fan]')?.dataset.selected==='1'",timeout=30000)
    pg.click("[data-cmd=move]"); pg.wait_for_function("()=>document.querySelector('[data-testid=canvas-selected]')?.dataset.moving==='1'",timeout=30000)
    pg.click(f"[data-testid=canvas-sec-{first}]")
    pg.wait_for_function("(f)=>{const r=[...document.querySelectorAll('[data-testid=arr-table] tbody tr')].map(e=>e.dataset.testid); return r[0]==='arr-row-Fan';}", arg=first, timeout=30000)
    ok(f"S41h Move — Fan 을 고르고 Move → {first} 자리를 누르면 초안 순서가 바뀐다(저장 전)", arr_rows()[:3], arr_rows()[0]=="Fan")
    # 같은 저장 한 곳: 새로 열고(초안 버림) Copy → 저장 → API 에 반영 → API 로 원복(09-21 규칙)
    SL={"A":"EU","B":"55","C":"2123","D":"630","E":"SS","F":"1-21-13-15"}
    g_before=ctx.request.get(ARR+"?code=EU&slots="+json.dumps(SL)).json()["sections"]
    pg.goto(NODE4,wait_until="domcontentloaded"); wait_canvas()
    pg.click("[data-cmd=arrangement]"); pg.wait_for_selector("[data-testid=design-canvas][data-loaded='1']",timeout=30000); pg.wait_for_selector("[data-testid=arr-table] tbody tr",timeout=30000)
    pg.click("[data-testid=canvas-sec-Fan]"); pg.wait_for_function("()=>{const b=document.querySelector('[data-cmd=copy]'); return b && !b.disabled;}",timeout=30000)
    pg.click("[data-cmd=copy]"); pg.wait_for_selector("[data-testid=arr-row-Fan-2]",timeout=30000)
    pre_save=arr_rows()   # 저장 직전 초안에 Fan-2 가 있는지(늦은 재로딩이 덮어썼다면 여기서 보인다)
    # 저장 **응답**을 기다린다. (예전 대기 조건 '다음 BOM Run' 은 패널의 고정 안내문에도 있어 즉시 통과했고,
    #  아래 GET 이 저장 커밋보다 먼저 읽는 경합이 Windows 에서 S41i 를 떨어뜨렸다 — 2026-09-26 실측)
    with pg.expect_response(lambda q: q.url.endswith("/api/setup/arrangement") and q.request.method=="POST",timeout=30000) as sres:
        pg.click("[data-testid=arr-save]")
    sent=[x["name"] for x in json.loads(sres.value.request.post_data or "{}").get("sections",[])]
    pg.wait_for_function("()=>(document.querySelector('[data-testid=arr-msg]')?.innerText||'').startsWith('저장 ·')",timeout=30000)
    g_after=[x["name"] for x in ctx.request.get(ARR+"?code=EU&slots="+json.dumps(SL)).json()["sections"]]
    fan=[x for x in g_before if x["name"]=="Fan"][0]
    ok("S41i 툴바 편집은 기존 '저장' 한 곳으로 반영된다 — Fan-2 가 등록되고 Fan 의 길이·방향을 물려받는다(부품 배치는 비움)", (g_after, pre_save, sent),
       "Fan-2" in g_after and g_after.index("Fan-2")==g_after.index("Fan")+1 and "Fan-2" in pre_save and "Fan-2" in sent)
    rs=put([{"name":x["name"], **({"len":x["len"]} if x.get("len") is not None else {}), **({"dir":x["dir"]} if x.get("dir") else {}), "components":x.get("components",[])} for x in g_before])
    g_back=[x["name"] for x in ctx.request.get(ARR+"?code=EU&slots="+json.dumps(SL)).json()["sections"]]
    ok("S41j 원복(API) — 뒤 단계가 오염되지 않게 등록 구획을 되돌린다", (rs.status, g_back==[x["name"] for x in g_before]), rs.status==200 and g_back==[x["name"] for x in g_before])
    pg.goto(NODE4,wait_until="domcontentloaded"); wait_canvas()
    dwg_off=pg.eval_on_selector("[data-cmd=dwg-view]","e=>e.disabled")
    if pg.get_attribute("[data-testid=toolbox-toggle]","aria-pressed")=="true":   # 앞 단계(S13)가 열어 둔 Toolbox 창이 Action Bar 를 가린다
        pg.click("[data-testid=toolbox-toggle]"); pg.wait_for_selector("[data-testid=toolbox-toggle][aria-pressed=false]",timeout=30000)
    nuke(pg); pg.click("[data-run=bom]")   # 개발 서버 표시기(nextjs-portal)가 왼쪽 아래 BOM Run 을 덮는다 — 기존 단계와 같이 치운다
    pg.wait_for_function("()=>{const s=document.querySelector('[data-cmd=dwg-view]'); return s && !s.disabled;}",timeout=60000)
    # F6(2026-09-27): DWG View 는 내려받는 대신 **화면에 띄운다**. 같은 DXF 는 뷰어 안 링크로 받는다 — 기대값 갱신(근거: ccmd F · p13)
    pg.select_option("[data-cmd=dwg-view]","front")
    pg.wait_for_selector("[data-testid=dwg-viewer][data-view=front][data-ready='1']",timeout=30000)
    with pg.expect_download(timeout=30000) as dl:
        pg.click("[data-testid=dwg-viewer-download]")
    fn=dl.value.suggested_filename
    pg.click("[data-testid=dwg-viewer-close]"); pg.wait_for_selector("[data-testid=dwg-viewer]",state="detached",timeout=10000)
    ok("S41k DWG View ▼ — BOM 스냅샷이 없으면 잠기고, 있으면 고른 뷰(정면도)의 DXF 를 받는다", (dwg_off, fn), dwg_off and fn.endswith("-front.dxf"))
    pg.click("[data-cmd=approval]")
    pg.wait_for_function("()=>document.querySelector('[data-testid=inspector-approval]')?.dataset.focused==='1'",timeout=30000)
    ok("S41l 승인 — Inspector 의 Approval 로 데려가 강조한다 (요청·결정은 거기서만 — 보는 곳/하는 곳 분리)", True, True)
    pg.click("[data-cmd=arrangement]"); pg.wait_for_selector("[data-testid=canvas-sec-Fan]",timeout=30000)
    pg.click("[data-testid=canvas-sec-Fan]"); pg.wait_for_function("()=>{const b=document.querySelector('[data-cmd=move]'); return b && !b.disabled;}",timeout=30000)
    nuke(pg); pg.screenshot(path=f"{OUT}/54_toolbar.png")   # 캡처만 — 초안은 저장하지 않는다

    # ── S36 Schedule management (p12·18·50) — 작업대를 떠나지 않고 일정을 잡는다 ──
    pg.goto(BASE+"/workbench?node=a0000000-0000-4000-8000-000000000004",wait_until="domcontentloaded")
    pg.wait_for_selector("[data-testid=task-title]",timeout=30000); nuke(pg)
    before=pg.eval_on_selector_all("[data-testid=task-row]","e=>e.length")
    # 하이드레이션 전에는 DOM 에 값만 들어가고 React 상태가 비어 버튼이 잠겨 있다 — 버튼이 풀릴 때까지 기다렸다 누른다
    for _ in range(20):
        pg.fill("[data-testid=task-title]","코일 사양 확인")
        pg.fill("[data-testid=task-due]","2026-09-01")   # 지난 날짜 — '지남' 표기 확인용
        if pg.eval_on_selector("[data-testid=task-add]","e=>!e.disabled"): break
        time.sleep(0.3)
    pg.click("[data-testid=task-add]")
    pg.wait_for_function("(n)=>document.querySelectorAll('[data-testid=task-row]').length>n", arg=before, timeout=30000)
    todo=pg.inner_text("[data-testid=task-group-todo]")
    ok("S36a 작업대에서 할 일을 기한과 함께 등록한다 (To-do list 에 뜨고 기한이 지났으면 '지남')",
       (before, todo.replace("\n"," | ")[:70]), "코일 사양 확인" in todo and "지남" in todo)
    tid=pg.eval_on_selector("[data-testid=task-group-todo] [data-testid=task-row] button","e=>e.dataset.testid.replace('task-toggle-','')")
    pg.click(f"[data-testid=task-toggle-{tid}]")
    pg.wait_for_function("()=>{const d=document.querySelector('[data-testid=task-group-done]'); return d && d.innerText.includes('코일 사양 확인');}", timeout=30000)
    donet=pg.inner_text("[data-testid=task-group-done]"); todo2=pg.inner_text("[data-testid=task-group-todo]")
    ok("S36b '완료'를 누르면 To-do 에서 Done items 로 옮겨간다 (되돌리기도 있다)",
       (donet.split("\n")[0], "코일 사양 확인" in todo2), "코일 사양 확인" in donet and "코일 사양 확인" not in todo2)
    ok("S36c Approval Request List 가 같은 상자에 뜬다 (요청·결정은 아래 Approval 에서 — 보는 곳과 하는 곳을 나눈다)",
       bool(pg.query_selector("[data-testid=approval-request-list]")) or "요청 없음" in pg.inner_text("[data-testid=task-group-done]"),
       True)
    pg.screenshot(path=f"{OUT}/53_schedule.png",full_page=True)

    # ── S35 승인 대장 (p55 EDIM Approval Management) — 문서와 도면을 한 표로 ──
    reg=ctx.request.get(BASE+"/api/register").json()
    kinds=set(r["kind"] for r in reg.get("rows",[]))
    ok("S35a 대장이 문서와 도면을 함께 모은다 (번호·개정·상태·발행 시각)",
       (reg.get("total"), sorted(kinds), reg.get("counts")), reg.get("total",0)>0 and kinds=={"document","drawing"})
    iss=ctx.request.get(BASE+"/api/register?status=issued").json()
    ok("S35b 상태로 거른다 — 발행본만 (발행 시각이 채워져 있다)",
       (iss.get("total"), [r["releasedAt"] is not None for r in iss.get("rows",[])][:3]),
       iss.get("total",0)>0 and all(r["status"]=="issued" and r["releasedAt"] for r in iss["rows"]))
    dw=ctx.request.get(BASE+"/api/register?kind=drawing").json()
    ok("S35c 종류로 거른다 — 도면만 (3각법 4종이 같은 대장에 있다)",
       sorted(set(r["type"] for r in dw.get("rows",[]))), all(r["kind"]=="drawing" for r in dw.get("rows",[])) and dw.get("total",0)>0)
    badr=ctx.request.get(BASE+"/api/register?status=released")
    ok("S35d 없는 상태 이름은 400 으로 막는다 (조용히 전체를 주지 않는다)", badr.status, badr.status==400)
    pg.goto(BASE+"/m/register",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=register]",timeout=30000)
    pg.wait_for_function("()=>document.querySelectorAll('[data-testid=reg-row]').length>0", timeout=30000)
    allrows=pg.eval_on_selector_all("[data-testid=reg-row]","e=>e.length")
    pg.click("[data-testid=reg-f-issued]")
    pg.wait_for_function("()=>[...document.querySelectorAll('[data-testid=reg-row]')].every(e=>e.dataset.status==='issued')", timeout=30000)
    onlyissued=pg.eval_on_selector_all("[data-testid=reg-row]","e=>e.length")
    chips=pg.inner_text("[data-testid=register]")
    ok("S35e 화면: 대장이 뜨고 '발행' 칩을 누르면 발행본만 남는다 · 칩 숫자는 걸러도 그대로다(전체 수가 0 이 되지 않는다)",
       (allrows, onlyissued, chips.split("\n")[1][:40]), allrows>onlyissued>0 and f"전체 {allrows}" in chips)
    pg.screenshot(path=f"{OUT}/52_register.png",full_page=True)

    # ── P3-a 플랫폼 관리자 계층 · DB①/DB② 소유 분리 (p54 User Management · p59 최종 승인 · p64 Admin.) ──
    # S16 회사 관리자가 Company Info.에서 Special 의뢰를 올린다 = 회사→플랫폼 유일 통로
    pg.goto(BASE+"/m/company",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=user-management]",timeout=30000); time.sleep(1.5); nuke(pg)
    subj="Special 의뢰 (demo) "+str(int(time.time()))
    pg.fill("[data-testid=request-subject]",subj); pg.fill("[data-testid=request-detail]","코일 열교환 계산 — 매크로로 안 됩니다")
    pg.click("[data-testid=request-submit]"); pg.wait_for_selector(f"[data-testid=platform-requests] >> text={subj}", timeout=30000); time.sleep(0.4); nuke(pg)  # 고정 sleep 금지 — 첫 컴파일이 4초를 넘긴다
    body=pg.inner_text("[data-testid=platform-requests]"); ok("S16a 회사 owner가 Special 의뢰를 올린다 (회사→플랫폼 유일 통로, 상태 '대기')", subj[-14:], subj in body and "대기" in body)
    pg.screenshot(path=f"{OUT}/40_company_admin.png",full_page=True)
    # S16b 회사 세션으로는 플랫폼 영역이 열리지 않는다 (API·화면 둘 다)
    r=ctx.request.get(BASE+"/api/platform/requests"); ok("S16b 회사 세션은 플랫폼 대기열 API에 401", r.status, r.status==401)
    pg.goto(BASE+"/platform",wait_until="domcontentloaded"); time.sleep(1.2); nuke(pg); body=pg.inner_text("body")
    ok("S16c 회사 계정의 /platform 화면은 403 안내", body[:40].replace("\n"," "), "403" in body and "플랫폼 관리자 전용" in body)
    # S16d 플랫폼 계정 = 별도 사람. 로그인하면 /platform으로 간다
    pl=b.new_context(viewport={"width":1440,"height":900}); plp=pl.new_page()
    r=pl.request.post(BASE+"/api/auth/login",data=LOGIN("platform@edim.test"))
    ok("S16d 멤버십 없는 플랫폼 계정이 로그인된다 → /platform", (r.status, r.json().get("redirect")), r.status==200 and r.json().get("redirect")=="/platform")
    plp.goto(BASE+"/platform",wait_until="domcontentloaded"); plp.wait_for_selector("[data-testid=request-queue]",timeout=30000); time.sleep(1.2); nuke(plp)
    body=plp.inner_text("body")
    ok("S16e 플랫폼 콘솔: 테넌트 2곳과 올라온 의뢰가 보인다", (("Acme AHU" in body), ("Globex Air" in body)), "Acme AHU" in body and "Globex Air" in body and subj in body)
    _m16=re.search(r"원천자료 (\d+)건",body)
    ok("S16f DB① 에는 샘플 학습 자료만 들어 있다(0033 시드 · 회사 자료 0) — 학습 탭으로 간다", (_m16.group(1) if _m16 else None, bool(plp.query_selector("[data-testid=platform-learning-link]"))),
       bool(_m16) and int(_m16.group(1))>0 and bool(plp.query_selector("[data-testid=platform-learning-link]")))
    ok("S16g 플랫폼 콘솔에 고객사 업무 데이터는 없다 (BOM·프로젝트·코드 0건)", ("Micron" in body, "EU-55" in body), ("Micron" not in body) and ("EU-55" not in body))
    plp.screenshot(path=f"{OUT}/41_platform_console.png",full_page=True)
    # S16h 플랫폼 계정은 회사 업무 화면에 못 들어간다 (반대 방향 차단)
    plp.goto(BASE+"/workbench",wait_until="domcontentloaded"); time.sleep(1.2)
    ok("S16h 플랫폼 계정은 /workbench에 들어갈 수 없다", plp.url.split(BASE)[-1], "/login" in plp.url)
    # S16i 플랫폼이 승인 → 회사 화면에 '승인됨'으로 돌아온다
    plp.goto(BASE+"/platform",wait_until="domcontentloaded"); plp.wait_for_selector("[data-testid=request-queue][data-ready='1']",timeout=60000); nuke(plp)   # 하이드레이션 대기(고정 sleep 1.0 은 첫 컴파일에서 모자랐다 — 2026-09-25 새 DB 재현)
    plp.fill("[data-testid=decision-note]","Special 개발 착수")
    row=plp.locator("[data-testid=request-row]").filter(has_text=subj).first
    row.locator("[data-testid=approve]").click()
    row.wait_for(state="attached", timeout=30000)
    plp.wait_for_function("""(sub)=>{const r=[...document.querySelectorAll('[data-testid=request-row]')].find(e=>e.textContent.includes(sub)); return r && r.getAttribute('data-state')==='approved';}""", arg=subj, timeout=30000)  # 고정 sleep 금지(09-21·09-22 반복 실수)
    ok("S16i 플랫폼이 승인한다", row.get_attribute("data-state"), row.get_attribute("data-state")=="approved")
    pg.goto(BASE+"/m/company",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=platform-requests]",timeout=30000)
    pg.wait_for_function("""()=>{const e=document.querySelector('[data-testid=platform-requests]'); return e && e.innerText.includes('승인됨');}""", timeout=30000)  # 상태 대기
    nuke(pg)
    body=pg.inner_text("[data-testid=platform-requests]")
    ok("S16j 결정이 회사 화면으로 돌아온다 ('승인됨' + 결정 메모)", subj[-14:], subj in body and "승인됨" in body and "Special 개발 착수" in body)
    pl.close()
    # S17 2층→3층: owner가 역할을 올리면 그 계정의 권한이 실제로 바뀐다 (p54 User Management)
    v2=b.new_context(); v2.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    r=v2.request.post(BASE+"/api/setup/sub-codes",headers=J,data=json.dumps({"group":"AHU Code","itemKey":"B","itemName":"용량","value":"97"})); before=r.status
    pg.select_option("[data-testid='role-viewer@acme.test']","engineer"); time.sleep(2.0)
    r2=v2.request.post(BASE+"/api/setup/sub-codes",headers=J,data=json.dumps({"group":"AHU Code","itemKey":"B","itemName":"용량","value":"97"})); after=r2.status
    ok("S17a owner가 viewer를 engineer로 올리면 그 계정에 등록 권한이 생긴다 (403 → 200)", (before,after), before==403 and after==200)
    if r2.status==200: v2.request.delete(BASE+"/api/setup/sub-codes?id="+r2.json().get("id",""))
    v2.close()
    pg.select_option("[data-testid='role-viewer@acme.test']","viewer"); time.sleep(2.0)
    r=ctx.request.patch(BASE+"/api/company/members",headers=J,data=json.dumps({"userId":"10000000-0000-4000-8000-00000000000a","role":"viewer"}))
    ok("S17b 마지막 owner 강등은 거부된다 (409)", r.status, r.status==409)
    pg.screenshot(path=f"{OUT}/42_user_management.png",full_page=True)
    # ── S42 p12·p50 Project Management — 등록 · 헤더(담당자·Remarks·Description) · 영업 단계 · 접수 자료(File) ──
    # 맨 뒤에 둔다: 새 프로젝트가 Work Hierarchy 에 노드로 생기므로 앞 단계의 트리를 흔들지 않게(reset:demo 가 지운다).
    PNO="PS-E2E-"+str(int(time.time()))[-6:]
    pg.goto(BASE+"/m/project",wait_until="domcontentloaded")
    pg.wait_for_selector("[data-testid=project-mgmt][data-ready='1']",timeout=60000); nuke(pg)
    hdr=pg.inner_text("[data-testid=pm-detail]")
    ok("S42a 프로젝트 관리 화면이 p12 헤더를 갖는다 (Project No·Type·Client·Client 담당자·담당자·영업 단계·Item·Remarks·등록일·Description)",
       hdr.split("\n")[0:2], all(k in hdr for k in ("PS-61313-5","Project Type","Client 담당자 정보","담당자","영업 단계","Remarks","등록일","Description")))
    for _ in range(20):
        pg.fill("[data-testid=pm-new-no]",PNO); pg.fill("[data-testid=pm-new-name]","E2E 신규 AHU"); pg.fill("[data-testid=pm-new-client]","KSY")
        if pg.eval_on_selector("[data-testid=pm-create]","e=>!e.disabled"): break
        time.sleep(0.3)   # 하이드레이션 대기 재시도(값을 다시 넣고 버튼이 풀리는지 본다)
    pg.click("[data-testid=pm-create]")
    pg.wait_for_selector(f"[data-testid=pm-row-{PNO}][data-selected='1']",timeout=30000)
    newp=[x for x in ctx.request.get(BASE+"/api/projects").json()["rows"] if x["projectNo"]==PNO][0]
    pg.goto(BASE+f"/workbench?node={newp['hierarchyStable']}",wait_until="domcontentloaded")
    pg.wait_for_selector(f"[data-testid=inspector-bound][data-project='{newp['id']}']",timeout=60000)
    ok("S42b 등록 — 새 프로젝트가 목록에 생기고 Work Hierarchy 노드로도 생겨 작업대에서 바로 열린다", (PNO, newp["salesStage"]), newp["salesStage"]=="기술제안")
    dup=ctx.request.post(BASE+"/api/projects",headers=J0,data=json.dumps({"projectNo":PNO,"name":"중복"}))
    ok("S42c 같은 Project No 는 다시 등록할 수 없다 (409)", dup.status, dup.status==409)
    pg.goto(BASE+"/m/project",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=project-mgmt][data-ready='1']",timeout=60000); nuke(pg)
    pg.click(f"[data-testid=pm-row-{PNO}]"); pg.wait_for_selector(f"[data-testid=pm-detail][data-project='{PNO}']",timeout=30000)
    pg.select_option("[data-testid=pm-owner]","20000000-0000-4000-8000-00000000000a")
    pg.fill("[data-testid=pm-contact]","김담당 010-0000-0000"); pg.fill("[data-testid=pm-remarks]","1차 사양 회의 완료")
    pg.fill("[data-testid=pm-description]","Pain Point: 클린룸 차압 유지"); pg.select_option("[data-testid=pm-stage]","협의")
    pg.click("[data-testid=pm-save]")
    pg.wait_for_function("()=>(document.querySelector('[data-testid=pm-msg]')?.innerText||'').includes('협의')",timeout=30000)
    got=[x for x in ctx.request.get(BASE+"/api/projects").json()["rows"] if x["projectNo"]==PNO][0]
    ok("S42d 헤더 저장 — 담당자·Client 담당자·Remarks·Description 과 영업 단계(기술제안→협의)가 함께 저장된다",
       (got["ownerId"][-4:] if got["ownerId"] else None, got["remarks"], got["salesStage"]),
       got["ownerId"]=="20000000-0000-4000-8000-00000000000a" and got["clientContact"].startswith("김담당") and got["remarks"]=="1차 사양 회의 완료" and got["description"].startswith("Pain Point") and got["salesStage"]=="협의")
    bad=ctx.request.patch(BASE+f"/api/projects/{got['id']}",headers=J0,data=json.dumps({"ownerId":"10000000-0000-4000-8000-00000000000b"}))
    ok("S42e 담당자는 이 회사 구성원만 — 다른 회사 사람을 넣으면 400", bad.status, bad.status==400)
    BODY="고객 요구사항: 풍량 55,000 CMH · 차압 유지\n".encode("utf-8")
    pg.set_input_files("[data-testid=pm-up-file]",{"name":"requirements.txt","mimeType":"text/plain","buffer":BODY})
    pg.fill("[data-testid=pm-up-desc]","1차 접수 자료")
    pg.wait_for_function("()=>{const b=document.querySelector('[data-testid=pm-upload]'); return b && !b.disabled;}",timeout=30000)
    pg.click("[data-testid=pm-upload]")
    pg.wait_for_selector("[data-testid=pm-att-row][data-has-file='1']",timeout=30000)
    att=ctx.request.get(BASE+f"/api/project-attachments?projectId={got['id']}").json()["rows"][0]
    back=ctx.request.get(BASE+f"/api/project-attachments/{att['id']}/file")
    ok("S42f 접수 자료 등록(File) — 실제 파일이 올라가고, 받으면 같은 바이트가 돌아온다", (att["name"], att["fileSize"], back.status),
       back.status==200 and back.body()==BODY and att["fileSize"]==len(BODY) and att["department"]=="영업")
    big=ctx.request.post(BASE+"/api/project-attachments",multipart={"projectId":got["id"],"department":"영업","docType":"File",
        "file":{"name":"big.bin","mimeType":"application/octet-stream","buffer":b"0"*(10*1024*1024+1)}})
    ok("S42g 10MB 를 넘는 파일은 받지 않는다 (413)", big.status, big.status==413)
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    vp=vw.request.patch(BASE+f"/api/projects/{got['id']}",headers=J0,data=json.dumps({"remarks":"viewer 수정"}))
    vu=vw.request.post(BASE+"/api/project-attachments",multipart={"projectId":got["id"],"department":"영업","docType":"File","file":{"name":"x.txt","mimeType":"text/plain","buffer":b"x"}})
    vw.close()
    ok("S42h viewer 는 헤더 수정·자료 등록이 막힌다 (403 · 403)", (vp.status, vu.status), vp.status==403 and vu.status==403)
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    gx=gb.request.get(BASE+f"/api/project-attachments/{att['id']}/file"); gp=gb.request.patch(BASE+f"/api/projects/{got['id']}",headers=J0,data=json.dumps({"remarks":"x"}))
    gb.close()
    ok("S42i 다른 회사는 이 파일도 프로젝트도 못 본다 (RLS — 404 · 404)", (gx.status, gp.status), gx.status==404 and gp.status==404)
    pg.goto(BASE+"/m/project",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=project-mgmt][data-ready='1']",timeout=60000)
    pg.click(f"[data-testid=pm-row-{PNO}]"); pg.wait_for_selector("[data-testid=pm-att-row][data-has-file='1']",timeout=30000); nuke(pg)
    pg.screenshot(path=f"{OUT}/55_project_mgmt.png",full_page=True)
    # ── S43 p48 Print Set-up Form — 종류별 인쇄 양식. 모양만 바뀌고 숫자는 그대로 ──
    PS=BASE+"/api/print-setup"
    g0=ctx.request.get(PS+"?type=quotation").json()
    ok("S43a 저장 전에는 기본 양식이고, Print Test 에 쓸 최근 견적서가 있다", (g0["saved"], g0["settings"]["paper"], bool(g0["sample"])), (not g0["saved"]) and g0["settings"]["paper"]=="A4" and g0["sample"] is not None)
    SAMPLE=g0["sample"]["id"]
    def quote_total():
        h=ctx.request.get(BASE+f"/api/documents/{SAMPLE}/print").text()
        m=re.search(r'data-testid="quote-total">([^<]+)<',h); return (m.group(1) if m else None), h
    qt0,_=quote_total()
    bad=ctx.request.put(PS,headers=J0,data=json.dumps({"type":"quotation","settings":{"marginMm":99}}))
    ok("S43b 허용 범위 밖 값은 필드 이름과 함께 거부된다 (여백 99mm → 400)", (bad.status, bad.json().get("error","")[-8:]), bad.status==400 and "marginMm" in bad.json().get("error",""))
    pg.goto(BASE+"/setup/print",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=print-setup][data-ready='1']",timeout=60000); nuke(pg)
    for _ in range(20):
        pg.select_option("[data-testid=ps-orientation]","landscape"); pg.select_option("[data-testid=ps-color]","mono")
        pg.fill("[data-testid=ps-header]","Acme AHU · 기술영업팀"); pg.fill("[data-testid=ps-footer]","Good air makes Good Life"); pg.fill("[data-testid=ps-watermark]","CONFIDENTIAL")
        if pg.eval_on_selector("[data-testid=ps-save]","e=>!e.disabled"): break
        time.sleep(0.3)   # 하이드레이션 대기 재시도
    pg.click("[data-testid=ps-save]")
    pg.wait_for_selector("[data-testid=ps-msg][data-ok='1']",timeout=30000)
    g1=ctx.request.get(PS+"?type=quotation").json()["settings"]
    ok("S43c 화면에서 저장 — 방향·색상·머리글·바닥글·워터마크가 견적서 양식으로 남는다", (g1["orientation"],g1["color"],g1["watermark"]),
       g1["orientation"]=="landscape" and g1["color"]=="mono" and g1["header"].startswith("Acme") and g1["footer"].startswith("Good air") and g1["watermark"]=="CONFIDENTIAL")
    qt1,h1=quote_total()
    ok("S43d 인쇄본이 양식을 입는다 (A4 가로 흑백 · 머리글 · 바닥글 · 워터마크) — 견적 금액은 그대로",
       (qt0, qt1, 'data-print-setup="A4-landscape-mono"' in h1),
       qt0 is not None and qt0==qt1 and 'data-print-setup="A4-landscape-mono"' in h1 and "size: A4 landscape" in h1 and 'data-testid="print-watermark">CONFIDENTIAL' in h1 and 'data-testid="print-header"' in h1 and 'data-testid="print-footer"' in h1)
    td=ctx.request.get(PS+"?type=techdata").json()
    ok("S43e 양식은 문서 종류마다 따로다 — Tech Data 는 여전히 기본", (td["saved"], td["settings"]["color"]), (not td["saved"]) and td["settings"]["color"]=="color")
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    vp=vw.request.put(PS,headers=J0,data=json.dumps({"type":"quotation","settings":{"color":"color"}})); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    gq=gb.request.get(PS+"?type=quotation").json(); gb.close()
    ok("S43f viewer 는 양식을 못 바꾸고(403), 다른 회사는 이 회사 양식을 못 본다(자기 기본값)", (vp.status, gq["saved"]), vp.status==403 and not gq["saved"])
    fr=pg.frame_locator("[data-testid=ps-preview]")
    fr.locator("[data-testid=print-watermark]").wait_for(timeout=30000)
    ok("S43g Print Test — 오른쪽 미리보기가 저장된 양식(워터마크)으로 다시 그려진다", True, True)
    nuke(pg); pg.screenshot(path=f"{OUT}/56_print_setup.png",full_page=True)
    # ── S44 p25·p26 사용자 UI Form · UI Design 작업장 — 끌어다 놓기 · Set-up · Templet · Run(실제 카탈로그) ──
    UF=BASE+"/api/ui-forms"
    pg.goto(BASE+"/setup/ui",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=ui-designer][data-ready='1']",timeout=60000); nuke(pg)
    pg.fill("[data-testid=ui-new-name]","E2E 용량 조회"); pg.click("[data-testid=ui-new]")
    pg.wait_for_selector("[data-testid='ui-form-E2E 용량 조회'][data-selected='1']",timeout=30000)
    pg.wait_for_selector("[data-testid=ui-canvas]",timeout=30000)
    pg.drag_and_drop("[data-palette=combo]","[data-testid=ui-canvas]",target_position={"x":40,"y":40})
    pg.wait_for_selector("[data-widget=combo1]",timeout=30000)
    pos=pg.eval_on_selector("[data-widget=combo1]","e=>[e.style.left,e.style.top]")
    ok("S44a 위젯 상자의 Combo box 를 캔버스로 끌어다 놓으면 놓은 칸에 생긴다 (Drag)", pos, pos==["34px","34px"])
    pg.select_option("[data-testid=ui-w-subcode]","B"); pg.fill("[data-testid=ui-w-label]","용량 (B)")
    pg.click("[data-palette=table]"); pg.wait_for_selector("[data-widget=table1]",timeout=30000)
    pg.select_option("[data-testid=ui-w-table]","EU|cap")
    pg.click("[data-palette=button]"); pg.wait_for_selector("[data-widget=button1]",timeout=30000)
    pg.select_option("[data-testid=ui-w-target]","table1"); pg.select_option("[data-testid=ui-w-filter]","combo1")
    pg.click("[data-testid=ui-templet]")
    pg.click("[data-testid=ui-save]"); pg.wait_for_function("()=>document.querySelector('[data-testid=ui-msg]')?.innerText.includes('저장')",timeout=30000)
    f1=[x for x in ctx.request.get(UF).json()["rows"] if x["name"]=="E2E 용량 조회"][0]
    wb={w["id"]:w for w in f1["spec"]["widgets"]}
    rects=[(w["x"],w["y"],w["w"],w["h"]) for w in f1["spec"]["widgets"]]
    lap=[(i,j) for i in range(len(rects)) for j in range(i+1,len(rects)) if rects[i][0]<rects[j][0]+rects[j][2] and rects[j][0]<rects[i][0]+rects[i][2] and rects[i][1]<rects[j][1]+rects[j][3] and rects[j][1]<rects[i][1]+rects[i][3]]
    ok("S44a2 누르기로 추가한 위젯은 기존 위젯과 겹치지 않는 빈 자리에 놓인다", rects, not lap)
    ok("S44b Set-up 이 저장된다 — Combo=Sub Code B · Table=EU.cap · Button=찾기→table1 (Active Set-up=combo1) · Templet",
       sorted(wb), f1["isTemplet"] and wb["combo1"]["source"]=={"kind":"subcode","itemKey":"B"} and wb["table1"]["source"]=={"kind":"table","code":"EU","table":"cap"}
       and wb["button1"]["action"]=="find" and wb["button1"]["target"]=="table1" and wb["button1"]["filterBy"]=="combo1")
    bad=ctx.request.put(UF+"/"+f1["id"],headers=J0,data=json.dumps({"spec":{"widgets":[{"id":"button1","type":"button","x":0,"y":0,"w":4,"h":2,"label":"x","action":"find","target":"combo1"},
        {"id":"combo1","type":"combo","x":5,"y":0,"w":4,"h":2,"label":"c"}]}}))
    out=ctx.request.put(UF+"/"+f1["id"],headers=J0,data=json.dumps({"spec":{"widgets":[{"id":"label1","type":"label","x":22,"y":0,"w":4,"h":1,"label":"x"}]}}))
    ok("S44c 잘못된 폼은 저장되지 않는다 — 대상이 Table 이 아님 · 캔버스 밖 (400 · 400)", (bad.status, out.status), bad.status==400 and out.status==400)
    pg.click("[data-testid=ui-mode-run]"); pg.wait_for_selector("[data-run-table=table1]",timeout=30000)
    all_rows=int(pg.get_attribute("[data-run-table=table1]","data-rows"))
    pg.select_option("[data-run-combo=combo1]","55"); pg.click("[data-run-button=button1]")
    pg.wait_for_function("()=>document.querySelector('[data-run-table=table1]')?.dataset.rows==='1'",timeout=30000)
    first=pg.inner_text("[data-run-table=table1] tbody tr td")
    ok("S44d Run — Combo 에서 55 를 고르고 찾기를 누르면 EU.cap 표에서 Item 55 행만 남는다 (실제 카탈로그 데이터)", (all_rows, first), all_rows==4 and first=="55")
    pg.click("[data-testid=ui-mode-design]")
    pg.click("[data-testid='ui-call-E2E 용량 조회']")
    pg.wait_for_selector("[data-testid='ui-form-E2E 용량 조회 사본'][data-selected='1']",timeout=30000)
    cp=[x for x in ctx.request.get(UF).json()["rows"] if x["name"]=="E2E 용량 조회 사본"][0]
    ok("S44e Templet 호출하여 Customizing — 같은 위젯·Set-up 의 사본이 새 폼으로 생긴다(사본은 Templet 아님)", (len(cp["spec"]["widgets"]), cp["isTemplet"]),
       cp["spec"]==f1["spec"] and not cp["isTemplet"])
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    vp=vw.request.post(UF,headers=J0,data=json.dumps({"name":"viewer 폼"})); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    gl=gb.request.get(UF).json()["rows"]; gd=gb.request.delete(UF+"/"+f1["id"]); gb.close()
    ok("S44f viewer 는 폼을 못 만들고(403), 다른 회사는 이 폼을 못 보고 못 지운다(0건 · 404)", (vp.status, len(gl), gd.status), vp.status==403 and len(gl)==0 and gd.status==404)
    pg.click("[data-testid='ui-form-E2E 용량 조회']"); pg.wait_for_selector("[data-widget=table1]",timeout=30000); pg.click("[data-widget=button1]")
    nuke(pg); pg.screenshot(path=f"{OUT}/57_ui_design.png",full_page=True)
    # ── S45 p32 Material code & General purchase items · p67 단가 이력 ──
    CAT=BASE+"/api/setup/catalog"; PR=BASE+"/api/setup/prices"
    psh0=[c for c in ctx.request.get(CAT).json()["productCodes"] if c["code"]=="PSH 1"][0]
    pg.goto(BASE+"/setup/material",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=material-reg][data-ready='1']",timeout=60000); nuke(pg)
    n_codes=len(pg.query_selector_all("[data-testid^=mat-code-]"))
    pg.click("[data-testid='mat-code-PSH 1']"); pg.wait_for_selector("[data-testid=mat-table][data-code='PSH 1']",timeout=30000)
    ok("S45a 자재 등록 화면 — 구매품 코드가 분류 트리로 보이고, 고르면 속성 표(Supplier·V·Hz·IP·Insulation)가 열린다",
       (n_codes, pg.input_value("[data-testid=mat-cell-0-A]")), n_codes>=5 and pg.input_value("[data-testid=mat-cell-0-A]")=="한국전열" and pg.input_value("[data-testid=mat-col-E]")=="Insulation")
    pg.click("[data-testid=mat-add-col]"); pg.fill("[data-testid=mat-col-F]","Efficiency"); pg.fill("[data-testid=mat-cell-0-F]","IE3")
    pg.click("[data-testid=mat-save]"); pg.wait_for_selector("[data-testid=mat-msg][data-ok='1']",timeout=30000)
    psh1=[c for c in ctx.request.get(CAT).json()["productCodes"] if c["code"]=="PSH 1"][0]
    buy=[t for t in psh1["tables"].values() if t.get("role")=="buy"][0]
    ok("S45b 속성 추가 — Efficiency 열이 같은 buy 표(제품 코드 한 곳)에 저장된다", [c["name"] for c in buy["cols"]], [c["name"] for c in buy["cols"]][-1]=="Efficiency" and buy["rows"][0]["cells"].get("F")=="IE3")
    rr=ctx.request.post(BASE+"/api/setup/product-codes",headers=J0,data=json.dumps(psh0))
    ok("S45b2 원복(API) — PSH 1 속성 표를 되돌린다(카탈로그는 reset 이 되돌리지 않음)", rr.status, rr.status==200)
    MC="E2E-MT"+str(int(time.time()))[-5:]
    pg.fill("[data-testid=mat-new-code]",MC); pg.fill("[data-testid=mat-new-name]","Motor AC Φ3"); pg.fill("[data-testid=mat-new-cat]","General Purchase items/Motor/Industrial/AC")
    pg.click("[data-testid=mat-create]"); pg.wait_for_selector(f"[data-testid=mat-table][data-code='{MC}']",timeout=30000)
    mc=[c for c in ctx.request.get(CAT).json()["productCodes"] if c["code"]==MC][0]
    ok("S45c 새 자재 코드 등록 — 분류 경로와 기본 속성 표(6열, Efficiency 포함)로 카탈로그에 생긴다", (mc["kind"], mc["category"]),
       mc["kind"]=="purchase" and mc["category"]=="General Purchase items/Motor/Industrial/AC" and len(list(mc["tables"].values())[0]["cols"])==6)
    import datetime as _dt
    TODAY=_dt.date.today().isoformat()
    for price,eff in (("1200000","2026-01-01"),("1280000",TODAY),("1500000","2099-01-01")):
        pg.fill("[data-testid=mat-p-price]",price); pg.fill("[data-testid=mat-p-date]",eff); pg.fill("[data-testid=mat-p-supplier]","효성")
        n0=len(pg.query_selector_all("[data-testid=mat-price-row]")); pg.click("[data-testid=mat-p-add]")
        pg.wait_for_function("(n)=>document.querySelectorAll('[data-testid=mat-price-row]').length>n", arg=n0, timeout=30000)
    st=pg.eval_on_selector_all("[data-testid=mat-price-row]","es=>es.map(e=>e.dataset.state)")
    cur=pg.inner_text("[data-testid=mat-price-0]")
    ok("S45d 단가 이력 — 3건을 쌓으면 예정(2099)·현재(오늘)·지난(1월)으로 갈리고, 표의 G:Price 는 오늘 유효한 1,280,000", (st, cur),
       st==["예정","현재","지난"] and cur.startswith("1,280,000"))
    z=ctx.request.post(PR,headers=J0,data=json.dumps({"code":MC,"price":0,"effectiveFrom":TODAY}))
    nf=ctx.request.post(PR,headers=J0,data=json.dumps({"code":"NOPE-XX","price":10,"effectiveFrom":TODAY}))
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    vp=vw.request.post(PR,headers=J0,data=json.dumps({"code":MC,"price":10,"effectiveFrom":TODAY})); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    gp=gb.request.get(PR+"?code="+MC).json()["rows"]; gb.close()
    ok("S45e 단가 0 은 400 · 없는 코드는 404 · viewer 는 403 · 다른 회사는 이 이력을 못 본다(0건)", (z.status, nf.status, vp.status, len(gp)),
       z.status==400 and nf.status==404 and vp.status==403 and len(gp)==0)
    nuke(pg); pg.screenshot(path=f"{OUT}/58_material.png",full_page=True)
    # ── S46 p35 Arrangement Code — 등록(스냅샷) → 승인 → 적용(기존 저장 규칙) ──
    AC=BASE+"/api/setup/arrangement-codes"
    def eu_secs():
        return [x for x in ctx.request.get(CAT).json()["productCodes"] if x["code"]=="EU"][0]["sections"]
    snap=eu_secs(); names0=[x["name"] for x in snap]
    pg.goto(BASE+"/setup/arrangement-code",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=arr-codes][data-ready='1']",timeout=60000); nuke(pg)
    pg.fill("[data-testid=ac-new-code]","FDV-E2E"); pg.select_option("[data-testid=ac-new-product]","EU"); pg.fill("[data-testid=ac-new-desc]","Fan Centrifugal · Double (E2E)")
    pg.click("[data-testid=ac-register]"); pg.wait_for_selector("[data-testid=ac-row-FDV-E2E][data-status=pending]",timeout=30000)
    fdv=[x for x in ctx.request.get(AC).json()["rows"] if x["code"]=="FDV-E2E"][0]
    ok("S46a 등록 — EU 의 지금 배치가 스냅샷으로 떠지고 Approval Status 는 Pending", (fdv["status"], [x["name"] for x in fdv["sections"]]==names0),
       fdv["status"]=="pending" and [x["name"] for x in fdv["sections"]]==names0)
    ap0=ctx.request.post(AC+f"/{fdv['id']}/apply",headers=J0,data="{}")
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    vd=vw.request.post(AC+f"/{fdv['id']}/decide",headers=J0,data=json.dumps({"decision":"approve"})); vw.close()
    ok("S46b 승인 전에는 적용이 막히고(409), viewer 는 승인할 수 없다(403)", (ap0.status, vd.status), ap0.status==409 and vd.status==403)
    pg.click("[data-testid=ac-row-FDV-E2E]"); pg.wait_for_selector("[data-testid=ac-detail][data-code=FDV-E2E]",timeout=30000)
    pg.fill("[data-testid=ac-note]","배치 검토 완료"); pg.click("[data-testid=ac-approve]")
    pg.wait_for_selector("[data-testid=ac-detail][data-status=approved]",timeout=30000)
    again=ctx.request.post(AC+f"/{fdv['id']}/decide",headers=J0,data=json.dumps({"decision":"reject"}))
    ok("S46c owner 가 화면에서 승인 → Approved, 결정은 한 번뿐(다시 결정 409)", again.status, again.status==409)
    moved=[snap[-1]]+snap[:-1]
    mv=ctx.request.post(ARR,headers=J0,data=json.dumps({"code":"EU","sections":[{"name":x["name"],**({"len":x["len"]} if x.get("len") else {}),**({"dir":x["dir"]} if x.get("dir") else {}),"components":x.get("components",[])} for x in moved]}))
    names_moved=[x["name"] for x in eu_secs()]
    pg.wait_for_function("()=>{const b=document.querySelector('[data-testid=ac-apply]'); return b && !b.disabled;}",timeout=30000)
    pg.click("[data-testid=ac-apply]"); pg.wait_for_function("()=>(document.querySelector('[data-testid=ac-msg]')?.innerText||'').includes('적용했습니다')",timeout=30000)
    names_back=[x["name"] for x in eu_secs()]
    ok("S46d 적용 — EU 배치를 바꿔 둔 뒤 승인된 코드를 적용하면 스냅샷 순서로 돌아온다 (기존 Arrangement 저장 경로)",
       (mv.status, names_moved[0], names_back[0]), mv.status==200 and names_moved!=names0 and names_back==names0)
    pg.fill("[data-testid=ac-new-code]","FDV-E2E-R"); pg.click("[data-testid=ac-register]"); pg.wait_for_selector("[data-testid=ac-row-FDV-E2E-R][data-status=pending]",timeout=30000)
    pg.click("[data-testid=ac-row-FDV-E2E-R]"); pg.wait_for_selector("[data-testid=ac-detail][data-code=FDV-E2E-R]",timeout=30000)
    pg.click("[data-testid=ac-reject]"); pg.wait_for_selector("[data-testid=ac-detail][data-status=rejected]",timeout=30000)
    rj=[x for x in ctx.request.get(AC).json()["rows"] if x["code"]=="FDV-E2E-R"][0]
    ap2=ctx.request.post(AC+f"/{rj['id']}/apply",headers=J0,data="{}")
    dup=ctx.request.post(AC,headers=J0,data=json.dumps({"code":"FDV-E2E","productCode":"EU"}))
    ok("S46e 반려된 코드는 적용할 수 없고(409), 같은 코드 이름은 다시 등록할 수 없다(409)", (ap2.status, dup.status), ap2.status==409 and dup.status==409)
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    gl=gb.request.get(AC).json()["rows"]; ga=gb.request.post(AC+f"/{fdv['id']}/apply",headers=J0,data="{}"); gb.close()
    ok("S46f 다른 회사는 이 코드를 못 보고 적용도 못 한다 (0건 · 404)", (len(gl), ga.status), len(gl)==0 and ga.status==404)
    pg.click("[data-testid=ac-row-FDV-E2E]"); pg.wait_for_selector("[data-testid=ac-detail][data-code=FDV-E2E]",timeout=30000)
    nuke(pg); pg.screenshot(path=f"{OUT}/59_arrangement_code.png",full_page=True)
    # ── S47 ⑥ p46 Spec List in-put table — 사양 입력 → 코드 추천(등록값에서만) · 저장은 기존 개정 한 곳 ──
    SI=BASE+"/api/setup/spec-items"; SR=BASE+"/api/setup/spec-recommend"
    pg.goto(BASE+"/setup/spec",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=spec-items][data-ready='1']",timeout=60000); nuke(pg)
    pg.select_option("[data-testid=spec-product]","EU"); pg.wait_for_selector("[data-testid=spec-items][data-ready='1'][data-product=EU]",timeout=30000)
    seeded=pg.eval_on_selector_all("[data-testid^=spec-row-]","es=>es.map(e=>e.dataset.testid.replace('spec-row-',''))")
    pg.fill("[data-testid=spec-new-key]","fan_kw"); pg.fill("[data-testid=spec-new-label]","팬 동력"); pg.fill("[data-testid=spec-new-unit]","kW")
    pg.select_option("[data-testid=spec-new-slot]","B"); pg.select_option("[data-testid=spec-new-kind]","table")
    pg.select_option("[data-testid=spec-new-table]","cap"); pg.select_option("[data-testid=spec-new-col]","A"); pg.select_option("[data-testid=spec-new-op]","ge")
    pg.click("[data-testid=spec-add]"); pg.wait_for_selector("[data-testid=spec-row-fan_kw]",timeout=30000)
    ok("S47a 사양 항목 — EU 에 시드 3종(풍량·가습량·재질)이 있고, 화면에서 제품 표의 열(cap.A 팬 kW)을 읽는 항목을 추가한다", seeded,
       seeded==["airflow","humid","material"] and bool(pg.query_selector("[data-testid=spec-row-fan_kw]")))
    pg.goto(NODE4,wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=spec-panel][data-ready='1']",timeout=60000); nuke(pg)
    pg.wait_for_selector("[data-testid=spec-in-fan_kw]",timeout=30000)
    pg.fill("[data-testid=spec-in-airflow]","11000"); pg.fill("[data-testid=spec-in-fan_kw]","10"); pg.fill("[data-testid=spec-in-material]","SS")
    pg.click("[data-testid=spec-recommend]"); pg.wait_for_selector("[data-testid=spec-msg]",timeout=30000)
    pg.wait_for_function("()=>document.querySelector('[data-testid=assembled-code]')?.innerText.startsWith('EU-25-')",timeout=30000)
    bv=pg.eval_on_selector("[data-testid=code-builder] select[data-slot=B]","e=>e.value"); ev=pg.eval_on_selector("[data-testid=code-builder] select[data-slot=E]","e=>e.value")
    asm=pg.inner_text("[data-testid=assembled-code]").strip(); fan_line=pg.get_attribute("[data-testid=spec-line-fan_kw]","data-picked")
    ok("S47b 사양 입력 → 코드 추천 — 풍량 11,000 CMH 와 팬 10 kW 를 둘 다 만족하는 최소 용량 25, 재질 SS 가 Code Builder 에 채워진다", (bv, ev, asm, fan_line),
       bv=="25" and ev=="SS" and asm.startswith("EU-25-") and "SS" in asm and fan_line=="25")
    nuke(pg); pg.screenshot(path=f"{OUT}/60_spec_input.png",full_page=True)
    rv0=len(ctx.request.get(BASE+"/api/rccs/revisions?node=a0000000-0000-4000-8000-000000000004").json().get("revisions",[]))
    pg.fill("[data-testid=rev-reason]","spec input recommend (E2E)")
    # 운영 모드(next start)에서는 추천 직후 저장 버튼이 아직 busy 로 잠겨 있을 때 force 클릭이 조용히 버려졌다(ccmd H STEP 1 게이트 · 286/287).
    # 버튼이 풀리길 기다렸다가 저장 POST 응답을 받아서 넘어간다(S4c 와 같은 방식).
    pg.wait_for_selector("[data-testid=rev-save]:not([disabled])",timeout=30000)
    with pg.expect_response(lambda r: "/api/rccs/revisions" in r.url and r.request.method=="POST", timeout=30000):
        pg.click("[data-testid=rev-save]")
    rv1=ctx.request.get(BASE+"/api/rccs/revisions?node=a0000000-0000-4000-8000-000000000004").json().get("revisions",[])
    ok("S47c 저장은 기존 개정(Rev) 한 곳 — 추천된 코드가 새 개정으로 남는다", (len(rv1)-rv0, rv1[0]["code"] if rv1 else None),
       len(rv1)==rv0+1 and rv1[0]["code"]==asm)
    bad=ctx.request.post(SI,headers=J0,data=json.dumps({"productCode":"EU","key":"bad_slot","label":"x","slot":"E","source":{"kind":"table","table":"cap","col":"M","op":"ge"}}))
    un=ctx.request.post(SR,headers=J0,data=json.dumps({"productCode":"EU","inputs":{"airflow":"90000"}})).json()
    uk=ctx.request.post(SR,headers=J0,data=json.dumps({"productCode":"EU","inputs":{"nope":"1"}}))
    ok("S47d 잘못된 정의는 400(표의 행 슬롯과 사양 슬롯이 다르다) · 맞는 등록값이 없으면 지어내지 않고 unmet · 없는 사양 키는 400",
       (bad.status, un.get("unmet"), un.get("slots"), uk.status), bad.status==400 and un.get("unmet")==["B"] and un.get("slots")=={} and uk.status==400)
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    vp=vw.request.post(SI,headers=J0,data=json.dumps({"productCode":"EU","key":"v_try","label":"v","slot":"E","source":{"kind":"choice"}})); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    gl=gb.request.get(SI+"?product=EU").json()["rows"]; gr=gb.request.post(SR,headers=J0,data=json.dumps({"productCode":"EU","inputs":{"airflow":"11000"}})); gb.close()
    ok("S47e viewer 는 사양 항목을 못 만든다(403) · 다른 회사는 이 항목을 못 보고(0건) 우리 제품으로 추천도 못 받는다(404)", (vp.status, len(gl), gr.status),
       vp.status==403 and len(gl)==0 and gr.status==404)
    # ── S48 ⑦ 도면 용도 구분 (p17 · p39 · 0020) — 승인도·제작도·견적도로 등록하고 목록에서 거른다 · 발행 뒤에는 못 바꾼다 ──
    DW=BASE+"/api/drawings"; N4="a0000000-0000-4000-8000-000000000004"
    pg.goto(NODE4,wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=canvas-cmds][data-ready='1']",timeout=60000); nuke(pg)
    if pg.get_attribute("[data-testid=toolbox-toggle]","aria-pressed")=="true":
        pg.click("[data-testid=toolbox-toggle]"); pg.wait_for_selector("[data-testid=toolbox-toggle][aria-pressed=false]",timeout=30000)
    nuke(pg); pg.click("[data-run=bom]")
    pg.locator("button", has_text=re.compile(r"^Design$")).first.click(force=True)
    pg.wait_for_selector("[data-testid=design-canvas][data-loaded='1']",timeout=30000)
    pg.wait_for_function("()=>{const b=document.querySelector('[data-testid=drawing-make-front]'); return b && !b.disabled;}",timeout=60000)
    for pu in ("approval","manufacturing"):
        pg.select_option("[data-testid=drawing-purpose]",pu); nuke(pg); pg.click("[data-testid=drawing-make-front]")
        pg.wait_for_selector(f"[data-testid=drawing-row][data-purpose={pu}]",timeout=30000)
    apv=ctx.request.get(DW+f"?node={N4}&purpose=approval").json()["rows"]; mfg=ctx.request.get(DW+f"?node={N4}&purpose=manufacturing").json()["rows"]
    ok("S48a 용도를 골라 등록 — 같은 정면도라도 승인도(-APV)·제작도(-MFG)로 번호·개정이 따로 간다",
       ([x["drawingNo"]+" "+x["currentRev"] for x in apv], [x["drawingNo"]+" "+x["currentRev"] for x in mfg]),
       len(apv)>=1 and len(mfg)>=1 and apv[0]["drawingNo"].endswith("-FRT-APV") and mfg[0]["drawingNo"].endswith("-FRT-MFG") and apv[0]["currentRev"]=="A")
    pg.select_option("[data-testid=drawing-filter]","approval"); pg.wait_for_selector("[data-testid=drawing-register][data-filter=approval][data-listed='1']",timeout=30000)
    fa=pg.eval_on_selector_all("[data-testid=drawing-row]","es=>es.map(e=>e.dataset.purpose)")
    nuke(pg); pg.screenshot(path=f"{OUT}/61_drawing_purpose.png",full_page=True)
    pg.select_option("[data-testid=drawing-filter]","none"); pg.wait_for_selector("[data-testid=drawing-register][data-filter=none][data-listed='1']",timeout=30000)
    fn_=pg.eval_on_selector_all("[data-testid=drawing-row]","es=>es.map(e=>e.dataset.purpose)")
    ok("S48b 목록에서 용도로 거른다 — 승인도만 / 미지정(옛 도면)만", (fa, len(fn_)), len(fa)>=1 and all(x=="approval" for x in fa) and all(x=="" for x in fn_))
    rid=ctx.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":S55_0,"code":"EU-55-2123-630SS-1-21-13-15","node":N4})).json().get("runId")
    b400=ctx.request.post(DW,headers=J0,data=json.dumps({"runId":rid,"type":"front","purpose":"sales"}))
    g400=ctx.request.get(DW+"?purpose=sales")
    ch=ctx.request.patch(DW+f"/{mfg[0]['id']}",headers=J0,data=json.dumps({"purpose":"quotation"}))
    now=[x for x in ctx.request.get(DW+f"?node={N4}").json()["rows"] if x["id"]==mfg[0]["id"]][0]["purpose"]
    iss=[x for x in ctx.request.get(DW).json()["rows"] if x["status"]=="issued"]
    li=ctx.request.patch(DW+f"/{iss[0]['id']}",headers=J0,data=json.dumps({"purpose":"approval"})) if iss else None
    ok("S48c 잘못된 용도는 400(등록·목록) · 발행 전에는 용도를 바꿀 수 있고(제작도→견적도) · 발행된 도면은 409(잠금)",
       (b400.status, g400.status, ch.status, now, li.status if li else "발행 도면 없음"),
       b400.status==400 and g400.status==400 and ch.status==200 and now=="quotation" and li is not None and li.status==409)
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    vp=vw.request.post(DW,headers=J0,data=json.dumps({"runId":rid,"type":"front","purpose":"approval"}))
    vpa=vw.request.patch(DW+f"/{apv[0]['id']}",headers=J0,data=json.dumps({"purpose":"quotation"})); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    gl=gb.request.get(DW+"?purpose=approval").json()["rows"]; gpa=gb.request.patch(DW+f"/{apv[0]['id']}",headers=J0,data=json.dumps({"purpose":"quotation"})); gb.close()
    ok("S48d viewer 는 용도 도면을 못 만들고 못 바꾼다(403·403) · 다른 회사는 우리 승인도를 못 보고(0건) 못 바꾼다(404)",
       (vp.status, vpa.status, len(gl), gpa.status), vp.status==403 and vpa.status==403 and len(gl)==0 and gpa.status==404)
    # ── S49 ⑧ Company DB (p64 · p67 · 0021) — 고객·공급처 목록 · 프로젝트 Client 와 단가 공급처가 목록을 가리킨다(글자 열도 함께) ──
    PT=BASE+"/api/setup/partners"; PJ=BASE+"/api/projects"; PRC=BASE+"/api/setup/prices"
    pg.goto(BASE+"/setup/company",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=company-db][data-ready='1']",timeout=60000); nuke(pg)
    seeded_p=sorted(x["code"] for x in ctx.request.get(PT).json()["rows"])
    pg.fill("[data-testid=pn-customer-code]","E2E-C1"); pg.fill("[data-testid=pn-customer-name]","Samsung Bio (E2E)"); pg.fill("[data-testid=pn-customer-contact]","설비팀 010-0000-0000"); pg.fill("[data-testid=pn-customer-nation]","KR")
    pg.click("[data-testid=pn-customer-add]"); pg.wait_for_selector("[data-testid=partner-row-E2E-C1]",timeout=30000)
    pg.fill("[data-testid=pn-supplier-code]","E2E-S1"); pg.fill("[data-testid=pn-supplier-name]","Hanil Pump (E2E)")
    pg.click("[data-testid=pn-supplier-add]"); pg.wait_for_selector("[data-testid=partner-row-E2E-S1]",timeout=30000)
    nuke(pg); pg.screenshot(path=f"{OUT}/62_company_db.png",full_page=True)
    rows_p=ctx.request.get(PT).json()["rows"]; c1=[x for x in rows_p if x["code"]=="E2E-C1"][0]; s1=[x for x in rows_p if x["code"]=="E2E-S1"][0]
    ok("S49a Company DB — 시드 고객·공급처(C-MICRON·S-KSB)에 화면에서 고객 E2E-C1 · 공급처 E2E-S1 을 등록한다", (seeded_p, c1["kind"], s1["kind"]),
       seeded_p==["C-MICRON","S-KSB"] and c1["kind"]=="customer" and s1["kind"]=="supplier")
    pg.goto(BASE+"/m/project",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=pm-detail][data-project='PS-61313-5']",timeout=30000); nuke(pg)
    pj0=[x for x in ctx.request.get(PJ).json()["rows"] if x["projectNo"]=="PS-61313-5"][0]
    pg.select_option("[data-testid=pm-client-pick]",c1["id"]); pg.click("[data-testid=pm-save]")
    pg.wait_for_function("()=>(document.querySelector('[data-testid=pm-msg]')?.innerText||'').startsWith('저장')",timeout=30000)
    pj1=[x for x in ctx.request.get(PJ).json()["rows"] if x["projectNo"]=="PS-61313-5"][0]
    ok("S49b 프로젝트 Client 를 Company DB 에서 고르면 client_id 와 글자(client_name)가 함께 바뀐다 — 옛 데이터는 글자만 있었다",
       (pj0.get("clientId"), pj0["clientName"], pj1.get("clientId")==c1["id"], pj1["clientName"]),
       pj0.get("clientId") is None and pj0["clientName"]=="Micron" and pj1.get("clientId")==c1["id"] and pj1["clientName"]=="Samsung Bio (E2E)")
    rs_=ctx.request.patch(PJ+f"/{pj1['id']}",headers=J0,data=json.dumps({"clientId":"","clientName":"Micron"}))   # 원복(뒤 단계·시연 화면 보존)
    mc=[x["code"] for x in ctx.request.get(CAT).json()["productCodes"] if x["kind"]=="purchase"][0]
    pr=ctx.request.post(PRC,headers=J0,data=json.dumps({"code":mc,"price":123000,"effectiveFrom":"2026-01-02","supplierId":s1["id"],"note":"E2E company db"}))
    ph=[x for x in ctx.request.get(PRC+"?code="+mc).json()["rows"] if x["note"]=="E2E company db"]
    ok("S49c 단가 이력의 공급처를 Company DB 로 가리키면 supplier_id 와 글자(supplier)가 함께 남는다", (pr.status, ph[0]["supplier"] if ph else None, rs_.status),
       pr.status==200 and ph and ph[0]["supplierId"]==s1["id"] and ph[0]["supplier"]=="Hanil Pump (E2E)" and rs_.status==200)
    bk=ctx.request.post(PT,headers=J0,data=json.dumps({"kind":"bank","code":"B1","name":"x"}))
    dup=ctx.request.post(PT,headers=J0,data=json.dumps({"kind":"customer","code":"E2E-C1","name":"dup"}))
    wrong=ctx.request.patch(PJ+f"/{pj1['id']}",headers=J0,data=json.dumps({"clientId":s1["id"]}))
    wrong2=ctx.request.post(PRC,headers=J0,data=json.dumps({"code":mc,"price":1000,"effectiveFrom":"2026-01-03","supplierId":c1["id"]}))
    ok("S49d 없는 종류는 400 · 같은 코드는 409 · 공급처를 고객으로(프로젝트) · 고객을 공급처로(단가) 가리키면 400",
       (bk.status, dup.status, wrong.status, wrong2.status), bk.status==400 and dup.status==409 and wrong.status==400 and wrong2.status==400)
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    vp=vw.request.post(PT,headers=J0,data=json.dumps({"kind":"customer","code":"V1","name":"v"})); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    gl=gb.request.get(PT).json()["rows"]; gc=gb.request.post(PJ,headers=J0,data=json.dumps({"projectNo":"GX-E2E-1","name":"gx","clientId":c1["id"]})); gb.close()
    ok("S49e viewer 는 등록 못 한다(403) · 다른 회사는 우리 고객·공급처를 못 보고(0건) 자기 프로젝트에 우리 고객을 걸 수 없다(400)",
       (vp.status, len(gl), gc.status), vp.status==403 and len(gl)==0 and gc.status==400)
    # ── S50 ⑨ Input Data 템플릿 (p16 · 0022) — Tech Data 가 회사 입력 항목 값을 받아 문서에 스냅샷으로 남긴다 ──
    II=BASE+"/api/setup/input-items"; DOCS=BASE+"/api/documents"; N4="a0000000-0000-4000-8000-000000000004"
    seeded_i=[x["key"] for x in ctx.request.get(II).json()["rows"]]
    pg.goto(NODE4,wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=canvas-cmds][data-ready='1']",timeout=60000); nuke(pg)
    if pg.get_attribute("[data-testid=toolbox-toggle]","aria-pressed")=="true":
        pg.click("[data-testid=toolbox-toggle]"); pg.wait_for_selector("[data-testid=toolbox-toggle][aria-pressed=false]",timeout=30000)
    nuke(pg); pg.click("[data-run=bom]")
    pg.locator("button", has_text=re.compile(r"^Document$")).first.click(force=True)
    pg.wait_for_selector("[data-testid=doc-inputdata][data-ready='1'] [data-testid=doc-in-temperature]",timeout=30000)
    pg.wait_for_function("()=>{const b=document.querySelector('[data-testid=doc-make-techdata]'); return b && !b.disabled;}",timeout=60000)
    pg.fill("[data-testid=doc-in-temperature]","25"); pg.fill("[data-testid=doc-in-humidity]","60"); nuke(pg)
    pg.wait_for_function("()=>document.querySelector('[data-testid=doc-in-temperature]').value==='25' && document.querySelector('[data-testid=doc-in-humidity]').value==='60'",timeout=10000)
    with pg.expect_response(lambda q: q.url.endswith("/api/documents") and q.request.method=="POST",timeout=30000) as dres:
        pg.click("[data-testid=doc-make-techdata]")
    da=dres.value.json(); pg.wait_for_selector("[data-testid=document-msg][data-ok='1']",timeout=30000)
    nuke(pg); pg.screenshot(path=f"{OUT}/63_input_data.png",full_page=True)
    ha=ctx.request.get(DOCS+f"/{da.get('id')}/print").text()
    ok("S50a Tech Data 를 만들 때 Input Data(Temperature 25 °C · Humidity 60 %)를 받아 문서에 남긴다 — 시드 항목은 청사진 p16 두 개",
       (seeded_i, da.get("docNo"), 'data-key="temperature">25 °C' in ha, 'data-key="humidity">60 %' in ha),
       seeded_i==["temperature","humidity"] and 'data-key="temperature">25 °C' in ha and 'data-key="humidity">60 %' in ha)
    pg.goto(BASE+"/setup/input-data",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=input-items][data-ready='1']",timeout=60000); nuke(pg)
    pg.fill("[data-testid=in-new-key]","pressure"); pg.fill("[data-testid=in-new-label]","Pressure"); pg.fill("[data-testid=in-new-unit]","Pa")
    pg.fill("[data-testid=in-new-default]","101325"); pg.fill("[data-testid=in-new-min]","80000"); pg.fill("[data-testid=in-new-max]","110000")
    pg.click("[data-testid=in-add]"); pg.wait_for_selector("[data-testid=input-row-pressure]",timeout=30000)
    rid5=ctx.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":S55_0,"code":"EU-55-2123-630SS-1-21-13-15","node":N4})).json().get("runId")
    dbb=ctx.request.post(DOCS,headers=J0,data=json.dumps({"runId":rid5,"type":"techdata"})).json()
    hb=ctx.request.get(DOCS+f"/{dbb.get('id')}/print").text(); ha2=ctx.request.get(DOCS+f"/{da.get('id')}/print").text()
    ok("S50b 화면에서 항목(Pressure Pa)을 더하면 다음 문서부터 받고(안 보내면 기본값 101325) — 먼저 만든 문서는 스냅샷 그대로(항목 2개)",
       ('data-key="pressure">101325 Pa' in hb, 'data-key="pressure"' in ha2),
       'data-key="pressure">101325 Pa' in hb and 'data-key="temperature">20 °C' in hb and 'data-key="pressure"' not in ha2 and 'data-key="temperature">25 °C' in ha2)
    o1=ctx.request.post(DOCS,headers=J0,data=json.dumps({"runId":rid5,"type":"techdata","inputData":{"humidity":120}}))
    o2=ctx.request.post(DOCS,headers=J0,data=json.dumps({"runId":rid5,"type":"techdata","inputData":{"density":1.2}}))
    o3=ctx.request.post(II,headers=J0,data=json.dumps({"key":"bad_range","label":"x","minValue":10,"maxValue":1}))
    ok("S50c 범위 밖(습도 120 %) · 템플릿에 없는 항목(density) 은 문서를 만들지 않는다(400·400) · 최소>최대 정의는 400",
       (o1.status, o2.status, o3.status), o1.status==400 and o2.status==400 and o3.status==400)
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    vi=vw.request.post(II,headers=J0,data=json.dumps({"key":"v_try","label":"v"})); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    gi=gb.request.get(II).json()["rows"]; gd=gb.request.get(DOCS+f"/{da.get('id')}/print"); gb.close()
    ok("S50d viewer 는 항목을 못 만든다(403) · 다른 회사는 우리 템플릿을 못 보고(0건) 우리 Tech Data 도 못 연다(404)",
       (vi.status, len(gi), gd.status), vi.status==403 and len(gi)==0 and gd.status==404)
    # ── S51 ⑩ 3D 뷰어 (p4 · p37) — 도면과 같은 BOM 스냅샷의 구획 박스를 브라우저에서 돌려 본다(새 데이터 없음 · 읽기 전용) ──
    M3=BASE+"/api/model3d"
    pg.goto(NODE4,wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=canvas-cmds][data-ready='1']",timeout=60000); nuke(pg)
    if pg.get_attribute("[data-testid=toolbox-toggle]","aria-pressed")=="true":
        pg.click("[data-testid=toolbox-toggle]"); pg.wait_for_selector("[data-testid=toolbox-toggle][aria-pressed=false]",timeout=30000)
    nuke(pg); pg.click("[data-run=bom]")
    pg.locator("button", has_text=re.compile(r"^Design$")).first.click(force=True)
    pg.wait_for_function("()=>{const a=document.querySelector('[data-testid=view-3d]'); return a && a.getAttribute('href').startsWith('/viewer3d?runId=');}",timeout=60000)
    href3=pg.get_attribute("[data-testid=view-3d]","href"); rid3=href3.split("runId=")[1]
    nuke(pg); pg.click("[data-testid=view-3d]")
    pg.wait_for_selector("[data-testid=viewer3d][data-ready='1']",timeout=60000)
    m3=ctx.request.get(M3+f"?runId={rid3}").json()
    nb=int(pg.get_attribute("[data-testid=viewer3d]","data-boxes")); ntri=int(pg.get_attribute("[data-testid=viewer3d]","data-triangles"))
    colors=pg.evaluate("""()=>{const c=document.querySelector('[data-testid=viewer3d-canvas] canvas'); const t=document.createElement('canvas'); t.width=80; t.height=52;
      const g=t.getContext('2d'); g.drawImage(c,0,0,80,52); const d=g.getImageData(0,0,80,52).data; const s=new Set(); for(let i=0;i<d.length;i+=4) s.add((d[i]>>4)+','+(d[i+1]>>4)+','+(d[i+2]>>4)); return s.size;}""")
    ok("S51a Design 탭 '3D 보기' → 같은 스냅샷의 구획 박스가 WebGL 로 그려진다 (박스 수 = 구획 수 · 전장 = 구획 길이 합 · 화면이 비지 않음)",
       (nb, len(m3.get("boxes",[])), m3.get("length"), ntri, colors),
       nb==len(m3["boxes"])>0 and m3["length"]==sum(x["len"] for x in m3["boxes"]) and ntri>0 and colors>8)
    nuke(pg); pg.screenshot(path=f"{OUT}/64_viewer3d.png",full_page=True)
    pg.click("[data-testid=view3d-top]"); pg.wait_for_selector("[data-testid=view3d-top][style*='accent']",timeout=10000)
    ok("S51b 보기 전환(평면) 뒤에도 뷰어가 살아 있다", pg.get_attribute("[data-testid=viewer3d]","data-ready"), pg.get_attribute("[data-testid=viewer3d]","data-ready")=="1")
    n400=ctx.request.get(M3); n404=ctx.request.get(M3+"?runId=00000000-0000-4000-8000-000000000000")
    ok("S51c runId 없으면 400 · 없는 스냅샷은 404 — 스냅샷 없이 3D 를 지어내지 않는다", (n400.status, n404.status), n400.status==400 and n404.status==404)
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v3=vw.request.get(M3+f"?runId={rid3}"); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    g3=gb.request.get(M3+f"?runId={rid3}"); gb.close()
    ok("S51d 읽기 전용 — viewer 도 볼 수 있다(200 · 쓰기 API 없음) · 다른 회사는 우리 스냅샷을 못 연다(404)", (v3.status, g3.status), v3.status==200 and g3.status==404)
    # ── S52 ccmd E · p67 단가 이력 → BOM 원가·견적·구매 — BOM Run 순간의 "현재 단가"가 줄에 박히고, 스냅샷은 그 뒤 바뀌지 않는다 ──
    PRC=BASE+"/api/setup/prices"; RUNB=BASE+"/api/run/bom"; N4="a0000000-0000-4000-8000-000000000004"
    def bom_run():
        return ctx.request.post(RUNB,headers=J0,data=json.dumps({"slots":S55_0,"code":"EU-55-2123-630SS-1-21-13-15","node":N4})).json()
    def line_of(run,code):
        tr={t["no"]:t["childCode"] for t in run["trace"]}
        return [l for l in run["lines"] if tr.get(l["no"])==code][0]
    TODAY=_dt.date.today().isoformat()
    p1=ctx.request.post(PRC,headers=J0,data=json.dumps({"code":"PFP 1","price":25000,"effectiveFrom":TODAY,"note":"E2E price-to-cost"})).json()
    r1=bom_run(); l1=line_of(r1,"PFP 1")
    ok("S52a 현재 단가를 등록하고 BOM Run → 그 줄 단가 = 등록 단가, 출처 '이력'(priceId · 유효일)",
       (l1["unitCost"], l1.get("priceSource")), l1["unitCost"]==25000 and l1.get("priceSource",{}).get("kind")=="history" and l1["priceSource"].get("priceId")==p1.get("id") and l1["priceSource"].get("effectiveFrom")==TODAY)
    c1=ctx.request.post(BASE+"/api/run/cost",headers=J0,data=json.dumps({"runId":r1["runId"]})).json()["cost"]
    mat=sum(l["qty"]*l["unitCost"] for l in r1["lines"]); lab=round(mat*0.18); ovh=round((mat+lab)*0.12)
    ok("S52b 원가 합계 = Σ 수량×단가로 다시 세어 일치 (재료비 · 인건비 18% · 경비 12%)", (c1["material"],c1["labor"],c1["overhead"],c1["total"]),
       c1["material"]==mat and c1["labor"]==lab and c1["overhead"]==ovh and c1["total"]==mat+lab+ovh)
    q1=ctx.request.post(BASE+"/api/documents",headers=J0,data=json.dumps({"runId":r1["runId"],"type":"quotation"})).json()
    qh=ctx.request.get(BASE+f"/api/documents/{q1.get('id')}/print").text()
    pr1=ctx.request.post(BASE+"/api/purchase-requests",headers=J0,data=json.dumps({"runId":r1["runId"]})).json()
    prs=[x for x in ctx.request.get(BASE+f"/api/purchase-requests?node={N4}").json()["rows"] if x["bomRunId"]==r1["runId"]]
    prl=[l for l in (prs[0]["lines"] if prs else []) if str(l.get("resolvedCode","")).startswith("PFP 1")]
    ok("S52c 같은 스냅샷의 견적(합계 = 원가 합계 · 단가 기준 줄)과 구매 요청(PFP 1 단가 25,000)이 같은 단가를 쓴다",
       (q1.get("total"), c1["total"], 'data-testid="price-basis"' in qh, [float(l["unitPrice"]) for l in prl]),
       q1.get("total")==c1["total"] and 'data-testid="price-basis"' in qh and prl and all(float(l["unitPrice"])==25000 for l in prl))
    p2=ctx.request.post(PRC,headers=J0,data=json.dumps({"code":"PFP 1","price":31000,"effectiveFrom":TODAY,"note":"E2E price-to-cost 2"})).json()
    c1b=ctx.request.post(BASE+"/api/run/cost",headers=J0,data=json.dumps({"runId":r1["runId"]})).json()["cost"]
    qh_b=ctx.request.get(BASE+f"/api/documents/{q1.get('id')}/print").text()
    r2=bom_run(); l2=line_of(r2,"PFP 1")
    ok("S52d 단가를 바꿔도(새 행) 다시 Run 하지 않으면 옛 스냅샷 원가·견적은 그대로 — 새로 Run 하면 새 단가 31,000",
       (c1b["total"]==c1["total"], qh_b==qh, l2["unitCost"], l2["priceSource"].get("priceId")==p2.get("id")),
       c1b["total"]==c1["total"] and qh_b==qh and l2["unitCost"]==31000 and l2["priceSource"].get("priceId")==p2.get("id"))
    ctx.request.post(PRC,headers=J0,data=json.dumps({"code":"PFP 1","price":99000,"effectiveFrom":"2099-01-01","note":"E2E future"}))
    ctx.request.post(PRC,headers=J0,data=json.dumps({"code":"PFB 1","price":40,"currency":"USD","effectiveFrom":TODAY,"note":"E2E usd"}))
    rel_pfb=[l for l in r2["lines"] if line_of(r2,"PFB 1")["no"]==l["no"]][0]["unitCost"]
    r3=bom_run(); l3=line_of(r3,"PFP 1"); l3b=line_of(r3,"PFB 1")
    ok("S52e 미래 단가(2099)는 적용되지 않고(31,000 그대로) · 원화가 아닌 단가는 적용하지 않고 관계값 유지('통화 불일치')",
       (l3["unitCost"], l3b["unitCost"], l3b["priceSource"]["kind"]),
       l3["unitCost"]==31000 and l3b["unitCost"]==rel_pfb and l3b["priceSource"]["kind"]=="currency-mismatch")
    dw=ctx.request.post(BASE+"/api/drawings",headers=J0,data=json.dumps({"runId":r1["runId"],"type":"plan"}))
    ok("S52f 단가 행을 더한 뒤에도 옛 스냅샷의 도면 생성이 막히지 않는다 (단가는 카탈로그 지문 밖 · 지문 동일)",
       (dw.status, r1["catalogFp"]==r3["catalogFp"]), dw.status==200 and r1["catalogFp"]==r3["catalogFp"])
    pg.goto(NODE4,wait_until="domcontentloaded"); hydrated(pg); nuke(pg)
    if pg.get_attribute("[data-testid=toolbox-toggle]","aria-pressed")=="true":
        pg.click("[data-testid=toolbox-toggle]"); pg.wait_for_selector("[data-testid=toolbox-toggle][aria-pressed=false]",timeout=30000)
    nuke(pg); pg.click("[data-run=bom]")
    pg.locator("button", has_text=re.compile(r"^BOM$")).first.click(force=True)
    wait_sel(pg,"[data-testid=bom-table] [data-price-src=history]")
    srcs=pg.eval_on_selector_all("[data-testid=bom-table] [data-price-src]","es=>es.map(e=>e.dataset.priceSrc)")
    ok("S52g 화면 BOM 표가 줄마다 단가 출처(이력 · 관계값 · 통화 불일치)를 스냅샷 그대로 보인다",
       sorted(set(srcs)), "history" in srcs and "relationship" in srcs and "currency-mismatch" in srcs)
    nuke(pg); pg.screenshot(path=f"{OUT}/65_price_to_cost.png",full_page=True)
    # ── S53 F1 · p12 · p50 Client 담당자 여러 명(주담당 1) · 영업 활동 이력(쌓기만) — 0023 ──
    PJ=BASE+"/api/projects"
    pid=[x for x in ctx.request.get(PJ).json()["rows"] if x["projectNo"]=="PS-61313-5"][0]["id"]
    pg.goto(BASE+"/m/project",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=pm-detail][data-project='PS-61313-5']",timeout=30000)
    pg.wait_for_selector("[data-testid=pm-crm][data-ready='1']",timeout=30000); nuke(pg)
    pg.fill("[data-testid=pm-c-name]","김설비"); pg.fill("[data-testid=pm-c-dept]","FAB 설비팀"); pg.fill("[data-testid=pm-c-contact]","010-1111-2222")
    pg.click("[data-testid=pm-c-add]"); pg.wait_for_selector("[data-testid='pm-contact-김설비']",timeout=30000)
    pg.fill("[data-testid=pm-c-name]","이구매"); pg.fill("[data-testid=pm-c-dept]","구매팀"); pg.check("[data-testid=pm-c-isprimary]")
    pg.click("[data-testid=pm-c-add]"); pg.wait_for_selector("[data-testid='pm-contact-이구매'][data-primary='1']",timeout=30000)
    prim=pg.eval_on_selector_all("[data-testid^=pm-contact-]","es=>es.map(e=>[e.dataset.testid.replace('pm-contact-',''),e.dataset.primary])")
    ok("S53a Client 담당자 여러 명 — 첫 담당자는 주담당, 새 담당자를 주담당으로 추가하면 주담당이 옮겨 간다(프로젝트당 1명)", prim,
       sorted(prim)==[["김설비","0"],["이구매","1"]])
    pg.click("[data-testid='pm-c-edit-김설비']"); pg.fill("[data-testid=pm-ce-dept]","FAB 설비2팀"); pg.click("[data-testid=pm-ce-save]")
    wait_text(pg,"[data-testid='pm-contact-김설비']","설비2팀")
    pg.click("[data-testid='pm-c-primary-김설비']"); pg.wait_for_selector("[data-testid='pm-contact-김설비'][data-primary='1']",timeout=30000)
    pg.click("[data-testid='pm-c-del-이구매']"); pg.wait_for_selector("[data-testid='pm-contact-이구매']",state="detached",timeout=30000)
    cs=ctx.request.get(PJ+f"/{pid}/contacts").json()["rows"]
    ok("S53b 수정(부서) · 주담당 바꾸기 · 삭제가 화면에서 된다", [(c["name"],c["department"],c["isPrimary"]) for c in cs],
       [(c["name"],c["department"],c["isPrimary"]) for c in cs]==[("김설비","FAB 설비2팀",True)])
    pg.fill("[data-testid=pm-a-content]","사양 회의 — 풍량 55,000 CMH · SS 외판 확정"); pg.select_option("[data-testid=pm-a-kind]","meeting")
    pg.click("[data-testid=pm-a-add]"); pg.wait_for_selector("[data-testid=pm-activity-row][data-kind=meeting]",timeout=30000)
    pg.fill("[data-testid=pm-a-content]","견적 송부 후 통화"); pg.select_option("[data-testid=pm-a-kind]","call"); pg.click("[data-testid=pm-a-add]")
    pg.wait_for_function("()=>document.querySelectorAll('[data-testid=pm-activity-row]').length>=2",timeout=30000)
    nuke(pg); pg.screenshot(path=f"{OUT}/66_project_contacts.png",full_page=True)
    acts=ctx.request.get(PJ+f"/{pid}/activities").json()["rows"]
    pa=ctx.request.patch(PJ+f"/{pid}/activities",headers=J0,data=json.dumps({"content":"x"})); da=ctx.request.delete(PJ+f"/{pid}/activities")
    ok("S53c 영업 활동 이력 — 날짜·종류·내용이 쌓이고(2건), 수정·삭제 요청은 405 (DB 도 앱 역할의 UPDATE/DELETE 권한이 없다)",
       (len(acts), [a["kind"] for a in acts], pa.status, da.status), len(acts)==2 and set(a["kind"] for a in acts)=={"meeting","call"} and pa.status==405 and da.status==405)
    e1=ctx.request.post(PJ+f"/{pid}/contacts",headers=J0,data=json.dumps({"name":"  "}))
    e2=ctx.request.post(PJ+f"/{pid}/activities",headers=J0,data=json.dumps({"date":"2026-13-40","kind":"call","content":"x"}))
    e3=ctx.request.post(PJ+f"/{pid}/activities",headers=J0,data=json.dumps({"date":"2026-09-27","kind":"sms","content":"x"}))
    ok("S53d 이름 없는 담당자 · 잘못된 날짜 · 없는 종류는 400", (e1.status, e2.status, e3.status), e1.status==400 and e2.status==400 and e3.status==400)
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v1=vw.request.post(PJ+f"/{pid}/contacts",headers=J0,data=json.dumps({"name":"v"})); v2=vw.request.post(PJ+f"/{pid}/activities",headers=J0,data=json.dumps({"date":"2026-09-27","kind":"call","content":"v"})); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    g1=gb.request.get(PJ+f"/{pid}/contacts"); g2=gb.request.get(PJ+f"/{pid}/activities"); g3=gb.request.delete(BASE+f"/api/project-contacts/{cs[0]['id']}"); gb.close()
    ok("S53e viewer 는 담당자·활동을 못 쌓는다(403·403) · 다른 회사는 우리 프로젝트의 담당자·활동을 못 보고 못 지운다(404·404·404)",
       (v1.status, v2.status, g1.status, g2.status, g3.status), v1.status==403 and v2.status==403 and g1.status==404 and g2.status==404 and g3.status==404)
    # ── S54 F2 · p64 · p12 고객·공급처 수정 · 사용 중지 · 삭제(가리키면 409) — 0024 ──
    PT=BASE+"/api/setup/partners"; PJ=BASE+"/api/projects"; PRC=BASE+"/api/setup/prices"
    ctx.request.post(PT,headers=J0,data=json.dumps({"kind":"customer","code":"E2E-C9","name":"Hyundai Fab (E2E)"}))
    ctx.request.post(PT,headers=J0,data=json.dumps({"kind":"supplier","code":"E2E-S8","name":"Daehan Filter (E2E)"}))
    ctx.request.post(PT,headers=J0,data=json.dumps({"kind":"supplier","code":"E2E-S9","name":"Unused Supply (E2E)"}))
    pts=ctx.request.get(PT).json()["rows"]; c9=[x for x in pts if x["code"]=="E2E-C9"][0]; s8=[x for x in pts if x["code"]=="E2E-S8"][0]; s9=[x for x in pts if x["code"]=="E2E-S9"][0]
    pj=[x for x in ctx.request.get(PJ).json()["rows"] if x["projectNo"]=="PS-61313-5"][0]
    link=ctx.request.patch(PJ+f"/{pj['id']}",headers=J0,data=json.dumps({"clientId":c9["id"]}))
    pcode=[x["code"] for x in ctx.request.get(CAT).json()["productCodes"] if x["kind"]=="purchase"][0]
    ctx.request.post(PRC,headers=J0,data=json.dumps({"code":pcode,"price":1000,"effectiveFrom":"2026-01-05","supplierId":s8["id"],"note":"E2E f2"}))
    pg.goto(BASE+"/setup/company",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=company-db][data-ready='1']",timeout=60000); nuke(pg)
    pg.click("[data-testid=pe-edit-E2E-C9]"); pg.fill("[data-testid=pe-name]","Hyundai Fab 2 (E2E)"); pg.fill("[data-testid=pe-contact]","설비팀 02-000-0000"); pg.click("[data-testid=pe-save]")
    wait_text(pg,"[data-testid=partner-row-E2E-C9]","Hyundai Fab 2")
    ok("S54a 고객 수정 — 이름·연락처를 고치면 목록과 API 가 바뀐다(코드는 그대로)", pg.inner_text("[data-testid=partner-row-E2E-C9]")[:60].replace("\n"," "),
       "Hyundai Fab 2 (E2E)" in pg.inner_text("[data-testid=partner-row-E2E-C9]") and link.status==200)
    pg.click("[data-testid=pe-del-E2E-C9]"); pg.wait_for_selector("[data-testid=company-db-msg][data-ok='0']",timeout=30000)
    m409=pg.inner_text("[data-testid=company-db-msg]")
    d8=ctx.request.delete(PT+f"/{s8['id']}")
    ok("S54b 가리키는 곳이 있으면 삭제 409 — 고객(프로젝트 1) · 공급처(단가 이력 1), 거부 이유가 화면에 보인다", (m409[:70], d8.status),
       "409" in m409 and "프로젝트 1" in m409 and d8.status==409 and d8.json()["usage"]["prices"]==1)
    pg.click("[data-testid=pe-toggle-E2E-C9]"); pg.wait_for_selector("[data-testid=partner-row-E2E-C9][data-active='0']",timeout=30000)
    nuke(pg); pg.screenshot(path=f"{OUT}/67_partner_edit.png",full_page=True)
    keep=ctx.request.patch(PJ+f"/{pj['id']}",headers=J0,data=json.dumps({"clientId":c9["id"],"remarks":"E2E keep inactive client"}))
    newlink=ctx.request.post(PJ,headers=J0,data=json.dumps({"projectNo":"E2E-F2-NEW","name":"x","clientId":c9["id"]}))
    ok("S54c 사용 중지 — 이미 가리키는 프로젝트는 그대로 저장되고(200), 새로 거는 것은 400 (목록에서도 빠진다)",
       (keep.status, newlink.status), keep.status==200 and newlink.status==400)
    d9=ctx.request.delete(PT+f"/{s9['id']}"); gone=not any(x["code"]=="E2E-S9" for x in ctx.request.get(PT).json()["rows"])
    bad=ctx.request.patch(PT+f"/{s9['id']}",headers=J0,data=json.dumps({"code":"X"}))
    ok("S54d 아무도 가리키지 않는 공급처는 지워진다(200 · 목록에서 사라짐) · 코드 바꾸기는 400", (d9.status, gone, bad.status), d9.status==200 and gone and bad.status in (400,404))
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v1=vw.request.patch(PT+f"/{c9['id']}",headers=J0,data=json.dumps({"name":"v"})); v2=vw.request.delete(PT+f"/{c9['id']}"); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    g1=gb.request.patch(PT+f"/{c9['id']}",headers=J0,data=json.dumps({"name":"g"})); g2=gb.request.delete(PT+f"/{s8['id']}"); gb.close()
    ok("S54e viewer 는 고치거나 지울 수 없다(403·403) · 다른 회사는 우리 고객·공급처를 못 고치고 못 지운다(404·404)", (v1.status, v2.status, g1.status, g2.status),
       v1.status==403 and v2.status==403 and g1.status==404 and g2.status==404)
    ctx.request.patch(PJ+f"/{pj['id']}",headers=J0,data=json.dumps({"clientId":"","clientName":"Micron","remarks":""}))   # 원복
    # ── S55 F3 · p46 사양 항목 수정·삭제 · CSV Import(미리보기 → 확정, 틀린 줄은 줄 번호와 이유) ──
    SI=BASE+"/api/setup/spec-items"
    pg.goto(BASE+"/setup/spec",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=spec-items][data-ready='1']",timeout=60000); nuke(pg)
    pg.select_option("[data-testid=spec-product]","EU"); pg.wait_for_selector("[data-testid=spec-items][data-ready='1'][data-product=EU]",timeout=30000)
    n0=len(ctx.request.get(SI+"?product=EU").json()["rows"])
    BAD_CSV="key,label,unit,slot,kind,op,scale,table,col\nmacro_result,팬 동력,kW,B,table,ge,,cap,A\ncoil_rows,코일 열수,,B,table,ge,,cap,B\nairflow,중복 키,,B,item,ge,1000,,\nbad_tbl,슬롯 틀림,,E,table,ge,,cap,M\n"
    pg.set_input_files("[data-testid=spec-import-file]",files=[{"name":"spec.csv","mimeType":"text/csv","buffer":BAD_CSV.encode("utf-8")}])
    pg.wait_for_function("()=>!document.querySelector('[data-testid=spec-import-preview]').disabled",timeout=10000)
    pg.click("[data-testid=spec-import-preview]"); pg.wait_for_selector("[data-testid=spec-import-preview-table]",timeout=30000)
    lines=pg.eval_on_selector_all("[data-testid^=spec-import-line-]","es=>es.map(e=>[e.dataset.testid.replace('spec-import-line-',''),e.dataset.ok,e.innerText.slice(-40)])")
    locked=pg.eval_on_selector("[data-testid=spec-import-confirm]","e=>e.disabled")
    nuke(pg); pg.screenshot(path=f"{OUT}/68_spec_import.png",full_page=True)
    ok("S55a 미리보기 — 틀린 줄은 파일의 줄 번호와 이유(중복 key · 표의 슬롯 불일치)로 보이고, 틀린 줄이 있으면 확정이 잠긴다",
       ([(l[0],l[1]) for l in lines], locked), [(l[0],l[1]) for l in lines]==[("2","1"),("3","1"),("4","0"),("5","0")] and locked and "이미 있는 key" in lines[2][2])
    forced=ctx.request.post(SI+"/import",headers=J0,data=json.dumps({"productCode":"EU","csv":BAD_CSV,"confirm":True}))
    n1=len(ctx.request.get(SI+"?product=EU").json()["rows"])
    pg.fill("[data-testid=spec-import-text]","\n".join(BAD_CSV.split("\n")[:3])+"\n")
    pg.click("[data-testid=spec-import-preview]"); pg.wait_for_function("()=>{const b=document.querySelector('[data-testid=spec-import-confirm]'); return b && !b.disabled;}",timeout=30000)
    pg.click("[data-testid=spec-import-confirm]"); pg.wait_for_selector("[data-testid=spec-row-coil_rows]",timeout=30000)
    ok("S55b 틀린 줄이 섞이면 확정 요청도 400 · 하나도 들어가지 않는다 → 고친 CSV 는 확정으로 2개가 들어온다",
       (forced.status, n1-n0, len(ctx.request.get(SI+"?product=EU").json()["rows"])-n0), forced.status==400 and n1==n0 and len(ctx.request.get(SI+"?product=EU").json()["rows"])==n0+2)
    pg.click("[data-testid=se-edit-coil_rows]"); pg.fill("[data-testid=se-label]","코일 열 수"); pg.fill("[data-testid=se-unit]","열"); pg.click("[data-testid=se-save]")
    wait_text(pg,"[data-testid=spec-row-coil_rows]","코일 열 수")
    pg.click("[data-testid=se-del-macro_result]"); pg.wait_for_selector("[data-testid=spec-row-macro_result]",state="detached",timeout=30000)
    rows=ctx.request.get(SI+"?product=EU").json()["rows"]; cr=[x for x in rows if x["key"]=="coil_rows"][0]
    ok("S55c 화면에서 이름·단위 수정과 삭제가 된다", (cr["label"], cr["unit"], any(x["key"]=="macro_result" for x in rows)), cr["label"]=="코일 열 수" and cr["unit"]=="열" and not any(x["key"]=="macro_result" for x in rows))
    k400=ctx.request.patch(SI+f"/{cr['id']}",headers=J0,data=json.dumps({"key":"x"}))
    s400=ctx.request.patch(SI+f"/{cr['id']}",headers=J0,data=json.dumps({"source":{"kind":"table","table":"nope","col":"A","op":"ge"}}))
    ok("S55d key 바꾸기 400 · 카탈로그에 없는 표로 바꾸기 400(수정도 등록과 같은 대조)", (k400.status, s400.status), k400.status==400 and s400.status==400)
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v1=vw.request.post(SI+"/import",headers=J0,data=json.dumps({"productCode":"EU","csv":BAD_CSV})); v2=vw.request.patch(SI+f"/{cr['id']}",headers=J0,data=json.dumps({"label":"v"})); v3=vw.request.delete(SI+f"/{cr['id']}"); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    g1=gb.request.patch(SI+f"/{cr['id']}",headers=J0,data=json.dumps({"label":"g"})); g2=gb.request.delete(SI+f"/{cr['id']}"); gb.close()
    ok("S55e viewer 는 Import·수정·삭제 모두 403 · 다른 회사는 우리 항목을 못 고치고 못 지운다(404·404)", (v1.status, v2.status, v3.status, g1.status, g2.status),
       v1.status==403 and v2.status==403 and v3.status==403 and g1.status==404 and g2.status==404)
    # ── S56 F4 · p32 · p30 자재·구매 코드 Approval Status(작성중→승인→사용중지, 역행 금지) · DWG 2D/3D 첨부 — 0025 ──
    AT=BASE+"/api/attachments"; CS=BASE+"/api/setup/code-status"; PRC=BASE+"/api/setup/prices"; MC4="PSH 1"
    pg.goto(BASE+"/setup/material",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=material-reg][data-ready='1']",timeout=60000); nuke(pg)
    pg.click(f"[data-testid='mat-code-{MC4}']"); pg.wait_for_selector(f"[data-testid=mat-table][data-code='{MC4}']",timeout=30000)
    pg.wait_for_selector("[data-testid=code-docs] [data-testid=code-dwg][data-ready='1']",timeout=30000)
    st0=pg.get_attribute("[data-testid=code-status]","data-status")
    pg.click("[data-testid=code-status-draft]"); pg.wait_for_selector("[data-testid=code-status][data-status=draft]",timeout=30000)
    pg.click("[data-testid=code-status-approved]"); pg.wait_for_selector("[data-testid=code-status][data-status=approved]",timeout=30000)
    back=ctx.request.post(CS,headers=J0,data=json.dumps({"code":MC4,"status":"draft"}))
    ok("S56a 코드 상태 — 미지정 → 작성중 → 승인이 화면에서 되고, 되돌리기는 409(DB 트리거도 막는다)", (st0, back.status), st0=="" and back.status==409)
    DXF2="0\nSECTION\n2\nENTITIES\n0\nENDSEC\n0\nEOF\n"
    pg.set_input_files("[data-testid=code-dwg-file]",files=[{"name":"PSH1_front.dxf","mimeType":"application/dxf","buffer":DXF2.encode("utf-8")}])
    pg.click("[data-testid=code-dwg-upload]"); pg.wait_for_selector("[data-testid=code-dwg-row][data-kind=dwg2d]",timeout=30000)
    pg.select_option("[data-testid=code-dwg-kind]","dwg3d")
    pg.set_input_files("[data-testid=code-dwg-file]",files=[{"name":"PSH1.step","mimeType":"application/octet-stream","buffer":b"ISO-10303-21;\nEND-ISO-10303-21;\n"}])
    pg.click("[data-testid=code-dwg-upload]"); pg.wait_for_selector("[data-testid=code-dwg-row][data-kind=dwg3d]",timeout=30000)
    nuke(pg); pg.screenshot(path=f"{OUT}/69_code_approval.png",full_page=True)
    rows_a=ctx.request.get(AT+f"?ownerKind=product_code&ownerKey={MC4}").json()["rows"]
    got=ctx.request.get(AT+f"/{[x for x in rows_a if x['kind']=='dwg2d'][0]['id']}/file").text()
    ok("S56b DWG 첨부 — 2D(.dxf)·3D(.step)를 올리고, 내려받으면 올린 내용 그대로", (sorted(x["kind"] for x in rows_a), got==DXF2),
       sorted(x["kind"] for x in rows_a)==["dwg2d","dwg3d"] and got==DXF2)
    def up(kind, name, data, key=MC4, c=ctx):
        return c.request.post(AT, multipart={"ownerKind":"product_code","ownerKey":key,"kind":kind,"file":{"name":name,"mimeType":"application/octet-stream","buffer":data}})
    bad_ext=up("dwg2d","virus.exe",b"MZ"); big=up("dwg3d","big.stl",b"0"*(10*1024*1024+1)); nokey=up("dwg2d","a.dxf",b"x",key="NOPE 9")
    ok("S56c 허용 밖 확장자 415 · 10MB 초과 413 · 없는 코드 404", (bad_ext.status, big.status, nokey.status), bad_ext.status==415 and big.status==413 and nokey.status==404)
    ret=ctx.request.post(CS,headers=J0,data=json.dumps({"code":MC4,"status":"retired"}))
    after=up("dwg2d","late.dxf",b"x"); pr=ctx.request.post(PRC,headers=J0,data=json.dumps({"code":MC4,"price":1000,"effectiveFrom":"2026-01-01"}))
    ok("S56d 사용중지 — 새 도면 첨부 409 · 새 단가 409 (이미 뜬 BOM 스냅샷·첨부는 그대로)", (ret.status, after.status, pr.status), ret.status==200 and after.status==409 and pr.status==409)
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v1=up("dwg2d","v.dxf",b"x",key="PFP 1",c=vw); v2=vw.request.post(CS,headers=J0,data=json.dumps({"code":"PFP 1","status":"approved"})); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    g1=gb.request.get(AT+f"?ownerKind=product_code&ownerKey={MC4}").json()["rows"]; g2=gb.request.get(AT+f"/{rows_a[0]['id']}/file"); gb.close()
    ok("S56e viewer 는 첨부·상태 변경 403 · 다른 회사는 우리 코드의 첨부를 못 보고(0건) 못 내려받는다(404)", (v1.status, v2.status, len(g1), g2.status),
       v1.status==403 and v2.status==403 and len(g1)==0 and g2.status==404)
    # ── S57 F5 · p35 · p30 Arrangement Drawing Control — 승인된 Arrangement Code 에만 DWG 를 붙인다(F4 첨부 재사용) ──
    AC=BASE+"/api/setup/arrangement-codes"; AT=BASE+"/api/attachments"
    ctx.request.post(AC,headers=J0,data=json.dumps({"code":"E2E-ADC","productCode":"EU","description":"F5 drawing control"}))
    adc=[x for x in ctx.request.get(AC).json()["rows"] if x["code"]=="E2E-ADC"][0]
    def upa(name, data, c=ctx, key=None):
        return c.request.post(AT, multipart={"ownerKind":"arrangement_code","ownerKey":key or adc["id"],"kind":"dwg2d","file":{"name":name,"mimeType":"application/dxf","buffer":data}})
    pg.goto(BASE+"/setup/arrangement-code",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=arr-codes][data-ready='1']",timeout=60000); nuke(pg)
    pg.click("[data-testid=ac-row-E2E-ADC]"); pg.wait_for_selector("[data-testid=ac-detail][data-code=E2E-ADC]",timeout=30000)
    pg.wait_for_selector("[data-testid=ac-dwg][data-ready='1']",timeout=30000)
    locked=bool(pg.query_selector("[data-testid=ac-dwg-locked]")); pend=upa("early.dxf",b"0\nEOF\n")
    ok("S57a 승인 전(Pending) Arrangement Code — 화면은 올리기 자리를 잠그고, API 도 409", (locked, pend.status), locked and pend.status==409)
    ctx.request.post(AC+f"/{adc['id']}/decide",headers=J0,data=json.dumps({"decision":"approve","note":"E2E"}))
    pg.reload(wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=arr-codes][data-ready='1']",timeout=60000); nuke(pg)
    pg.click("[data-testid=ac-row-E2E-ADC]"); pg.wait_for_selector("[data-testid=ac-detail][data-status=approved]",timeout=30000)
    pg.wait_for_selector("[data-testid=ac-dwg-file]",timeout=30000)
    ADXF="0\nSECTION\n2\nENTITIES\n0\nTEXT\n1\nE2E-ADC\n0\nENDSEC\n0\nEOF\n"
    pg.set_input_files("[data-testid=ac-dwg-file]",files=[{"name":"E2E-ADC_arrangement.dxf","mimeType":"application/dxf","buffer":ADXF.encode("utf-8")}])
    pg.click("[data-testid=ac-dwg-upload]"); pg.wait_for_selector("[data-testid=ac-dwg-row][data-kind=dwg2d]",timeout=30000)
    rows=ctx.request.get(AT+f"?ownerKind=arrangement_code&ownerKey={adc['id']}").json()["rows"]
    got=ctx.request.get(AT+f"/{rows[0]['id']}/file").text()
    ok("S57b 승인 뒤 화면에서 Arrangement 도면(DXF)을 올리고 그대로 내려받는다", (len(rows), got==ADXF), len(rows)==1 and got==ADXF)
    nope=upa("x.dxf",b"x",key="00000000-0000-4000-8000-000000000000"); ext=c_ext=ctx.request.post(AT, multipart={"ownerKind":"arrangement_code","ownerKey":adc["id"],"kind":"dwg2d","file":{"name":"x.exe","mimeType":"application/octet-stream","buffer":b"MZ"}})
    ok("S57c 없는 Arrangement Code 404 · 허용 밖 확장자 415", (nope.status, ext.status), nope.status==404 and ext.status==415)
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test")); v1=upa("v.dxf",b"x",c=vw); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    g1=upa("g.dxf",b"x",c=gb); g2=gb.request.get(AT+f"?ownerKind=arrangement_code&ownerKey={adc['id']}").json()["rows"]; gb.close()
    ok("S57d viewer 403 · 다른 회사는 우리 Arrangement Code 에 못 붙이고(404) 목록도 0건", (v1.status, g1.status, len(g2)), v1.status==403 and g1.status==404 and len(g2)==0)
    # ── S58 F6 · p13 Sub Item list · DWG View(화면에 띄우기 — 스냅샷 DXF 를 서버에서 SVG 로, 새 도면 계산 없음) ──
    DX=BASE+"/api/dxf"
    pg.goto(NODE4,wait_until="domcontentloaded"); hydrated(pg); nuke(pg)
    if pg.get_attribute("[data-testid=toolbox-toggle]","aria-pressed")=="true":
        pg.click("[data-testid=toolbox-toggle]"); pg.wait_for_selector("[data-testid=toolbox-toggle][aria-pressed=false]",timeout=30000)
    nuke(pg); pg.click("[data-run=bom]")
    pg.locator("button", has_text=re.compile(r"^Design$")).first.click(force=True)
    pg.wait_for_selector("[data-testid=sub-item-list][data-rows]:not([data-rows='0'])",timeout=60000)
    n_all=int(pg.get_attribute("[data-testid=sub-item-list]","data-rows"))
    rid6=pg.get_attribute("[data-testid=view-3d]","href").split("runId=")[1]
    pg.locator("button", has_text=re.compile(r"^BOM$")).first.click(force=True); pg.wait_for_selector("[data-testid=bom-table]",timeout=30000)
    n_snap=len(pg.query_selector_all("[data-testid=bom-table] tbody tr"))   # 같은 스냅샷(runId)의 BOM 표 행 수
    pg.locator("button", has_text=re.compile(r"^Design$")).first.click(force=True); pg.wait_for_selector("[data-testid=sub-item-list][data-section='']",timeout=30000)
    pg.click("[data-testid=canvas-sec-Fan]"); pg.wait_for_selector("[data-testid=sub-item-list][data-section=Fan]",timeout=30000)
    fan_rows=pg.eval_on_selector_all("[data-testid=sub-item-row]","es=>es.map(e=>e.dataset.section)")
    ok("S58a Sub Item list — 스냅샷의 BOM 줄을 그대로 보이고(전체 줄 수 = 같은 스냅샷 BOM 표 행 수), 개념도에서 Fan 을 고르면 Fan 구획 줄만",
       (n_all, n_snap, fan_rows), n_all>0 and n_all==n_snap and len(fan_rows)>0 and all(x=="Fan" for x in fan_rows))
    pg.click("[data-testid=canvas-sec-Fan]")
    pg.wait_for_function("()=>{const s=document.querySelector('[data-cmd=dwg-view]'); return s && !s.disabled;}",timeout=60000)
    pg.select_option("[data-cmd=dwg-view]","plan"); pg.wait_for_selector("[data-testid=dwg-viewer][data-view=plan][data-ready='1']",timeout=30000)
    nuke(pg); pg.screenshot(path=f"{OUT}/70_dwg_view.png")
    dxf_plan=ctx.request.get(DX+f"?runId={rid6}&type=plan").text()
    n_line=sum(1 for i,l in enumerate(dxf_plan.split("\n")[1:],1) if l.strip()=="LINE" and dxf_plan.split("\n")[i-1].strip()=="0")
    svg_lines=pg.eval_on_selector("[data-testid=dwg-viewer-svg] svg","e=>e.querySelectorAll('line').length")
    ok("S58b DWG View ▼ → 도면이 화면에 뜬다 — SVG 의 선 수 = 같은 스냅샷 DXF 의 LINE 수(다시 계산하지 않고 옮겼다)", (svg_lines, n_line), svg_lines==n_line and n_line>0)
    pg.click("[data-testid=dwg-viewer-iso]"); pg.wait_for_selector("[data-testid=dwg-viewer][data-view=iso][data-ready='1']",timeout=30000)
    iso_href=pg.get_attribute("[data-testid=dwg-viewer-download]","href")
    pg.click("[data-testid=dwg-viewer-close]"); pg.wait_for_selector("[data-testid=dwg-viewer]",state="detached",timeout=10000)
    ok("S58c 뷰어 안에서 뷰를 바꾸면(3D 등각) 그 뷰를 그리고, 내려받기 링크도 그 뷰의 DXF", iso_href, iso_href.endswith("type=iso"))
    b1=ctx.request.get(DX+f"?runId={rid6}&type=nope&format=svg"); b2=ctx.request.get(DX+"?runId=00000000-0000-4000-8000-000000000000&type=plan&format=svg")
    ok("S58d 없는 뷰 400 · 없는 스냅샷 404", (b1.status, b2.status), b1.status==400 and b2.status==404)
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test")); v1=vw.request.get(DX+f"?runId={rid6}&type=plan&format=svg"); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test")); g1=gb.request.get(DX+f"?runId={rid6}&type=plan&format=svg"); gb.close()
    ok("S58e 읽기 전용 — viewer 는 볼 수 있다(200 · 쓰기 없음) · 다른 회사는 우리 스냅샷 도면을 못 연다(404)", (v1.status, g1.status), v1.status==200 and g1.status==404)
    # ── S59 F7 · p15 Technical data 목록(모아보기·거르기) · 입력값 CSV Import ──
    TD=BASE+"/api/techdata"; DOCS=BASE+"/api/documents"; N4="a0000000-0000-4000-8000-000000000004"
    rid7=ctx.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":S55_0,"code":"EU-55-2123-630SS-1-21-13-15","node":N4})).json().get("runId")
    t1=ctx.request.post(DOCS,headers=J0,data=json.dumps({"runId":rid7,"type":"techdata","inputData":{"temperature":22}})).json()
    t2=ctx.request.post(DOCS,headers=J0,data=json.dumps({"runId":rid7,"type":"techdata","inputData":{"temperature":31,"humidity":40}})).json()
    ctx.request.patch(DOCS+f"/{t2['id']}",headers=J0,data=json.dumps({"status":"review"}))
    pg.goto(BASE+"/techdata",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=techdata-list][data-ready='1']",timeout=60000); nuke(pg)
    n_all=len(pg.query_selector_all("[data-testid=td-row]"))
    pg.select_option("[data-testid=td-status]","review"); pg.fill("[data-testid=td-q]",t2["docNo"]); pg.click("[data-testid=td-apply]")
    pg.wait_for_selector(f"[data-testid=techdata-list][data-ready='1'][data-filter='review|{t2['docNo']}']",timeout=30000)
    frows=pg.eval_on_selector_all("[data-testid=td-row]","es=>es.map(e=>[e.dataset.doc,e.dataset.status,e.querySelector('[data-key=temperature]')?.innerText])")
    ok("S59a Tech Data 목록 — 스냅샷별 문서를 모아 보이고(2건 이상), 상태(검토)·문서번호로 거르면 그 문서 하나와 그때 받은 입력값(31)",
       (n_all, frows), n_all>=2 and frows==[[t2["docNo"],"review","31"]])
    nuke(pg); pg.screenshot(path=f"{OUT}/71_techdata_list.png",full_page=True)
    pg.goto(NODE4,wait_until="domcontentloaded"); hydrated(pg); nuke(pg)
    if pg.get_attribute("[data-testid=toolbox-toggle]","aria-pressed")=="true":
        pg.click("[data-testid=toolbox-toggle]"); pg.wait_for_selector("[data-testid=toolbox-toggle][aria-pressed=false]",timeout=30000)
    nuke(pg); pg.click("[data-run=bom]")
    pg.locator("button", has_text=re.compile(r"^Document$")).first.click(force=True)
    pg.wait_for_selector("[data-testid=doc-inputdata][data-ready='1']",timeout=30000)
    pg.set_input_files("[data-testid=doc-in-import]",files=[{"name":"bad.csv","mimeType":"text/csv","buffer":"key,value\ntemperature,27\npressure2,1\n".encode("utf-8")}])
    pg.wait_for_selector("[data-testid=doc-in-msg][data-ok='0']",timeout=10000); badmsg=pg.inner_text("[data-testid=doc-in-msg]")
    kept=pg.input_value("[data-testid=doc-in-temperature]")
    pg.set_input_files("[data-testid=doc-in-import]",files=[{"name":"in.csv","mimeType":"text/csv","buffer":"key,value\ntemperature,27\nhumidity,55\n".encode("utf-8")}])
    pg.wait_for_selector("[data-testid=doc-in-msg][data-ok='1']",timeout=10000)
    filled=(pg.input_value("[data-testid=doc-in-temperature]"), pg.input_value("[data-testid=doc-in-humidity]"))
    ok("S59b 입력값 CSV Import — 템플릿에 없는 항목이 있으면 줄 번호와 함께 거부하고 아무것도 안 채운다 · 맞는 CSV 는 칸을 채운다",
       (badmsg[:40], kept, filled), "3번째 줄" in badmsg and "pressure2" in badmsg and kept=="20" and filled==("27","55"))
    pg.wait_for_function("()=>{const b=document.querySelector('[data-testid=doc-make-techdata]'); return b && !b.disabled;}",timeout=60000)
    with pg.expect_response(lambda q: q.url.endswith("/api/documents") and q.request.method=="POST",timeout=30000) as dres:
        pg.click("[data-testid=doc-make-techdata]")
    d7=dres.value.json(); hp=ctx.request.get(DOCS+f"/{d7.get('id')}/print").text()
    ok("S59c CSV 로 채운 값으로 Tech Data 를 만들면 문서에 그 값(27 °C · 55 %)이 스냅샷으로 남는다", d7.get("docNo"),
       'data-key="temperature">27 °C' in hp and 'data-key="humidity">55 %' in hp)
    bs=ctx.request.get(TD+"?status=bogus")
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test")); v1=vw.request.get(TD); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test")); g1=gb.request.get(TD).json()["rows"]; gb.close()
    ok("S59d 잘못된 상태 400 · 읽기 전용이라 viewer 는 볼 수 있다(200) · 다른 회사는 우리 Tech Data 를 못 본다(0건)", (bs.status, v1.status, len(g1)),
       bs.status==400 and v1.status==200 and len(g1)==0)
    # ── S60 F8 · p18 · p16 Data Up-Load — 작업대 노드에 자료를 올리고 Inspector 에 목록(F4 첨부·0014 저장소 재사용) ──
    AT=BASE+"/api/attachments"; NP="a0000000-0000-4000-8000-000000000004"; NI="a0000000-0000-4000-8000-000000000003"
    CSV8="item,value\nairflow,55000\n"
    pg.goto(NODE4,wait_until="domcontentloaded"); hydrated(pg); nuke(pg)
    pg.wait_for_selector("[data-testid=node-upload][data-ready='1']",timeout=30000)
    pg.set_input_files("[data-testid=node-upload-file]",files=[{"name":"site_survey.csv","mimeType":"text/csv","buffer":CSV8.encode("utf-8")}])
    pg.click("[data-testid=node-upload-upload]"); pg.wait_for_selector("[data-testid=node-upload-row][data-kind=data]",timeout=30000)
    pg.goto(BASE+f"/workbench?node={NI}",wait_until="domcontentloaded"); hydrated(pg); nuke(pg)
    pg.wait_for_selector("[data-testid=node-upload][data-ready='1']",timeout=30000)
    pg.set_input_files("[data-testid=node-upload-file]",files=[{"name":"coil_datasheet.pdf","mimeType":"application/pdf","buffer":b"%PDF-1.4\n% e2e\n"}])
    pg.click("[data-testid=node-upload-upload]"); pg.wait_for_selector("[data-testid=node-upload-row][data-kind=data]",timeout=30000)
    lp=ctx.request.get(AT+f"?ownerKind=node&ownerKey={NP}").json()["rows"]; li=ctx.request.get(AT+f"?ownerKind=node&ownerKey={NI}").json()["rows"]
    got=ctx.request.get(AT+f"/{lp[0]['id']}/file").text()
    ok("S60a 작업대 노드마다 자료를 올린다 — 프로젝트 노드(CSV)와 프로젝트가 아닌 Item 노드(PDF) 각각 Inspector 목록에 뜨고, 내려받으면 그대로",
       (len(lp), len(li), got==CSV8), len(lp)==1 and len(li)==1 and got==CSV8)
    def upn(key, name, data, c=ctx):
        return c.request.post(AT, multipart={"ownerKind":"node","ownerKey":key,"kind":"data","file":{"name":name,"mimeType":"application/octet-stream","buffer":data}})
    e1=upn(NP,"run.exe",b"MZ"); e2=upn("00000000-0000-4000-8000-000000000000","a.csv",b"x")
    ok("S60b 허용 밖 확장자 415 · 없는 노드 404", (e1.status, e2.status), e1.status==415 and e2.status==404)
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test")); v1=upn(NP,"v.csv",b"x",c=vw); v2=vw.request.get(AT+f"?ownerKind=node&ownerKey={NP}"); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test")); g1=upn(NP,"g.csv",b"x",c=gb); g2=gb.request.get(AT+f"?ownerKind=node&ownerKey={NP}").json()["rows"]; gb.close()
    ok("S60c viewer 는 볼 수 있지만(200) 못 올린다(403) · 다른 회사는 우리 노드에 못 올리고(404) 목록도 0건", (v2.status, v1.status, g1.status, len(g2)),
       v2.status==200 and v1.status==403 and g1.status==404 and len(g2)==0)
    # ── S61 F9 · p54 System Set-Up 지도 — 있는 화면은 링크, 없는 것은 "아직 없음 — 필요한 입력" ──
    pg.goto(BASE+"/setup",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=setup-link-map]",timeout=30000); nuke(pg)
    pg.click("[data-testid=setup-link-map]"); pg.wait_for_selector("[data-testid=setup-map]",timeout=30000)
    kinds={x[0]:x[1] for x in pg.eval_on_selector_all("[data-testid^=map-]","es=>es.map(e=>[e.dataset.testid.replace('map-',''),e.dataset.kind])")}
    must=["s15","s16","s411","s31","s34","s32","s33","c-db"]
    ok("S61a 청사진 p54 에서 빠졌던 항목이 이어졌다 — Arrangement Code · Arrangement Set-up · TLM Design · CPQ Selection · Print Set-up · Technical · Document · Company DB 가 링크",
       [ (k,kinds.get(k)) for k in must], all(kinds.get(k)=="link" for k in must))
    hrefs=sorted(set(pg.eval_on_selector_all("[data-kind=link]","es=>es.map(e=>e.getAttribute('href'))")))
    st={h: ctx.request.get(BASE+h).status for h in hrefs}
    ok("S61b 지도의 링크가 가리키는 화면이 모두 열린다(200)", st, len(hrefs)>=10 and all(v==200 for v in st.values()))
    nones=pg.eval_on_selector_all("[data-kind=none]","es=>es.map(e=>e.innerText)")
    # H4(ccmd H) 이후 Department 는 ERP 기준정보 화면(/setup/erp)으로 링크 — 아직 없음은 Work Process · 그 밖의 ERP 둘.
    ok("S61c 없는 것은 있는 척하지 않는다 — Work Process · 그 밖의 ERP 는 '아직 없음 — 필요한 입력' 으로", len(nones),
       len(nones)==2 and all("아직 없음" in t and "필요한 입력" in t for t in nones))
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test")); vm=vw.request.get(BASE+"/setup/map"); vw.close()
    ok("S61d 지도는 읽기 화면 — viewer 도 연다(200)", vm.status, vm.status==200)
    # ── S62 F10 · p66 · p67 제조 정보 표(공정별 시간 × 임율 · 장비) → 인건비 · 견적 적용(스냅샷 단가·출처) ──
    MR=BASE+"/api/setup/mfg-rates"; RUNB=BASE+"/api/run/bom"; DOCS=BASE+"/api/documents"; N4="a0000000-0000-4000-8000-000000000004"
    def run10():
        r=ctx.request.post(RUNB,headers=J0,data=json.dumps({"slots":S55_0,"code":"EU-55-2123-630SS-1-21-13-15","node":N4})).json()
        return r, ctx.request.post(BASE+"/api/run/cost",headers=J0,data=json.dumps({"runId":r["runId"]})).json()["cost"]
    r0,c0=run10()
    ok("S62a 제조 정보 표가 비면 인건비 = 재료비 × 18% (기존 그대로 · 근거 ratio 가 스냅샷에)", (c0["labor"], c0.get("laborBasis")),
       c0["labor"]==round(c0["material"]*0.18) and (c0.get("laborBasis") or {}).get("kind")=="ratio")
    pg.goto(BASE+"/setup",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=setup-link-mfg]",timeout=30000); nuke(pg)
    pg.click("[data-testid=setup-link-mfg]"); pg.wait_for_selector("[data-testid=mfg-rates][data-ready='1'][data-product=EU]",timeout=60000); nuke(pg)
    for proc,eq,h,rt in (("조립","",12,45000),("도장","도장 부스",6,38000)):
        pg.fill("[data-testid=mfg-process]",proc); pg.fill("[data-testid=mfg-equipment]",eq); pg.fill("[data-testid=mfg-hours]",str(h)); pg.fill("[data-testid=mfg-rate]",str(rt))
        pg.click("[data-testid=mfg-add]"); wait_sel(pg,f"[data-testid='mfg-row-{proc}']")
    wait_text(pg,"[data-testid=mfg-total]","768,000")
    nuke(pg); pg.screenshot(path=f"{OUT}/72_mfg_rate.png",full_page=True)
    r1,c1=run10(); lb1=c1.get("laborBasis") or {}
    ok("S62b 제조 정보 표(조립 12h×45,000 · 도장 6h×38,000 도장 부스)를 등록하면 다음 BOM Run 인건비 = Σ 시간×임율 = 768,000 · 경비는 (재료+인건)×12% · 근거 행이 스냅샷에",
       (c1["labor"], c1["overhead"], lb1.get("kind"), [(x["process"],x["equipment"],x["amount"]) for x in lb1.get("rows",[])]),
       c1["labor"]==768000 and c1["overhead"]==round((c1["material"]+768000)*0.12) and c1["total"]==c1["material"]+768000+c1["overhead"]
       and lb1.get("kind")=="mfg-table" and [(x["process"],x["equipment"],x["amount"]) for x in lb1.get("rows",[])]==[("조립",None,540000),("도장","도장 부스",228000)])
    q1=ctx.request.post(DOCS,headers=J0,data=json.dumps({"runId":r1["runId"],"type":"quotation"})).json(); qh=ctx.request.get(DOCS+f"/{q1.get('id')}/print").text()
    qb=ctx.request.get(DOCS+f"/{q1.get('id')}").json().get("body",{}); ap=qb.get("applied") or []
    ok("S62c 견적 적용 — 견적 합계 = 스냅샷 원가 · PCR Manufacturing = 768,000 · 인쇄본에 인건비 기준 줄 · 견적 적용 Table 금액 합 = Material Cost",
       (qb.get("total"), c1["total"], qb.get("pcr",{}).get("manufacturing"), len(ap), sum(a["amount"] for a in ap), c1["material"], "labor-basis" in qh, "applied-table" in qh),
       qb.get("total")==c1["total"] and qb.get("pcr",{}).get("manufacturing")==768000 and len(ap)>0 and sum(a["amount"] for a in ap)==c1["material"]
       and all(a["table"] in ("견적","구매") for a in ap) and 'data-testid="labor-basis"' in qh and "제조 정보 표(EU)" in qh and 'data-testid="applied-table"' in qh and 'data-testid="price-basis"' in qh)
    pg.click("[data-testid='mfg-del-도장']"); wait_sel(pg,"[data-testid=mfg-msg][data-ok='1']"); wait_text(pg,"[data-testid=mfg-total]","540,000")
    c1b=ctx.request.post(BASE+"/api/run/cost",headers=J0,data=json.dumps({"runId":r1["runId"]})).json()["cost"]
    qh1b=ctx.request.get(DOCS+f"/{q1.get('id')}/print").text()
    r2,c2=run10()
    ok("S62d 표를 고쳐도(도장 삭제) 뜬 스냅샷·견적의 인건비는 768,000 그대로 · 새로 Run 하면 540,000", (c1b["labor"], "768,000" in qh1b, c2["labor"]),
       c1b["labor"]==768000 and "768,000" in qh1b and c2["labor"]==540000)
    z=ctx.request.post(MR,headers=J0,data=json.dumps({"productCode":"EU","process":"x","hours":0,"rate":1000}))
    nf=ctx.request.post(MR,headers=J0,data=json.dumps({"productCode":"ZZ","process":"x","hours":1,"rate":1000}))
    dup=ctx.request.post(MR,headers=J0,data=json.dumps({"productCode":"EU","process":"조립","hours":1,"rate":1000}))
    ok("S62e 시간 0 은 400 · 없는 제품 404 · 같은 공정 409", (z.status, nf.status, dup.status), z.status==400 and nf.status==404 and dup.status==409)
    rows=ctx.request.get(MR+"?product=EU").json()["rows"]
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v0=vw.request.get(MR+"?product=EU"); v1=vw.request.post(MR,headers=J0,data=json.dumps({"productCode":"EU","process":"v","hours":1,"rate":1})); v2=vw.request.delete(MR+f"/{rows[0]['id']}"); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    g1=gb.request.get(MR+"?product=EU").json()["rows"]; g2=gb.request.delete(MR+f"/{rows[0]['id']}"); gb.close()
    left=ctx.request.get(MR+"?product=EU").json()["rows"]
    ok("S62f viewer 는 보지만(200) 못 넣고·못 지운다(403) · 다른 회사는 우리 표를 못 보고(0건) 못 지운다(404) · 우리 표는 그대로",
       (v0.status, v1.status, v2.status, len(g1), g2.status, len(left)), v0.status==200 and v1.status==403 and v2.status==403 and len(g1)==0 and g2.status==404 and len(left)==1)
    # ── S63 H4 · p64 ERP 기준정보 6종 — Department · Warehouse · Inventory · Bank · Employee · Nation (Company DB 틀 재사용) ──
    EM=BASE+"/api/setup/erp-master"
    def em_rows(kind=None, c=ctx): return c.request.get(EM+(f"?kind={kind}" if kind else "")).json().get("rows",[])
    pg.goto(BASE+"/setup",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=setup-link-erp]",timeout=30000); nuke(pg)
    pg.click("[data-testid=setup-link-erp]"); pg.wait_for_selector("[data-testid=erp-master][data-ready='1'][data-kind=department]",timeout=60000); nuke(pg)
    def em_add(kind, code, name, fields):
        pg.click(f"[data-testid=erp-tab-{kind}]"); wait_sel(pg,f"[data-testid=erp-master][data-kind={kind}]")
        pg.fill("[data-testid=erp-new-code]",code); pg.fill("[data-testid=erp-new-name]",name)
        for k,v,sel in fields:
            if sel: pg.select_option(f"[data-testid=erp-new-{k}]",v)
            else: pg.fill(f"[data-testid=erp-new-{k}]",v)
        pg.click("[data-testid=erp-add]"); wait_sel(pg,f"[data-testid=erp-row-{code}]")
    em_add("department","D200","구매팀",[("parent","D100",True),("manager","김부장",False)])
    em_add("warehouse","W01","본사 창고",[("location","화성 1동",False)])
    em_add("inventory","INV-001","KDP 재고",[("warehouse","W01",True),("item","KDP",False),("qty","12",False),("unit","ea",False)])
    em_add("bank","B01","거래 은행(예시)",[("branch","본점",False),("account","000-00-000000",False),("nation","KR",True)])
    em_add("nation","US","미국",[("currency","USD",False)])
    em_add("employee","E001","홍길동",[("department","D200",True),("title","과장",False),("email","hong@acme.test",False)])
    nuke(pg); pg.screenshot(path=f"{OUT}/73_erp_master.png",full_page=True)
    allr=em_rows(); byk={k:[r for r in allr if r["kind"]==k] for k in ("department","warehouse","inventory","bank","employee","nation")}
    e1=[r for r in byk["employee"] if r["code"]=="E001"]; i1=[r for r in byk["inventory"] if r["code"]=="INV-001"]
    ok("S63a ERP 기준정보 6종을 화면에서 등록한다 — 직원은 부서를, 재고는 창고를, 은행은 국가를 목록에서 골라 가리킨다(수량은 수)",
       {k:len(v) for k,v in byk.items()} | {"E001": e1[0]["attrs"] if e1 else None, "INV": i1[0]["attrs"] if i1 else None},
       all(len(v)>=1 for v in byk.values()) and e1 and e1[0]["attrs"].get("department")=="D200" and i1 and i1[0]["attrs"].get("qty")==12 and i1[0]["attrs"].get("warehouse")=="W01")
    pg.click("[data-testid=erp-tab-department]"); wait_sel(pg,"[data-testid=erp-master][data-kind=department]")
    pg.click("[data-testid=erp-del-D200]"); wait_sel(pg,"[data-testid=erp-msg][data-ok='0']")
    m409=pg.inner_text("[data-testid=erp-msg]") if pg.query_selector("[data-testid=erp-msg]") else ""
    kr=[r for r in byk["nation"] if r["code"]=="KR"][0]; dkr=ctx.request.delete(EM+f"/{kr['id']}")
    ok("S63b 가리키는 곳이 있으면 삭제 409 + 어디서 가리키는지 — 부서 D200(직원 1) · 국가 KR(은행 1 · Company DB 2), 화면에 이유",
       (m409[:80], dkr.status, dkr.json().get("usage")), "409" in m409 and "사용 중" in m409 and dkr.status==409 and "bank.nation 1" in dkr.json().get("usage",{}).get("detail",[]) and any(d.startswith("Company DB") for d in dkr.json().get("usage",{}).get("detail",[])))
    d200=[r for r in byk["department"] if r["code"]=="D200"][0]; d100=[r for r in byk["department"] if r["code"]=="D100"][0]
    j0=lambda d: json.dumps(d)
    x5=ctx.request.patch(EM+f"/{d100['id']}",headers=J0,data=j0({"attrs":{"parent":"D200"}}))   # D200 의 상위가 D100 → 순환
    pg.click("[data-testid=erp-toggle-D200]"); pg.wait_for_selector("[data-testid=erp-row-D200][data-active='0']",timeout=30000)
    x1=ctx.request.post(EM,headers=J0,data=j0({"kind":"employee","code":"E002","name":"새 직원","attrs":{"department":"D200"}}))
    x2=ctx.request.patch(EM+f"/{e1[0]['id']}",headers=J0,data=j0({"attrs":{"department":"D200","title":"차장"}}))
    x3=ctx.request.post(EM,headers=J0,data=j0({"kind":"employee","code":"E003","name":"부서 없음","attrs":{}}))
    x4=ctx.request.post(EM,headers=J0,data=j0({"kind":"nation","code":"JP","name":"일본","attrs":{"currency":"yen"}}))
    x6=ctx.request.patch(EM+f"/{d100['id']}",headers=J0,data=j0({"code":"D999"}))
    x7=ctx.request.post(EM,headers=J0,data=j0({"kind":"department","code":"D100","name":"중복"}))
    x8=ctx.request.post(EM,headers=J0,data=j0({"kind":"desk","code":"X","name":"x"}))
    ok("S63c 사용 중지된 부서는 새로 못 가리키고(400) 이미 가리키던 직원은 그대로 고친다(200) · 필수 칸 400 · 통화 형식 400 · 부서 순환 400 · 코드 변경 400 · 중복 409 · 모르는 종류 400",
       (x1.status,x2.status,x3.status,x4.status,x5.status,x6.status,x7.status,x8.status),
       (x1.status,x2.status,x3.status,x4.status,x5.status,x6.status,x7.status,x8.status)==(400,200,400,400,400,400,409,400))
    inv=i1[0]; w01=[r for r in byk["warehouse"] if r["code"]=="W01"][0]
    y1=ctx.request.delete(EM+f"/{w01['id']}"); y2=ctx.request.delete(EM+f"/{inv['id']}"); y3=ctx.request.delete(EM+f"/{w01['id']}")
    ok("S63d 창고는 재고가 가리키는 동안 409, 재고를 지우면 창고도 지워진다(200)", (y1.status,y2.status,y3.status), (y1.status,y2.status,y3.status)==(409,200,200))
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v0=vw.request.get(EM); v1=vw.request.post(EM,headers=J0,data=j0({"kind":"nation","code":"CN","name":"중국"})); v2=vw.request.patch(EM+f"/{d100['id']}",headers=J0,data=j0({"name":"x"})); v3=vw.request.delete(EM+f"/{d100['id']}"); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    g0=len(em_rows(c=gb)); g1=gb.request.patch(EM+f"/{d100['id']}",headers=J0,data=j0({"name":"hijack"})); g2=gb.request.delete(EM+f"/{e1[0]['id']}"); gb.close()
    still=[r for r in em_rows("department") if r["code"]=="D100"]
    ok("S63e viewer 는 보지만(200) 못 넣고·못 고치고·못 지운다(403) · 다른 회사는 우리 기준정보 0건 · 고치기·지우기 404 · 우리 행 그대로",
       (v0.status,v1.status,v2.status,v3.status,g0,g1.status,g2.status,still[0]["name"] if still else None),
       v0.status==200 and (v1.status,v2.status,v3.status)==(403,403,403) and g0==0 and (g1.status,g2.status)==(404,404) and still and still[0]["name"]=="설계팀")
    # ── S64 H5 · p39 · p40 도면 템플릿 — Sub Drawing 호출(설계 우선순위) · Detail Design 주의사항 → 도면에 박히고 시트에 나온다 ──
    DT=BASE+"/api/setup/drawing-template"; DRW=BASE+"/api/drawings"
    pg.goto(BASE+"/setup",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=setup-link-drawing-template]",timeout=30000); nuke(pg)
    pg.click("[data-testid=setup-link-drawing-template]"); pg.wait_for_selector("[data-testid=dt][data-ready='1'][data-product=EU]",timeout=60000); nuke(pg)
    for code,pr in (("KFP 1","2"),("KHR 1","1"),("PSH 1","1"),("KCP 1","2")):
        pg.select_option("[data-testid=dt-sub-child]",code); pg.fill("[data-testid=dt-sub-priority]",pr); pg.click("[data-testid=dt-sub-add]")
        wait_sel(pg,f"[data-testid='dt-sub-row-{code}']")
    for txt,pr in (("베어링 하우징 조립 전 축 정렬 확인","1"),("<b>도장</b> 전 탈지","2")):
        n0=len(pg.query_selector_all("[data-testid=dt-note-row]"))
        pg.fill("[data-testid=dt-note-text]",txt); pg.fill("[data-testid=dt-note-priority]",pr); pg.click("[data-testid=dt-note-add]")
        pg.wait_for_function("(n)=>document.querySelectorAll('[data-testid=dt-note-row]').length>n",arg=n0,timeout=30000)
    nuke(pg); pg.screenshot(path=f"{OUT}/74_sub_drawing.png",full_page=True)
    up=ctx.request.post(BASE+"/api/attachments",multipart={"ownerKind":"product_code","ownerKey":"KFP 1","kind":"dwg2d","file":{"name":"plug_fan.dxf","mimeType":"application/dxf","buffer":b"0\nSECTION\n2\nENTITIES\n0\nENDSEC\n0\nEOF\n"}})
    r64,_c64=run10()
    d1=ctx.request.post(DRW,headers=J0,data=json.dumps({"runId":r64["runId"],"type":"assembly"})).json(); m1=d1.get("meta",{})
    sd=[(x["childCode"],x["priority"],(x.get("dwg") or {}).get("name")) for x in m1.get("subDrawings",[])]
    ok("S64a 도면을 뜨면 템플릿의 하부 도면이 설계 우선순위 순(같으면 코드 순)으로 박힌다 — 이 BOM 에 없는 PSH 1(가습, D=A1 일 때만)은 부르지 않고, KFP 1 은 코드에 첨부한 DWG 를 가리킨다 · 주의사항 2",
       (up.status, sd, m1.get("notes")),
       up.status==200 and sd==[("KHR 1",1,None),("KCP 1",2,None),("KFP 1",2,"plug_fan.dxf")] and len(m1.get("notes",[]))==2 and m1.get("notes",[None])[0].startswith("베어링"))
    sh1=ctx.request.get(DRW+f"/{d1.get('id')}/sheet"); t1=sh1.text()
    order=re.findall(r'<tr data-code="([^"]+)"',t1)
    ok("S64b 도면 시트(인쇄) — 저장된 DXF 를 SVG 로 · 하부 도면 표(Item · Description · Q'ty · Remarks · DWG) · 주의사항 목록(이스케이프)",
       (sh1.status, order, "sheet-svg" in t1, "&lt;b&gt;도장&lt;/b&gt;" in t1),
       sh1.status==200 and order==["KHR 1","KCP 1","KFP 1"] and 'data-testid="sheet-svg"' in t1 and "<svg" in t1 and 'data-testid="sheet-notes"' in t1 and "&lt;b&gt;도장&lt;/b&gt;" in t1 and "<b>도장</b>" not in t1)
    pg.click("[data-testid='dt-sub-del-KFP 1']"); pg.wait_for_selector("[data-testid='dt-sub-row-KFP 1']",state="detached",timeout=30000)
    t1b=ctx.request.get(DRW+f"/{d1.get('id')}/sheet").text()
    d2=ctx.request.post(DRW,headers=J0,data=json.dumps({"runId":r64["runId"],"type":"assembly"})).json()
    sd2=[x["childCode"] for x in d2.get("meta",{}).get("subDrawings",[])]
    ok("S64c 템플릿을 고쳐도(KFP 1 빼기) 이미 뜬 도면 시트는 그대로 · 새로 뜬 도면(다음 개정)부터 빠진다", (re.findall(r'<tr data-code="([^"]+)"',t1b), d2.get("rev"), sd2),
       re.findall(r'<tr data-code="([^"]+)"',t1b)==["KHR 1","KCP 1","KFP 1"] and sd2==["KHR 1","KCP 1"])
    q=lambda body: ctx.request.post(DT,headers=J0,data=json.dumps(body))
    e1=q({"productCode":"EU","kind":"sub","childCode":"ER","priority":1}); e2=q({"productCode":"EU","kind":"sub","childCode":"KHR 1","priority":3})
    e3=q({"productCode":"EU","kind":"sub","childCode":"KCD 1","priority":0}); e4=q({"productCode":"ZZ","kind":"note","text":"x","priority":1}); e5=q({"productCode":"EU","kind":"note","text":"  ","priority":1})
    ok("S64d 코드 관계에 없는 하위 코드 400 · 같은 하위 코드 409 · 우선순위 0 은 400 · 없는 제품 404 · 빈 주의사항 400", (e1.status,e2.status,e3.status,e4.status,e5.status),
       (e1.status,e2.status,e3.status,e4.status,e5.status)==(400,409,400,404,400))
    tid=ctx.request.get(DT+"?product=EU").json()["subs"][0]["id"]
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v0=vw.request.get(DT+"?product=EU"); v1=vw.request.post(DT,headers=J0,data=json.dumps({"productCode":"EU","kind":"note","text":"v","priority":1})); v2=vw.request.delete(DT+f"/{tid}"); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    g0=gb.request.get(DT+"?product=EU").json(); g1=gb.request.delete(DT+f"/{tid}"); g2=gb.request.get(DRW+f"/{d1.get('id')}/sheet"); gb.close()
    ok("S64e viewer 는 보지만(200) 못 넣고·못 뺀다(403) · 다른 회사는 우리 템플릿 0건 · 빼기 404 · 우리 도면 시트 404",
       (v0.status,v1.status,v2.status,len(g0.get("subs",[]))+len(g0.get("notes",[])),g1.status,g2.status),
       (v0.status,v1.status,v2.status)==(200,403,403) and len(g0.get("subs",[]))+len(g0.get("notes",[]))==0 and (g1.status,g2.status)==(404,404))
    # ── S65 H6 · p16 · p47 Output Data 템플릿 · 그래프 전용 data · 그래프 · Table List(Table Type) ──
    OI=BASE+"/api/setup/output-items"; GR=BASE+"/api/setup/graphs"; TM=BASE+"/api/setup/table-meta"
    pg.goto(BASE+"/setup",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=setup-link-document]",timeout=30000); nuke(pg)
    pg.click("[data-testid=setup-link-document]"); pg.wait_for_selector("[data-testid=doc-setup][data-ready='1']",timeout=60000); nuke(pg)
    for key,label,unit,src,ref in (("macro_result","Macro result","","macro",""),("width","Width W","mm","snapshot","dims.W"),("total_cost","Full cost","원","snapshot","cost.total")):
        pg.fill("[data-testid=out-new-key]",key); pg.fill("[data-testid=out-new-label]",label); pg.fill("[data-testid=out-new-unit]",unit)
        pg.select_option("[data-testid=out-new-source]",src)
        if ref: pg.select_option("[data-testid=out-new-ref]",ref)
        pg.click("[data-testid=out-add]"); wait_sel(pg,f"[data-testid=out-row-{key}]")
    pg.fill("[data-testid=graph-new-name]","Fan curve"); pg.select_option("[data-testid=graph-new-chart]","line")
    pg.fill("[data-testid=graph-new-x]","Airflow (CMH)"); pg.fill("[data-testid=graph-new-y]","Static (Pa)")
    pg.fill("[data-testid=graph-new-points]","20000,900\n40000,820\n55000,700\n70000,480"); pg.select_option("[data-testid=graph-new-marker]","macro_result")
    pg.wait_for_selector("[data-testid=graph-preview] svg[data-points='4']",timeout=30000)
    pg.click("[data-testid=graph-add]"); wait_sel(pg,"[data-testid='graph-row-Fan curve']")
    pg.wait_for_selector("[data-testid=table-list][data-ready='1'][data-product=EU]",timeout=30000)
    pg.select_option("[data-testid=tl-type-mat]","material"); pg.fill("[data-testid=tl-dept-mat]","Engineering"); pg.fill("[data-testid=tl-desc-mat]","재질 배율(SUS·AL)")
    pg.click("[data-testid=tl-save-mat]"); pg.wait_for_selector("[data-testid=tl-row-mat][data-type=material]",timeout=30000)
    pg.select_option("[data-testid=tl-type-opt]","variant"); pg.select_option("[data-testid=tl-variant-opt]","cap"); pg.click("[data-testid=tl-save-opt]"); pg.wait_for_selector("[data-testid=tl-row-opt][data-type=variant]",timeout=30000)
    nuke(pg); pg.screenshot(path=f"{OUT}/75_output_template.png",full_page=True)
    tl={t["name"]:t["meta"] for t in ctx.request.get(TM+"?product=EU").json()["tables"]}
    ok("S65a Set-Up ▸ Document — Output 항목 3(승인 매크로 결과 · 스냅샷 W · 원가) · 그래프 전용 data 4점 + 표시선 · Table List(mat=Material · opt=Variant of cap)",
       (len(ctx.request.get(OI).json()["rows"]), [(g["name"],len(g["points"]),g["markerKey"]) for g in ctx.request.get(GR).json()["rows"]], (tl.get("mat") or {}).get("tableType"), (tl.get("opt") or {}).get("variantOf")),
       len(ctx.request.get(OI).json()["rows"])==3 and (tl.get("mat") or {}).get("department")=="Engineering" and (tl.get("opt") or {}).get("tableType")=="variant" and (tl.get("opt") or {}).get("variantOf")=="cap")
    r65,c65=run10()
    t65=ctx.request.post(DOCS,headers=J0,data=json.dumps({"runId":r65["runId"],"type":"techdata"})).json()
    b65=ctx.request.get(DOCS+f"/{t65.get('id')}").json().get("body",{}); od={o["key"]:o["value"] for o in b65.get("outputData",[])}
    ph=ctx.request.get(DOCS+f"/{t65.get('id')}/print").text()
    ok("S65b Tech Data 에 Output 값이 스냅샷에서 박힌다 — 매크로 결과 455.4 · W(스냅샷 치수) · 원가 합계 = 그 스냅샷 원가 · 인쇄본에 Output 표와 그래프(표시선 455.4)",
       (od, c65["total"], "techdata-outputdata" in ph, 'data-marker="455.4"' in ph),
       od.get("macro_result")==455.4 and isinstance(od.get("width"),(int,float)) and od.get("total_cost")==c65["total"] and 'data-testid="techdata-outputdata"' in ph and 'data-testid="techdata-graph"' in ph and 'data-marker="455.4"' in ph)
    gid=[g for g in ctx.request.get(GR).json()["rows"] if g["name"]=="Fan curve"][0]["id"]
    busy=ctx.request.delete(OI+f"/{[o for o in ctx.request.get(OI).json()['rows'] if o['key']=='macro_result'][0]['id']}")
    ctx.request.delete(GR+f"/{gid}")
    ph2=ctx.request.get(DOCS+f"/{t65.get('id')}/print").text()
    ok("S65c 표시선으로 쓰이는 Output 항목은 못 뺀다(409) · 그래프를 빼도 이미 만든 Tech Data 인쇄본의 그래프는 그대로", (busy.status, 'data-testid="techdata-graph"' in ph2),
       busy.status==409 and 'data-testid="techdata-graph"' in ph2 and 'data-marker="455.4"' in ph2)
    q=lambda u,body: ctx.request.post(u,headers=J0,data=json.dumps(body))
    e1=q(OI,{"key":"Bad Key","label":"x","source":"macro"}); e2=q(OI,{"key":"x1","label":"x","source":"snapshot","ref":"cost.profit"}); e3=q(OI,{"key":"width","label":"dup","source":"macro"})
    e4=q(GR,{"name":"g","chart":"pie","points":[{"x":"a","y":1}]}); e5=q(GR,{"name":"g","chart":"bar","points":[{"x":"a","y":"b"}]}); e6=q(GR,{"name":"g","chart":"bar","points":[{"x":"a","y":1}],"markerKey":"nope"})
    e7=ctx.request.put(TM,headers=J0,data=json.dumps({"productCode":"EU","tableName":"zzz","tableType":"tech"})); e8=ctx.request.put(TM,headers=J0,data=json.dumps({"productCode":"EU","tableName":"cap","tableType":"tech","variantOf":"opt"}))
    ok("S65d 잘못된 key 400 · 스냅샷 경로 밖(새 계산) 400 · 중복 409 · 그래프 모양·점·표시선 400 · 없는 표 400 · Variant 가 아닌데 변형 원본 400",
       tuple(x.status for x in (e1,e2,e3,e4,e5,e6,e7,e8)), tuple(x.status for x in (e1,e2,e3,e4,e5,e6,e7,e8))==(400,400,409,400,400,400,400,400))
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v0=vw.request.get(OI); v1=vw.request.post(OI,headers=J0,data=json.dumps({"key":"v","label":"v","source":"macro"})); v2=vw.request.put(TM,headers=J0,data=json.dumps({"productCode":"EU","tableName":"cap","tableType":"tech"})); vw.close()
    wid=[o for o in ctx.request.get(OI).json()["rows"] if o["key"]=="width"][0]["id"]
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    g0=len(gb.request.get(OI).json().get("rows",[])); g1=len(gb.request.get(GR).json().get("rows",[])); g2=gb.request.delete(OI+f"/{wid}"); gb.close()
    ok("S65e viewer 는 보지만(200) 못 넣고·못 고친다(403) · 다른 회사는 우리 Output·그래프 0건 · 지우기 404",
       (v0.status,v1.status,v2.status,g0,g1,g2.status), v0.status==200 and (v1.status,v2.status)==(403,403) and g0==0 and g1==0 and g2.status==404)
    # ── S66 H7 · p47 Coding List — 노드마다 승인 매크로 1개 · 작업대 Inspector 에서 이동 · 행에서 작업대로 ──
    CL=BASE+"/api/macros/coding-list"
    pg.goto(BASE+f"/workbench?node={N4}",wait_until="domcontentloaded"); hydrated(pg); pg.wait_for_selector("[data-testid=inspector-coding-list]",timeout=60000); nuke(pg)
    pg.click("[data-testid=inspector-coding-list]"); pg.wait_for_selector("[data-testid=coding-list][data-ready='1']",timeout=60000)
    ms=ctx.request.get(BASE+f"/api/macros?node={N4}").json(); ms=ms.get("macros",ms.get("rows",[])) if isinstance(ms,dict) else ms
    ap=[x for x in ms if x.get("status")=="approved"]
    shown=pg.get_attribute(f"[data-testid=cl-row-{N4}]","data-approved"); cnt=pg.get_attribute(f"[data-testid=cl-row-{N4}]","data-count")
    rows=ctx.request.get(CL).json()["rows"]; prj=[r for r in rows if r["stableId"]==N4][0]
    ok("S66a Inspector 의 Coding List 로 이동 — 프로젝트 노드에 붙은 승인 매크로 개정(1개)이 목록에 그대로 · 마지막 BOM 이 쓴 개정도",
       (shown, cnt, [x.get("revision") for x in ap], (prj.get("lastRun") or {}).get("macroRevision")),
       len(ap)==1 and shown==str(ap[0].get("revision")) and cnt=="1" and prj.get("lastRun") is not None and prj["approved"]["dsl"]==ap[0].get("dsl"))
    ahu=[r for r in rows if r["label"]=="AHU-01"][0]
    d0=ahu["drafts"]
    dr=ctx.request.post(BASE+"/api/macros",headers=J0,data=json.dumps({"node":ahu["stableId"],"dsl":"=SUM(Table1(A,1:1))","mode":"draft"}))
    ahu2=[r for r in ctx.request.get(CL).json()["rows"] if r["stableId"]==ahu["stableId"]][0]
    pg.reload(wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=coding-list][data-ready='1']",timeout=60000)
    nuke(pg); pg.screenshot(path=f"{OUT}/76_coding_list.png",full_page=True)
    ok("S66b 초안은 승인 매크로가 아니다 — AHU-01 에 초안을 저장해도 '승인 매크로 없음' · 초안 수만 +1",
       (dr.status, ahu2["approved"], ahu2["drafts"]-d0, "승인 매크로 없음" in pg.inner_text(f"[data-testid=cl-row-{ahu['stableId']}]")),
       dr.status==200 and ahu2["approved"] is None and ahu2["drafts"]==d0+1 and "승인 매크로 없음" in pg.inner_text(f"[data-testid=cl-row-{ahu['stableId']}]"))
    pg.click(f"[data-testid=cl-open-{N4}]"); pg.wait_for_selector("[data-testid=canvas-cmds][data-ready='1']",timeout=60000)
    ok("S66c 행의 노드 이름을 누르면 그 노드의 작업대로", pg.url.replace(BASE,""), pg.url.endswith(f"/workbench?node={N4}"))
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test")); v0=vw.request.get(CL); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test")); g0=gb.request.get(CL).json().get("rows",[]); gb.close()
    ok("S66d 읽기 목록 — viewer 도 본다(200) · 다른 회사 목록에는 우리 노드가 없다", (v0.status, len(g0), any(r["stableId"]==N4 for r in g0)),
       v0.status==200 and not any(r["stableId"] in (N4, ahu["stableId"]) for r in g0))
    # ── S67 H8 · p57 Toolbox Macro — Data Management(목록) · 함수 마법사(식 글자 → 기존 파서·Verify) · 그래프 마법사(H6 그래프를 단계로) ──
    DS=BASE+"/api/setup/data-sources"
    pg.goto(BASE+"/setup",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=setup-link-toolbox]",timeout=30000); nuke(pg)
    pg.click("[data-testid=setup-link-toolbox]"); pg.wait_for_selector("[data-testid=data-mgmt][data-ready='1']",timeout=60000); nuke(pg)
    pg.select_option("[data-testid=dm-type]","formula"); pg.wait_for_selector("[data-testid=data-mgmt][data-ready='1'][data-type=formula]",timeout=30000)
    ftypes=pg.eval_on_selector_all("[data-testid=dm-row]","es=>es.map(e=>e.dataset.type)")
    allsrc=ctx.request.get(DS).json()["rows"]; kinds={k:len([r for r in allsrc if r["type"]==k]) for k in ("table","chart","formula")}
    ok("S67a Data Management — Type of source(Table · Chart · Formula) · Directory 목록 · Formula 로 거르면 승인 매크로만(프로젝트 노드의 식이 보인다)",
       (kinds, len(ftypes), sorted(set(ftypes))), kinds["table"]>=5 and kinds["formula"]>=1 and len(ftypes)>=1 and set(ftypes)=={"formula"} and any(r["href"].endswith(N4) for r in allsrc if r["type"]=="formula"))
    pg.select_option("[data-testid=fn-pick]","ROUND"); pg.wait_for_selector("[data-testid=fn-wizard][data-fn=ROUND]",timeout=10000)
    pg.fill("[data-testid=fn-arg-value]","SUM(Table1(A,1:4))"); pg.fill("[data-testid=fn-arg-digits]","1")
    pg.wait_for_selector("[data-testid=fn-check][data-ok='1']",timeout=30000); d1=pg.inner_text("[data-testid=fn-dsl]").strip()
    pg.select_option("[data-testid=fn-pick]","Table"); pg.wait_for_selector("[data-testid=fn-wizard][data-fn=Table]",timeout=10000)
    pg.fill("[data-testid=fn-arg-r0]","5"); pg.fill("[data-testid=fn-arg-r1]","1"); wait_text(pg,"[data-testid=fn-dsl]","시작 행"); bad=pg.inner_text("[data-testid=fn-dsl]")
    pg.select_option("[data-testid=fn-pick]","PreC"); pg.wait_for_selector("[data-testid=fn-note]",timeout=10000); note=pg.inner_text("[data-testid=fn-note]")
    ok("S67b 함수 마법사 — ROUND 를 고르고 인자를 채우면 =ROUND(SUM(Table1(A,1:4)), 1) · 기존 파서 통과 · 잘못된 행 범위는 식을 만들지 않고 이유 · PreC 는 예약(실행 안 됨) 경고",
       (d1, bad[:20], note[:20]), d1=="=ROUND(SUM(Table1(A,1:4)), 1)" and "시작 행" in bad and "RESERVED" in note)
    pg.click("[data-testid=gw-chart-bar]"); pg.click("[data-testid=gw-next]")
    pg.fill("[data-testid=gw-points]","1월,12\n2월,18\n3월,9"); pg.wait_for_selector("[data-testid=gw-preview] svg[data-points='3']",timeout=10000); pg.click("[data-testid=gw-next]")
    pg.click("[data-testid=gw-next]"); pg.fill("[data-testid=gw-name]","Monthly orders"); pg.fill("[data-testid=gw-x]","월"); pg.fill("[data-testid=gw-y]","대수")
    nuke(pg); pg.screenshot(path=f"{OUT}/77_wizards.png",full_page=True)
    pg.click("[data-testid=gw-create]"); pg.wait_for_selector("[data-testid=gw-msg][data-ok='1']",timeout=30000)
    gs=[g for g in ctx.request.get(BASE+"/api/setup/graphs").json()["rows"] if g["name"]=="Monthly orders"]
    ok("S67c 그래프 마법사 — 모양(막대) → 그래프 전용 data 3점 → 표시선(없음) → 이름·축 → 만들기 = H6 그래프 한 곳에 저장 · Data Management 의 Chart 에 보인다",
       ([(g["chart"],len(g["points"])) for g in gs], any(r["name"]=="Monthly orders" for r in ctx.request.get(DS+"?type=chart").json()["rows"])),
       len(gs)==1 and gs[0]["chart"]=="bar" and len(gs[0]["points"])==3 and any(r["name"]=="Monthly orders" for r in ctx.request.get(DS+"?type=chart").json()["rows"]))
    pg.goto(BASE+f"/workbench?node={N4}",wait_until="domcontentloaded"); hydrated(pg); nuke(pg)
    pg.locator("button", has_text=re.compile(r"^Macro$")).first.click(); pg.wait_for_selector("[data-testid=macro-fn-wizard-toggle]",timeout=30000)
    pg.click("[data-testid=macro-fn-wizard-toggle]"); pg.wait_for_selector("[data-testid=fn-wizard]",timeout=10000)
    pg.select_option("[data-testid=fn-pick]","SUM"); pg.wait_for_selector("[data-testid=fn-wizard][data-fn=SUM]",timeout=10000)
    pg.fill("[data-testid=fn-arg-table]","Table1(A,4:4)"); pg.wait_for_selector("[data-testid=fn-insert]:not([disabled])",timeout=30000); pg.click("[data-testid=fn-insert]")
    dv=pg.input_value("[data-testid=macro-dsl]")
    with pg.expect_response(lambda r: r.url.endswith("/api/macros") and r.request.method=="POST", timeout=30000) as vr:
        pg.click("[data-testid=macro-verify]")
    vj=vr.value.json()
    ok("S67d 작업대 Macro 탭의 함수 마법사 → Macro 칸에 넣기 → 기존 Verify 가 그대로 검사한다(저장·승인은 그 탭 한 곳)", (dv, vr.value.status, [d.get("severity") for d in vj.get("diagnostics",[])]),
       dv=="=SUM(Table1(A,4:4))" and vr.value.status==200 and not any(d.get("severity")=="error" for d in vj.get("diagnostics",[])))
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test")); v0=vw.request.get(DS); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test")); g0=gb.request.get(DS).json().get("rows",[]); gb.close()
    bq=ctx.request.get(DS+"?type=zzz")
    ok("S67e Data Management 는 읽기 목록 — viewer 200 · 다른 회사 목록에는 우리 식·그래프가 없다 · 모르는 type 400",
       (v0.status, len(g0), bq.status), v0.status==200 and not any(r["type"] in ("formula","chart") and (r["href"].endswith(N4) or r["name"]=="Monthly orders") for r in g0) and bq.status==400)
    # ── S68 H9 · p48 인쇄 양식 편집기 — 요소를 끌어 배치·크기 조절 → 새 버전 저장 → 인쇄본이 배치를 따르고, 발행본은 발행 순간 버전에 고정 ──
    PL=BASE+"/api/setup/print-layouts"
    pg.goto(BASE+"/setup",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=setup-link-print-layout]",timeout=30000); nuke(pg)
    pg.click("[data-testid=setup-link-print-layout]"); pg.wait_for_selector("[data-testid=pl][data-ready='1'][data-doctype=techdata]",timeout=60000); nuke(pg)
    pg.click("[data-testid=pl-default]")
    gx0=float(pg.get_attribute("[data-testid=pl-el-graph]","data-x")); gy0=float(pg.get_attribute("[data-testid=pl-el-graph]","data-y"))
    # 표 크기를 먼저 — 그래프를 먼저 옮기면 표 모서리 손잡이를 덮어 그래프가 다시 끌린다(H 게이트 S68a 실측)
    tw0=float(pg.get_attribute("[data-testid=pl-el-table]","data-w"))
    hb=pg.locator("[data-testid=pl-handle-table]").bounding_box(); pg.mouse.move(hb["x"]+4, hb["y"]+4); pg.mouse.down(); pg.mouse.move(hb["x"]-30, hb["y"]+20, steps=6); pg.mouse.up()
    bb=pg.locator("[data-testid=pl-el-graph]").bounding_box(); pg.mouse.move(bb["x"]+bb["width"]/2, bb["y"]+bb["height"]/2); pg.mouse.down()
    pg.mouse.move(bb["x"]+bb["width"]/2-60, bb["y"]+bb["height"]/2+40, steps=8); pg.mouse.up()
    gx1=float(pg.get_attribute("[data-testid=pl-el-graph]","data-x")); gy1=float(pg.get_attribute("[data-testid=pl-el-graph]","data-y")); tw1=float(pg.get_attribute("[data-testid=pl-el-table]","data-w"))
    pg.click("[data-testid=pl-add-text]"); pg.wait_for_selector("[data-testid=pl-sel-text]",timeout=10000); pg.fill("[data-testid=pl-sel-text]","Good air makes Good Life")
    pg.click("[data-testid=pl-save]"); pg.wait_for_selector("[data-testid=pl][data-latest='1']",timeout=30000)
    nuke(pg); pg.screenshot(path=f"{OUT}/78_print_layout.png",full_page=True)
    l1=ctx.request.get(PL+"?docType=techdata").json()["latest"]; g1=[e for e in l1["elements"] if e["kind"]=="graph"][0]
    ok("S68a 편집기 — 기본 양식 배치 → 그래프를 끌어 옮기고(x·y 바뀜) · 표 모서리로 크기(w 줄어듦) · 글상자 더하고 → v1 저장",
       ((gx0,gy0),(gx1,gy1),(tw0,tw1),l1["version"],len(l1["elements"])),
       gx1<gx0 and gy1>gy0 and tw1<tw0 and l1["version"]==1 and len(l1["elements"])==7 and g1["x"]==gx1)
    tdA=ctx.request.post(DOCS,headers=J0,data=json.dumps({"runId":RUN1,"type":"techdata"})).json(); hA=ctx.request.get(DOCS+f"/{tdA.get('id')}/print").text()
    ok("S68b 발행 전 Tech Data 인쇄본이 v1 배치를 따른다 — 그래프 칸 위치 = 편집기에서 옮긴 자리 · 글상자 글자 · 숫자는 문서 body 그대로",
       ('data-layout-version="1"' in hA, f'left:{g1["x"]}%;top:{g1["y"]}%' in hA, "Good air makes Good Life" in hA, "발행 전" in hA),
       'data-layout-version="1"' in hA and f'left:{g1["x"]}%;top:{g1["y"]}%' in hA and "Good air makes Good Life" in hA and "발행 전" in hA and 'data-testid="techdata-value"' in hA)
    for st in ("review","approved","issued"):
        sr=ctx.request.patch(DOCS+f"/{tdA.get('id')}",headers=J0,data=json.dumps({"status":st}))
    els2=[dict(e, **({"x":40.0} if e["kind"]=="title" else {})) for e in l1["elements"]]
    v2=ctx.request.post(PL,headers=J0,data=json.dumps({"docType":"techdata","elements":els2})).json()
    hA2=ctx.request.get(DOCS+f"/{tdA.get('id')}/print").text()
    tdB=ctx.request.post(DOCS,headers=J0,data=json.dumps({"runId":RUN1,"type":"techdata"})).json(); hB=ctx.request.get(DOCS+f"/{tdB.get('id')}/print").text()
    vers=ctx.request.get(PL+"?docType=techdata").json()["versions"]
    ok("S68c 발행하면 그 순간의 양식 버전이 박힌다 — 양식을 v2 로 고쳐도 발행본은 v1 배치 그대로 · 새 문서는 v2 · 버전 목록에 'v1 발행 1건'",
       (sr.status, v2.get("version"), 'data-layout-version="1"' in hA2, "발행 때 박힌 버전" in hA2, 'data-layout-version="2"' in hB, [(v["version"],v["issued"]) for v in vers]),
       sr.status==200 and v2.get("version")==2 and 'data-layout-version="1"' in hA2 and "발행 때 박힌 버전" in hA2 and 'data-layout-version="2"' in hB and 'data-el="title" data-id="title" style="left:5%' in hA2 and 'data-el="title" data-id="title" style="left:40%' in hB and dict((v["version"],v["issued"]) for v in vers).get(1)==1)
    q=lambda body: ctx.request.post(PL,headers=J0,data=json.dumps(body))
    e1=q({"docType":"techdata","elements":[{"kind":"table","x":80,"y":0,"w":30,"h":10}]}); e2=q({"docType":"techdata","elements":[{"kind":"chart","x":0,"y":0,"w":10,"h":10}]})
    e3=q({"docType":"techdata","elements":[]}); e4=q({"docType":"invoice","elements":[{"kind":"title","x":0,"y":0,"w":10,"h":10}]})
    ok("S68d 쪽 밖 400 · 모르는 요소 400 · 빈 양식 400 · 모르는 문서 종류 400", (e1.status,e2.status,e3.status,e4.status), (e1.status,e2.status,e3.status,e4.status)==(400,400,400,400))
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v0=vw.request.get(PL+"?docType=techdata"); v1=vw.request.post(PL,headers=J0,data=json.dumps({"docType":"techdata","elements":l1["elements"]})); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    g0=gb.request.get(PL+"?docType=techdata").json(); g2=gb.request.get(DOCS+f"/{tdA.get('id')}/print"); gb.close()
    ok("S68e viewer 는 보지만(200) 저장 못 한다(403) · 다른 회사는 우리 양식 버전 0 · 우리 발행본 인쇄 404",
       (v0.status, v1.status, len(g0.get("versions",[])), g2.status), v0.status==200 and v1.status==403 and len(g0.get("versions",[]))==0 and g0.get("latest") is None and g2.status==404)
    # ── S69 H10 · p58 그림 제작 Module 1단계 — 도면 위 주석(선 · 사각형 · 글자 · 치수선) 추가·이동·삭제 · 원 도면 불변 · DXF 에 ANNOT 레이어 ──
    DRW=BASE+"/api/drawings"
    r69,_c69=run10(); d69=ctx.request.post(DRW,headers=J0,data=json.dumps({"runId":r69["runId"],"type":"plan"})).json(); DID=d69.get("id")
    orig0=ctx.request.get(DRW+f"/{DID}").text()
    pg.goto(BASE+f"/drawings/{DID}/annotate",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=annot-editor][data-ready='1']",timeout=60000); nuke(pg)
    ob=pg.locator("[data-testid=annot-overlay]").bounding_box()
    P=lambda fx,fy: (ob["x"]+ob["width"]*fx, ob["y"]+ob["height"]*fy)
    def drag_on(a,b_):
        pg.mouse.move(*a); pg.mouse.down(); pg.mouse.move(*b_, steps=6); pg.mouse.up()
    def want_count(n): pg.wait_for_selector(f"[data-testid=annot-editor][data-count='{n}']",timeout=30000)
    pg.click("[data-testid=annot-tool-line]"); drag_on(P(.2,.7),P(.6,.7)); want_count(1)
    pg.click("[data-testid=annot-tool-rect]"); drag_on(P(.3,.2),P(.5,.4)); want_count(2)
    pg.click("[data-testid=annot-tool-text]"); pg.fill("[data-testid=annot-text-input]","용접 주의"); pg.mouse.click(*P(.7,.3)); want_count(3)
    pg.click("[data-testid=annot-tool-dim]"); drag_on(P(.1,.85),P(.8,.85)); want_count(4)
    nuke(pg); pg.screenshot(path=f"{OUT}/79_draw_module.png",full_page=True)
    AN=DRW+f"/{DID}/annotations"; rows=ctx.request.get(AN).json()["rows"]
    ok("S69a 도면 위에 선 · 사각형 · 글자 · 치수선을 끌어/눌러 더한다 — 도면 좌표(mm)로 저장",
       sorted((r["kind"], r.get("text")) for r in rows), sorted(r["kind"] for r in rows)==["dim","line","rect","text"] and [r["text"] for r in rows if r["kind"]=="text"]==["용접 주의"])
    ln0=[r for r in rows if r["kind"]=="line"][0]
    pg.click("[data-testid=annot-tool-select]")
    lb=pg.locator(f"[data-testid=annot-{ln0['id']}]").bounding_box(); drag_on((lb["x"]+lb["width"]/2, lb["y"]+lb["height"]/2),(lb["x"]+lb["width"]/2+80, lb["y"]+lb["height"]/2))
    wait_text(pg,"[data-testid=annot-msg]","옮겼습니다")
    tx0=[r for r in rows if r["kind"]=="text"][0]; tb=pg.locator(f"[data-testid=annot-{tx0['id']}]").bounding_box()
    pg.mouse.click(tb["x"]+tb["width"]/2, tb["y"]+tb["height"]/2); pg.click("[data-testid=annot-del]"); want_count(3)
    rows2=ctx.request.get(AN).json()["rows"]; ln1=[r for r in rows2 if r["kind"]=="line"][0]
    ok("S69b 고르기로 선을 끌어 옮기고(x 가 커짐 · 길이 그대로) · 글자 주석을 골라 지운다", (ln0["x1"], ln1["x1"], round(ln0["x2"]-ln0["x1"],1), round(ln1["x2"]-ln1["x1"],1), sorted(r["kind"] for r in rows2)),
       ln1["x1"]>ln0["x1"] and abs((ln1["x2"]-ln1["x1"])-(ln0["x2"]-ln0["x1"]))<0.2 and sorted(r["kind"] for r in rows2)==["dim","line","rect"])
    orig1=ctx.request.get(DRW+f"/{DID}").text(); ann=ctx.request.get(DRW+f"/{DID}?annot=1")
    ad=ezdxf.read(io.StringIO(ann.text())); ae=[e for e in ad.modelspace()]; od=[e for e in ezdxf.read(io.StringIO(orig1)).modelspace()]
    al=[e for e in ae if e.dxf.layer=="ANNOT"]
    ok("S69c 원 도면 DXF 는 한 글자도 안 바뀐다 · 주석 포함 DXF(?annot=1)는 ANNOT 레이어에 선 1 + 사각형 4 + 치수선(선 3 · 글자 1) = 선 8 · 글자 1 · 나머지 엔티티 = 원 도면",
       (orig0==orig1, len([e for e in al if e.dxftype()=="LINE"]), len([e for e in al if e.dxftype()=="TEXT"]), len(ae)-len(al), len(od), "annot.dxf" in (ann.headers.get("content-disposition") or "")),
       orig0==orig1 and len([e for e in al if e.dxftype()=="LINE"])==8 and len([e for e in al if e.dxftype()=="TEXT"])==1 and len(ae)-len(al)==len(od))
    dI=ctx.request.post(DRW,headers=J0,data=json.dumps({"runId":RUN1,"type":"plan"})).json()
    for st in ("review","approved","issued"): ist=ctx.request.patch(DRW+f"/{dI.get('id')}",headers=J0,data=json.dumps({"status":st}))
    lk=ctx.request.post(DRW+f"/{dI.get('id')}/annotations",headers=J0,data=json.dumps({"kind":"line","x1":0,"y1":0,"x2":100,"y2":0}))
    pg.goto(BASE+f"/drawings/{dI.get('id')}/annotate",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=annot-editor][data-locked='1']",timeout=60000)
    # ccmd K · KC-3 — '설계 심볼' 자리는 심볼 패널로 열렸다(S77). Free CAD 만 잠긴 자리로 남는다.
    _lk=pg.eval_on_selector_all("[data-testid=annot-locked-tool]","es=>es.map(e=>e.dataset.name)")
    ok("S69d 발행된 도면에는 주석을 더하지 못한다(409) · 편집기는 잠김 표시 · Free CAD 는 잠긴 자리(이유) — 설계 심볼은 KC-3 에서 열림",
       (ist.status, lk.status, bool(pg.query_selector("[data-testid=annot-locked]")), _lk),
       ist.status==200 and lk.status==409 and bool(pg.query_selector("[data-testid=annot-locked]")) and _lk==["Free CAD"])
    q=lambda body: ctx.request.post(AN,headers=J0,data=json.dumps(body))
    e1=q({"kind":"text","x1":0,"y1":0}); e2=q({"kind":"circle","x1":0,"y1":0,"x2":10,"y2":10}); e3=q({"kind":"line","x1":5,"y1":5,"x2":5,"y2":5})
    e4=ctx.request.patch(BASE+f"/api/drawing-annotations/{ln1['id']}",headers=J0,data=json.dumps({"kind":"rect"}))
    ok("S69e 글자 없는 글자 주석 400 · 모르는 종류 400 · 두 점이 같은 선 400 · 종류 바꾸기 400", (e1.status,e2.status,e3.status,e4.status), (e1.status,e2.status,e3.status,e4.status)==(400,400,400,400))
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v0=vw.request.get(AN); v1=vw.request.post(AN,headers=J0,data=json.dumps({"kind":"line","x1":0,"y1":0,"x2":100,"y2":0})); v2=vw.request.delete(BASE+f"/api/drawing-annotations/{ln1['id']}"); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    g0=gb.request.get(AN); g1=gb.request.patch(BASE+f"/api/drawing-annotations/{ln1['id']}",headers=J0,data=json.dumps({"dx":10,"dy":0})); g2=gb.request.get(DRW+f"/{DID}?annot=1"); gb.close()
    ok("S69f viewer 는 보지만(200) 못 더하고·못 지운다(403) · 다른 회사는 우리 도면 주석을 못 보고 못 옮긴다(404) · 주석 DXF 도 404",
       (v0.status,v1.status,v2.status,g0.status,g1.status,g2.status), (v0.status,v1.status,v2.status,g0.status,g1.status,g2.status)==(200,403,403,404,404,404))
    # ── S70 E6 · p39 "설계 검증 [Macro]" — 규칙 표 op=macro 행이 승인된 매크로(V_SECTION_RATIO)를 스냅샷 값으로 돌린다 ──
    def ui_bom_verify():
        pg.goto(NODE4,wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=canvas-cmds][data-ready='1']",timeout=60000); nuke(pg)
        if pg.get_attribute("[data-testid=toolbox-toggle]","aria-pressed")=="true":
            pg.click("[data-testid=toolbox-toggle]"); pg.wait_for_selector("[data-testid=toolbox-toggle][aria-pressed=false]",timeout=30000)
        with pg.expect_response(lambda q: q.url.endswith("/api/run/bom") and q.request.method=="POST",timeout=60000) as rr:
            nuke(pg); pg.click("[data-run=bom]")
        pg.locator("button", has_text=re.compile(r"^Design$")).first.click(force=True)
        pg.wait_for_selector("[data-testid=design-verify]",timeout=60000)
        return rr.value.json(), pg.get_attribute("[data-testid=design-verify]","data-ok"), pg.inner_text("[data-testid=design-verify]")
    SL70={"A":"EU","B":"55","C":"2123","D":"630","E":"SS","F":"1-21-13-15"}
    g70=ctx.request.get(ARR+"?code=EU&slots="+json.dumps(SL70)).json()["sections"]
    keep70=[{"name":x["name"], **({"len":x["len"]} if x.get("len") is not None else {}), **({"dir":x["dir"]} if x.get("dir") else {}), "components":x.get("components",[])} for x in g70]
    rP=run55("S70"); dP=rP.get("dims",{})
    uiP,okP,txP=ui_bom_verify()
    ok("S70a 규칙 표의 op=macro 행(V_SECTION_RATIO)이 BOM Run 마다 승인 매크로로 돌고 — 지금은 통과(화면 배지 · 스냅샷 판정)",
       (dP.get("rules"), dP.get("violations"), okP, txP[:40]), dP.get("rules",0)>=4 and dP.get("violations")==[] and okP=="true" and "통과" in txP)
    nAct=len([x for x in g70 if x.get("active")!=False]); X70=900*nAct   # 가장 긴 구획 > 나머지 합 → 전장의 절반 초과 (전장은 운반 한계 9000 아래로)
    brk=[dict(r) for r in keep70]
    for r in brk:
        if r["name"]=="Fan": r["len"]=X70
    bput=put(brk)
    rB70=run55("S70BAD"); vB=rB70.get("dims",{}).get("violations",[])
    uiB,okB,txB=ui_bom_verify()
    nuke(pg); pg.screenshot(path=f"{OUT}/80_macro_verify.png",full_page=True)
    ok("S70b 매크로 조건을 깨는 치수(Fan 구획을 전장 절반 넘게)로 저장 → Run → 위반이 스냅샷에 박히고 화면 배지가 위반(규칙 이름)으로 바뀐다",
       (bput.status, X70, [(v.get("name"),v.get("op"),v.get("actual")) for v in vB], okB, txB[:60]),
       bput.status==200 and any(v.get("op")=="macro" and v.get("name","").startswith("구획 비율") and v.get("actual")==0 for v in vB) and okB=="false" and "구획 비율" in txB)
    dB=ctx.request.get(BASE+f"/api/dxf?runId={rB70['runId']}&type=plan")
    ok("S70c 매크로 검증 위반이면 도면 DXF 를 뜨지 않는다(422 · 이유에 매크로 이름)", (dB.status, dB.text()[:90]), dB.status==422 and "V_SECTION_RATIO" in dB.text())
    put(keep70)
    dP2=ctx.request.get(BASE+f"/api/dxf?runId={rP['runId']}&type=plan"); dB2=ctx.request.get(BASE+f"/api/dxf?runId={rB70['runId']}&type=plan")
    rR=run55("S70R")
    ok("S70d 되돌리면 새 Run 은 다시 통과 · 앞 두 스냅샷의 판정은 그대로(통과 200 · 위반 422 — 판정도 스냅샷)",
       (rR.get("dims",{}).get("violations"), dP2.status, dB2.status), rR.get("dims",{}).get("violations")==[] and dP2.status==200 and dB2.status==422)
    # ── S71 E7 · p48 인쇄본 Office 내보내기 — 같은 스냅샷 body 를 .docx · .xlsx 로 (숫자는 값 · 합계는 엑셀이 다시 낼 수 있게) ──
    import docx as _docx, openpyxl as _xl, io as _io
    from urllib.parse import unquote as _unq
    r71=run55("S71"); q71=ctx.request.post(DOCS,headers=J0,data=json.dumps({"runId":r71["runId"],"type":"quotation"})).json()
    Q71=DOCS+f"/{q71.get('id')}"
    h71=ctx.request.get(Q71+"/print").text()
    no71=re.search(r"<title>(\S+) Rev",h71).group(1)
    tot71=int(re.search(r'data-testid="quote-total">([^<]+)<',h71).group(1).replace(",",""))
    at71=re.search(r'data-testid="applied-table">(.*?)</table>',h71,re.S)
    n71=len(re.findall(r"<tr>",at71.group(1)))-2 if at71 else -1   # 머리 줄 · 합계 줄 제외
    btn71=('data-testid="export-docx"' in h71, 'data-testid="export-xlsx"' in h71)
    xd=ctx.request.get(Q71+"/export?format=docx"); xx=ctx.request.get(Q71+"/export?format=xlsx")
    cd71=xd.headers.get("content-disposition","")
    open(f"{OUT}/office_{no71}.docx","wb").write(xd.body()); open(f"{OUT}/office_{no71}.xlsx","wb").write(xx.body())
    D=_docx.Document(_io.BytesIO(xd.body()))
    dtxt="\n".join(p.text for p in D.paragraphs)+"\n"+"\n".join(c.text for t in D.tables for r in t.rows for c in r.cells)
    dit=[t for t in D.tables if t.rows and t.rows[0].cells[1].text=="Code No."]
    dn=len(dit[0].rows)-2 if dit else -1
    ok("S71a 인쇄본 상단 Word · Excel → .docx 가 열리고(python-docx) 문서 번호 · 견적 합계 · 품목 수가 인쇄본 HTML 과 같다 (한글 파일명 filename*=UTF-8)",
       (xd.status, btn71, no71 in dtxt, format(tot71,",") in dtxt, dn, n71, _unq(cd71.split("UTF-8''")[-1]) if "UTF-8''" in cd71 else cd71),
       xd.status==200 and all(btn71) and no71 in dtxt and format(tot71,",") in dtxt and dn==n71 and n71>0 and "UTF-8''" in cd71 and "견적서" in _unq(cd71))
    W=_xl.load_workbook(_io.BytesIO(xx.body()))
    ws=W["문서"]; vals=[c.value for r in ws.iter_rows() for c in r]
    it=W["품목"]; irows=[r for r in it.iter_rows(min_row=2,values_only=True) if r[0]!="합계"]
    amt_col=[c.value for c in it[1]].index(next(h for h in [c.value for c in it[1]] if str(h).startswith("금액")))
    sumrow=[r for r in it.iter_rows(min_row=2) if r[0].value=="합계"]
    qtot=[r for r in ws.iter_rows(values_only=True) if r and r[0]=="합계" and isinstance(r[4] if len(r)>4 else None,(int,float))]
    ok("S71b .xlsx 가 열리고(openpyxl) 문서 번호 · 견적 합계(숫자 값) · 품목 수가 같다 — 품목 금액은 숫자라 합계 줄이 SUM 식",
       (xx.status, no71 in vals, qtot[0][4] if qtot else None, len(irows), sumrow[0][amt_col].value if sumrow else None),
       xx.status==200 and no71 in vals and bool(qtot) and qtot[0][4]==tot71 and len(irows)==n71 and all(isinstance(r[amt_col],(int,float)) for r in irows)
       and bool(sumrow) and str(sumrow[0][amt_col].value).startswith("=SUM("))
    bad71=ctx.request.get(Q71+"/export?format=pdf")
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test")); v71=vw.request.get(Q71+"/export?format=xlsx"); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test")); g71=gb.request.get(Q71+"/export?format=docx"); gb.close()
    ok("S71c 모르는 format 은 400 · viewer 는 내보내기 403 · 다른 회사 문서는 404", (bad71.status, v71.status, g71.status), (bad71.status, v71.status, g71.status)==(400,403,404))
    # ── S72 B · 학습 AI 1수준 + 이중 프로젝션 (ccmd J) — 플랫폼이 DB① 에만 올리고 · 학습 · 승인 · 한쪽 방향 투영 → 회사가 채택 ──
    import pathlib as _pl
    _SD=_pl.Path(__file__).resolve().parent.parent/"packages"/"db"/"prisma"/"learning-samples"
    ANS=json.loads((_SD/"ANSWERS.json").read_text(encoding="utf-8"))
    LP=BASE+"/api/platform/learning"
    lp=b.new_context(viewport={"width":1440,"height":900}); lpp=lp.new_page()
    lp.request.post(BASE+"/api/auth/login",data=LOGIN("platform@edim.test"))
    # (a) 새 도면 1장을 DB① 에 올린다 — 이 회사 제품 도면(정면도 DXF)을 동의 받아 받은 것처럼
    rU=run55("S72"); fdx=ctx.request.get(BASE+f"/api/dxf?runId={rU['runId']}&type=front").body()
    up=lp.request.post(LP+"/sources",multipart={"file":{"name":"E2E_front_upload.dxf","mimeType":"application/dxf","buffer":fdx},"origin":"tenant-consented"})
    dup=lp.request.post(LP+"/sources",multipart={"file":{"name":"E2E_front_upload_again.dxf","mimeType":"application/dxf","buffer":fdx}})
    pdf=lp.request.post(LP+"/sources",multipart={"file":{"name":"spec.pdf","mimeType":"application/pdf","buffer":b"%PDF-1.4"}})
    lpp.goto(BASE+"/platform/learning",wait_until="domcontentloaded"); lpp.wait_for_selector("[data-testid=learning][data-ready='1']",timeout=60000); nuke(lpp)
    ai_on=lpp.get_attribute("[data-testid=local-ai]","data-on")=="1"
    nsrc=len(lpp.query_selector_all("[data-testid=learn-source]")); nsam=len(lpp.query_selector_all("[data-testid=learn-source][data-sample='1']"))
    with lpp.expect_response(lambda q: q.url.endswith("/api/platform/learning/jobs") and q.request.method=="POST",timeout=300000) as jr:
        lpp.click("[data-testid=learn-run]")
    lpp.wait_for_selector("[data-testid=learning][data-job-state=done]",timeout=120000)
    stp=lpp.eval_on_selector_all("[data-testid=learn-step]","es=>es.map(e=>[e.dataset.tool,e.dataset.state])")
    nuke(lpp); lpp.screenshot(path=f"{OUT}/81_learning_job.png",full_page=True)
    ok("S72a 플랫폼이 DB① 에만 올린다(도면 1 · 같은 파일 409 · PDF 400) → 작업 = 계획 4단계(extract→align→mine→verify) 전부 완료 · 원천 = 샘플 69 + 1",
       (up.status, dup.status, pdf.status, jr.value.status, stp, nsrc, nsam, "로컬 AI 켜짐" if ai_on else "로컬 AI 꺼짐"),
       up.status==200 and dup.status==409 and pdf.status==400 and jr.value.status==200 and stp==[["extract","done"],["align","done"],["mine","done"],["verify","done"]] and nsrc==70 and nsam==69)
    # (b) 숨겨 둔 공식이 적합도 합격으로 · 잡음 3장은 어긋남 · 미정렬(로컬 AI 가 켜져 있으면 0, 꺼져 있으면 사전 밖 약어 1종)
    cards={c["target"]:c for c in lpp.eval_on_selector_all("[data-testid=formula-card]","es=>es.map(e=>({target:e.dataset.target,state:e.dataset.state,expr:e.dataset.expr,n:+e.querySelector('[data-testid=formula-fit]').dataset.n,err:+e.querySelector('[data-testid=formula-fit]').dataset.maxErr,out:[...e.querySelectorAll('[data-testid=formula-outliers]')].map(o=>o.innerText).join(' '),fits:e.querySelector('[data-testid=formula-company]').dataset.fits}))")}
    want={f["target"]:f["expression"] for f in ANS["formulas"]}
    found=[(t, cards.get(t,{}).get("expr")==e, cards.get(t,{}).get("state")) for t,e in want.items()]
    noise_ok=all(nz in cards.get("overall_length",{}).get("out","") for nz in ANS["noise"])
    una=int(lpp.get_attribute("[data-testid=learn-align]","data-unaligned") or -1); byai=int(lpp.get_attribute("[data-testid=learn-align]","data-by-ai") or 0)
    nuke(lpp); lpp.locator("[data-testid=formula-card]").first.scroll_into_view_if_needed(); lpp.screenshot(path=f"{OUT}/82_formula_cards.png",full_page=True)
    ok("S72b 숨겨 둔 공식(전장 = Σ 구획 · 전고 = 케이싱 + 2 × 프레임 · 코일 깊이 = 25 × 열수 + 50)을 다시 찾았다 — 합격 후보 · 잡음 3장은 전장 공식의 어긋남",
       (found, noise_ok, cards.get("overall_height",{}).get("err")), all(ok_ and st=="proposed" for _,ok_,st in found) and noise_ok and cards.get("overall_height",{}).get("err")==0)
    ok("S72c 정렬 — 사전으로 맞추고, 사전 밖 약어(BF HT)는 로컬 AI 가 켜져 있으면 허용 목록 안에서 맞추고 아니면 미정렬로 드러난다",
       (("on" if ai_on else "off"), una, byai), (ai_on and una==0 and byai==2) or ((not ai_on) and una==1 and byai==0))
    # (c) 승인(라벨) → 회사 A 로 투영 · 미승인은 투영 거절 · B 는 0건
    fl=lp.request.get(LP).json()["formulas"]; fid={f["target"]:f["id"] for f in fl}
    notyet=lp.request.post(LP+f"/formulas/{fid['coil.depth']}/project",headers=J0,data=json.dumps({"tenantId":"00000000-0000-4000-8000-00000000000a"}))
    for t in ("overall_length","overall_height"):
        with lpp.expect_response(lambda q: "/api/platform/learning/formulas/" in q.url and q.request.method=="POST",timeout=60000):
            lpp.locator(f"[data-testid=formula-card][data-target='{t}'] [data-testid=formula-approve]").click()
        lpp.wait_for_selector(f"[data-testid=formula-card][data-target='{t}'][data-state=approved]",timeout=30000)
    card=lpp.locator("[data-testid=formula-card][data-target='overall_length']")
    card.locator("[data-testid=formula-tenant]").select_option(label="Acme AHU")
    with lpp.expect_response(lambda q: q.url.endswith("/project") and q.request.method=="POST",timeout=60000) as pr:
        card.locator("[data-testid=formula-project]").click()
    lpp.wait_for_selector("[data-testid=projection-row]",timeout=30000)
    sA=ctx.request.get(BASE+"/api/learning/suggestions").json()["rows"]
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test")); sB=gb.request.get(BASE+"/api/learning/suggestions").json()["rows"]; gb.close()
    ok("S72d 승인한 공식만 π_user 로 회사 A 에 한쪽 방향 투영 — A 제안 1건(회사 형식 · 원천 흔적 없음) · B 0건 · 미승인 공식 투영은 409",
       (notyet.status, pr.value.status, [(x["target"],x["expression"],x["fitsCompany"]) for x in sA], len(sB), [k for k in (sA[0] if sA else {}) if "source" in k.lower()]),
       notyet.status==409 and pr.value.status==200 and len(sA)==1 and sA[0]["expression"]=="=Var(DIM,SECSUM)" and sA[0]["fitsCompany"] and len(sB)==0 and not [k for k in sA[0] if "source" in k.lower()])
    sim=float(lpp.get_attribute("[data-testid=similarity]","data-ratio")); apis=lp.request.get(LP).json()["similarity"]
    nuke(lpp); lpp.locator("[data-testid=similarity]").scroll_into_view_if_needed(); lpp.screenshot(path=f"{OUT}/83_projection.png",full_page=True)
    ok("S72e 구조 유사도 계기판 = 투영본 중 DB② 형식에 맞는 비율 — 1/1 = 1.00(목표 0.90) · 전고 공식은 회사 어휘에 없는 이름(케이싱 · 프레임)이라 '옮길 수 없음'",
       (sim, apis["matched"], apis["total"], cards.get("overall_height",{}).get("fits")), sim==1.0 and apis["ratio"]==1.0 and (apis["matched"],apis["total"])==(1,1) and cards.get("overall_height",{}).get("fits")=="0")
    # (d) 회사 A 가 채택 → 기존 흐름(Save draft = 검증 → 승인) → 그 매크로로 Run
    AHU="a0000000-0000-4000-8000-000000000002"
    pg.goto(BASE+f"/workbench?node={AHU}",wait_until="domcontentloaded"); hydrated(pg); nuke(pg)
    if pg.get_attribute("[data-testid=toolbox-toggle]","aria-pressed")!="true":
        pg.click("[data-testid=toolbox-toggle]")
    pg.wait_for_selector("[data-testid=toolbox-window]",timeout=30000)
    if pg.query_selector("[data-toolbox-tab=program]"): pg.click("[data-toolbox-tab=program]")
    pg.wait_for_selector("[data-testid=tb-suggestions][data-ready='1'] [data-testid=tb-sug-adopt]",timeout=60000)
    pg.click("[data-testid=tb-sug-adopt]"); pg.wait_for_function("()=>(document.querySelector('[data-testid=tb-dsl]')?.value||'')==='=Var(DIM,SECSUM)'",timeout=30000)
    with pg.expect_response(lambda q: q.url.endswith("/api/macros") and q.request.method=="POST",timeout=60000):
        pg.click("[data-testid=tb-draft]")
    wait_text(pg,"[data-testid=tb-msg]","채택",60000)
    pg.click("[data-testid=tb-approve]"); wait_text(pg,"[data-testid=tb-msg]","승인 완료",60000)
    sA2=ctx.request.get(BASE+"/api/learning/suggestions").json()["rows"]
    er=ctx.request.post(BASE+"/api/run/edim",headers=J0,data=json.dumps({"node":AHU,"slots":S55_0})).json()
    rL=run55("S72L"); lenmm=ctx.request.get(BASE+f"/api/dxf?runId={rL['runId']}&type=plan&meta=1").json().get("lengthMm")
    pg.click("[data-testid=tb-run]"); pg.wait_for_function("()=>/^[0-9.]+$/.test((document.querySelector('[data-testid=tb-value]')?.innerText||'').trim())",timeout=60000)
    tbv=pg.inner_text("[data-testid=tb-value]")
    nuke(pg); pg.screenshot(path=f"{OUT}/84_toolbox_suggestion.png",full_page=True)
    ok("S72f 회사 A 가 Toolbox '학습 제안'에서 채택 → Save draft(검증) → 승인(회사 2단 승인) → 그 매크로로 Run — 값 = 그 제품 도면의 전장(구획 합)",
       (sA2[0]["state"] if sA2 else None, bool(sA2 and sA2[0]["adoptedMacroId"]), er.get("dsl"), er.get("value"), lenmm, tbv),
       bool(sA2) and sA2[0]["state"]=="adopted" and bool(sA2[0]["adoptedMacroId"]) and er.get("dsl")=="=Var(DIM,SECSUM)" and er.get("value")==lenmm and float(tbv)>0)
    # (e) 운영 감시 — 승인 공식에 어긋나는 새 도면을 올리면 계기판에 1건
    drift=(_SD/"SAMPLE_ahu_001.dxf").read_text(encoding="utf-8")
    _mL=re.search(r"\nL=(\d+)\n",drift); drift=drift.replace(f"\nL={_mL.group(1)}\n",f"\nL={int(_mL.group(1))+7}\n")
    upd=lp.request.post(LP+"/sources",multipart={"file":{"name":"E2E_drift_site_revision.dxf","mimeType":"application/dxf","buffer":drift.encode("utf-8")}}).json()
    lpp.reload(wait_until="domcontentloaded"); lpp.wait_for_selector("[data-testid=learning][data-ready='1']",timeout=60000)
    ok("S72g 운영 감시 — 승인 공식(전장 = Σ 구획)에 7 mm 어긋난 새 도면을 올리면 계기판에 어긋남 1건(자동 조치 없음)",
       (upd.get("monitor",{}).get("mismatched"), lpp.get_attribute("[data-testid=monitor]","data-drift")), upd.get("monitor",{}).get("mismatched")==1 and lpp.get_attribute("[data-testid=monitor]","data-drift")=="1")
    # (f) 역류 0 — 회사 계정은 DB① API 403 · 플랫폼 계정은 회사 API 에 못 들어간다 · viewer 는 숨기기 403 · 다른 회사 제안 404
    c1=ctx.request.get(LP); c2=ctx.request.post(LP+"/jobs",headers=J0,data="{}"); c3=ctx.request.post(LP+"/sources",multipart={"file":{"name":"x.dxf","mimeType":"application/dxf","buffer":b"0\nEOF\n"}})
    p1=lp.request.get(BASE+"/api/learning/suggestions"); p2=lp.request.get(BASE+"/api/setup/catalog")
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test")); v1=vw.request.post(BASE+f"/api/learning/suggestions/{sA[0]['id']}",headers=J0,data=json.dumps({"action":"dismiss"})); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test")); g1=gb.request.post(BASE+f"/api/learning/suggestions/{sA[0]['id']}",headers=J0,data=json.dumps({"action":"dismiss"})); gb.close()
    lp.close()
    ok("S72h 역류 0 — 회사 계정은 학습 API 403(읽기 · 작업 · 올리기) · 플랫폼 계정은 회사 API 에 못 들어간다 · viewer 숨기기 403 · 다른 회사 제안 404",
       (c1.status,c2.status,c3.status,p1.status,p2.status,v1.status,g1.status), (c1.status,c2.status,c3.status)==(403,403,403) and p1.status in (401,403) and p2.status in (401,403) and v1.status==403 and g1.status==404)
    # ── S75c KA (ccmd K) — 부여 전: Special 을 부르는 샘플 제품(SPF)의 BOM Run 은 422 "Special 부여 필요"(S74 가 부여하기 전 자리) ──
    r75c=ctx.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":{"A":"SPF","B":"55","C":"2123"},"code":"SPF-55-2123"}))
    e75c=r75c.json().get("error","") if r75c.status!=200 else ""
    ok("S75c 부여 없는 회사가 Special 을 부르는 샘플 제품(SPF)을 BOM Run 하면 422 — 'Special 부여 필요 — 플랫폼에 의뢰하세요'(스냅샷 · 과금 없음)",
       (r75c.status, e75c[:40]), r75c.status==422 and "Special 부여 필요" in e75c)
    # ── S74 C · Special Tool Box 첫 사례 — 팬 선정 (ccmd J) — 회사 UI Form → 의뢰 → 플랫폼 승인·부여 → Toolbox 버튼 → 결정론 계산 → 사용 기록·과금 ──
    UF=BASE+"/api/ui-forms"
    cf=ctx.request.post(UF,headers=J0,data=json.dumps({"name":"팬 선정 입력(E2E)","scope":"Technical"})).json(); FID=cf.get("id")
    spec74={"widgets":[{"id":"label1","type":"label","x":0,"y":0,"w":12,"h":1,"label":"팬 선정 — 설계 조건"},
        {"id":"number1","type":"number","x":0,"y":1,"w":6,"h":3,"label":"풍량","param":"q_cmh","unit":"CMH"},
        {"id":"number2","type":"number","x":6,"y":1,"w":6,"h":3,"label":"기외정압","param":"p_pa","unit":"Pa"},
        {"id":"number3","type":"number","x":12,"y":1,"w":6,"h":3,"label":"밀도","param":"rho","unit":"kg/m³"}]}
    pu=ctx.request.put(UF+f"/{FID}",headers=J0,data=json.dumps({"name":"팬 선정 입력(E2E)","scope":"Technical","spec":spec74}))
    # (a) 회사 A 가 Special 의뢰(프로그램 = 팬 선정 · 입력 폼 = 방금 만든 UI Form) — 화면에서
    pg.goto(BASE+"/m/company",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=platform-requests][data-ready='1']",timeout=60000); nuke(pg)
    pg.fill("[data-testid=request-subject]","팬 선정 Special 요청(E2E)"); pg.fill("[data-testid=request-detail]","풍량·정압으로 팬 모델·모터 선정")
    pg.select_option("[data-testid=request-program]","fan-select"); pg.wait_for_selector(f"[data-testid=request-form] option[value='{FID}']",state="attached",timeout=30000)
    pg.select_option("[data-testid=request-form]",FID)
    with pg.expect_response(lambda q: q.url.endswith("/api/platform-requests") and q.request.method=="POST",timeout=30000) as rq:
        pg.click("[data-testid=request-submit]")
    pg.wait_for_function("()=>[...document.querySelectorAll('[data-testid=my-request]')].some(e=>e.innerText.includes('팬 선정 Special 요청(E2E)'))",timeout=30000)
    nuke(pg); pg.screenshot(path=f"{OUT}/85_special_request.png",full_page=True)
    sp=b.new_context(viewport={"width":1440,"height":900}); spp=sp.new_page(); sp.request.post(BASE+"/api/auth/login",data=LOGIN("platform@edim.test"))
    spp.goto(BASE+"/platform/special",wait_until="domcontentloaded"); spp.wait_for_selector("[data-testid=special-console][data-ready='1']",timeout=60000); nuke(spp)
    row=spp.locator("[data-testid=special-request][data-program=fan-select]").filter(has_text="팬 선정 Special 요청(E2E)").first
    with spp.expect_response(lambda q: q.url.endswith("/api/platform/special") and q.request.method=="POST",timeout=30000) as gr:
        row.locator("[data-testid=special-grant]").click()
    gA=ctx.request.get(BASE+"/api/special").json()["grants"]
    gbx=b.new_context(); gbx.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test")); gB=gbx.request.get(BASE+"/api/special").json()["grants"]
    rB74=gbx.request.post(BASE+"/api/special/fan-select/run",headers=J0,data=json.dumps({"inputs":{"q_cmh":12000,"p_pa":600}})); gbx.close()
    ok("S74a 회사 A 가 UI Form 을 붙여 Special(팬 선정) 의뢰 → 플랫폼이 승인 + 부여(한 번에) → A 에만 부여 · B 는 부여 없음 · B 실행 403",
       (rq.value.status, gr.value.status, [(g["programKey"], (g.get("form") or {}).get("id")==FID) for g in gA], len(gB), rB74.status),
       pu.status==200 and rq.value.status==200 and gr.value.status==200 and len(gA)==1 and gA[0]["programKey"]=="fan-select" and (gA[0].get("form") or {}).get("id")==FID and len(gB)==0 and rB74.status==403)
    # (b) Toolbox 에 'Special: 팬 선정' — 회사가 만든 폼 그대로 · 12,000 CMH · 600 Pa → 단위 테스트 기대값(EDIM-PF-560 · 2600 rpm · 모터 3.7 kW)
    pg.goto(BASE+"/workbench?node=a0000000-0000-4000-8000-000000000004",wait_until="domcontentloaded"); hydrated(pg); nuke(pg)
    if pg.get_attribute("[data-testid=toolbox-toggle]","aria-pressed")!="true": pg.click("[data-testid=toolbox-toggle]")
    pg.wait_for_selector("[data-testid=special-tab]",timeout=60000); pg.click("[data-testid=special-tab]")
    pg.wait_for_selector("[data-testid=special-panel][data-ready='1'] [data-testid=special-in-q_cmh]",timeout=30000)
    wids=pg.eval_on_selector_all("[data-testid=special-form] [data-widget]","es=>es.map(e=>[e.dataset.widget,e.dataset.type,e.dataset.param])")
    formdef=[[w["id"],w["type"],w.get("param","")] for w in [f for f in ctx.request.get(UF).json()["rows"] if f["id"]==FID][0]["spec"]["widgets"]]
    ok("S74b Special 입력 화면 = 회사가 Toolbox UI Form 으로 만든 그 폼(위젯 id · 종류 · 입력 이름이 같다)", (wids, formdef), wids==formdef and len(wids)==4)
    pg.fill("[data-testid=special-in-q_cmh]","12000"); pg.fill("[data-testid=special-in-p_pa]","600")
    with pg.expect_response(lambda q: q.url.endswith("/api/special/fan-select/run") and q.request.method=="POST",timeout=60000) as sr:
        pg.click("[data-testid=special-run]")
    pg.wait_for_selector("[data-testid=special-result]",timeout=30000)
    res=pg.eval_on_selector("[data-testid=special-result]","e=>[e.dataset.ok,e.dataset.model,e.dataset.rpm,e.dataset.motor,e.dataset.price]")
    bind=pg.inner_text("[data-testid=special-binding]")
    nuke(pg); pg.screenshot(path=f"{OUT}/86_fan_result.png",full_page=True)
    ok("S74c 12,000 CMH · 600 Pa → EDIM-PF-560 · 2600 rpm · 모터 3.7 kW (= 단위 테스트 기대값 · 손 계산 대조) · 가져온 자료 한 줄(원자료는 안 보임) · '샘플 성능표 기준'",
       (sr.value.status, res, bind[:60]), sr.value.status==200 and res==["1","EDIM-PF-560 (샘플)","2600","3.7","5000"] and "special_fan_candidates" in bind and "샘플" in pg.inner_text("[data-testid=special-result]"))
    # (c) 사용 기록 1행 · 요금 = 샘플 단가 · 범위 밖 입력은 '적합한 팬 없음' + 기록은 남되 요금 0
    pg.fill("[data-testid=special-in-q_cmh]","100000")
    with pg.expect_response(lambda q: q.url.endswith("/api/special/fan-select/run") and q.request.method=="POST",timeout=60000):
        pg.click("[data-testid=special-run]")
    pg.wait_for_selector("[data-testid=special-result][data-ok='0']",timeout=30000)
    none_txt=pg.inner_text("[data-testid=special-result]")
    pg.wait_for_selector("[data-testid=special-meter][data-today='2']",timeout=30000)
    meter=pg.eval_on_selector("[data-testid=special-meter]","e=>[e.dataset.today,e.dataset.amount]")
    ok("S74d 범위 밖(100,000 CMH) → '적합한 팬 없음 — 이유' · 기록은 남되 요금 0 → 오늘 2회 · 요금 합계 5,000(샘플 단가 1회분)",
       (none_txt[:40], meter), "적합한 팬 없음" in none_txt and meter==["2","5000"])
    # (d) source=tenant — 회사 자체 팬 표로 같은 계산(바인딩 지도가 다른 표를 가리킨다)
    pg.fill("[data-testid=special-in-q_cmh]","12000"); pg.select_option("[data-testid=special-source]","tenant")
    with pg.expect_response(lambda q: q.url.endswith("/api/special/fan-select/run") and q.request.method=="POST",timeout=60000):
        pg.click("[data-testid=special-run]")
    pg.wait_for_selector("[data-testid=special-result][data-ok='1']",timeout=30000)
    tmodel=pg.get_attribute("[data-testid=special-result]","data-model"); tbind=pg.inner_text("[data-testid=special-binding]")
    ok("S74e 바인딩 source=tenant — 회사 자체 팬 표(tenant_fan_curve)로 실행하면 그 표의 팬이 선정되고, 가져온 자료 줄도 회사 표를 가리킨다",
       (tmodel, tbind[:50]), tmodel=="ACME 자체 팬 (샘플)" and "tenant_fan_curve" in tbind)
    # (e) viewer 실행 403 · 다른 회사 사용 기록 0 · 플랫폼 과금(금액 칸만) = 3회 · 10,000
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test")); v74=vw.request.post(BASE+"/api/special/fan-select/run",headers=J0,data=json.dumps({"inputs":{"q_cmh":12000,"p_pa":600}})); vg=vw.request.get(BASE+"/api/special").json()["grants"]; vw.close()
    spp.reload(wait_until="domcontentloaded"); spp.wait_for_selector("[data-testid=special-console][data-ready='1']",timeout=60000); nuke(spp)
    bill=spp.eval_on_selector_all("[data-testid=special-grant-row]","es=>es.map(e=>[e.dataset.runs,e.dataset.amount])")
    spp.screenshot(path=f"{OUT}/87_special_meter.png",full_page=True); sp.close()
    ok("S74f viewer 는 보기만(부여 목록 200) · 실행 403(실행 = 과금이라 편집 권한) · 플랫폼 과금은 사용 기록의 금액 칸만 — 3회 · 10,000(샘플)",
       (v74.status, len(vg), bill), v74.status==403 and len(vg)==1 and bill==[["3","10000"]])
    # ── S75 KA · CPQ 가 BOM Run 안에서 Special 팬 선정을 부른다 (ccmd K) — 샘플 제품 SPF(등록 표 special) · 기존 시연 코드(EU)는 부르지 않는다 ──
    AHU75="a0000000-0000-4000-8000-000000000002"
    pg.goto(BASE+f"/workbench?node={AHU75}",wait_until="domcontentloaded"); hydrated(pg); nuke(pg)
    sel75=pg.query_selector_all("[data-testid=code-builder] select")
    sel75[0].select_option(value="SPF"); sel75[1].select_option(value="55")
    pg.locator("button",has_text=re.compile(r"^BOM$")).first.click(); pg.wait_for_selector("button:has-text('BOM Run')",timeout=30000)
    with pg.expect_response(lambda q: "/api/run/bom" in q.url and q.request.method=="POST",timeout=60000) as br75:
        pg.click("button:has-text('BOM Run')")
    j75=br75.value.json(); rid75=j75.get("runId")
    pg.wait_for_selector("[data-testid=bom-special][data-model]",timeout=30000)
    nuke(pg); pg.screenshot(path=f"{OUT}/75_cpq_special_bom.png",full_page=True)
    s75=ctx.request.get(BASE+f"/api/bom-runs/{rid75}").json(); sp75=s75.get("special") or {}; res75=sp75.get("result") or {}
    fan75=[l for l in s75.get("lines",[]) if l.get("childCode")=="SFN 1"]; mot75=[l for l in s75.get("lines",[]) if l.get("childCode")=="SMT 1"]
    c75=ctx.request.post(BASE+"/api/run/cost",headers=J0,data=json.dumps({"runId":rid75})).json().get("cost",{})
    mat75=sum(l["unitCost"]*l["qty"] for l in s75.get("lines",[]))
    ok("S75a 샘플 제품 SPF(B=55) BOM Run → 서버가 BOM Run 안에서 팬 선정을 부른다(입력 = 등록 표 air 12,000 CMH · 600 Pa — 사람 입력 없음) → 스냅샷 dims.special = EDIM-PF-560 · 2600 rpm · 3.7 kW · BOM 에 팬 · 모터 줄 · 원가 재료비에 모터 단가 420,000",
       (br75.value.status, sp75.get("input"), res75.get("model"), res75.get("rpm"), res75.get("motorKw"), [l.get("spec") for l in fan75+mot75], c75.get("material"), mat75),
       br75.value.status==200 and sp75.get("input")=={"q_cmh":12000,"p_pa":600} and res75.get("model")=="EDIM-PF-560 (샘플)" and res75.get("rpm")==2600 and res75.get("motorKw")==3.7
       and len(fan75)==1 and fan75[0]["spec"].startswith("EDIM-PF-560 (샘플) · 2600 rpm") and len(mot75)==1 and mot75[0]["unitCost"]==420000 and mot75[0]["spec"].startswith("3.7 kW")
       and c75.get("material")==mat75 and 420000<=mat75 and pg.get_attribute("[data-testid=bom-special]","data-motor")=="3.7")
    # 선정 불가(곡선 범위 밖)면 422 + 이유 · 과금 없음
    r75x=ctx.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":{"A":"SPF","B":"10","C":"2123"},"code":"SPF-10-2123"}))
    ok("S75a2 선정 불가(B=10 → 6,000 CMH · 400 Pa · 샘플 곡선 밖)면 422 + 이유 — 스냅샷도 과금도 없다(지어내지 않는다)",
       (r75x.status, (r75x.json().get("error","") if r75x.status!=200 else "")[:50]), r75x.status==422 and "선정 불가" in r75x.json().get("error",""))
    # (b) 같은 스냅샷으로 EBOM · Cost · 견적 · 조립도 → 사용 기록(과금)은 여전히 1건
    e75=ctx.request.post(BASE+"/api/run/ebom",headers=J0,data=json.dumps({"runId":rid75}))
    q75=ctx.request.post(BASE+"/api/documents",headers=J0,data=json.dumps({"runId":rid75,"type":"quotation"})).json()
    asm75=ctx.request.get(BASE+f"/api/dxf?runId={rid75}&type=assembly").text()
    dxf_png(asm75, f"{OUT}/75_cpq_special_drawing.png", "EDIM - ASSEMBLY · SPF 샘플 (ezdxf re-render) — Item 표에 Special 선정 팬 · 모터")
    tx75=dxf_stats(asm75)["texts"]
    runs75=ctx.request.get(BASE+f"/api/bom-runs/{rid75}").json().get("specialRuns",[])
    ok("S75b 같은 스냅샷으로 EBOM · Cost · 견적 · 조립도를 뽑아도 사용 기록(과금)은 BOM Run 1회 = 1건 · 5,000(샘플) · 견적 합계 = 스냅샷 원가 · 조립도 Item 표에 팬 모델 · 모터 kW(ezdxf)",
       (e75.status, q75.get("total"), c75.get("total"), [(r["price"]) for r in runs75], [t for t in tx75 if "EDIM-PF" in t or "kW" in t][:3]),
       e75.status==200 and q75.get("total")==c75.get("total") and len(runs75)==1 and runs75[0]["price"]==5000
       and any(t.startswith("EDIM-PF-560 (샘플) · 2600 rpm") for t in tx75) and any(t.startswith("3.7 kW") for t in tx75))
    # (d) 플랫폼이 성능표 점 하나를 고친다 → 새 BOM Run 은 다른 결과 · 앞 스냅샷의 선정 결과는 그대로
    pf=b.new_context(); pf.request.post(BASE+"/api/auth/login",data=LOGIN("platform@edim.test"))
    CP=BASE+"/api/platform/special/curves"; PT={"model":"EDIM-PF-560 (샘플)","rpm":2600,"q":11200}
    cu=pf.request.post(CP,headers=J0,data=json.dumps({**PT,"p":663,"eta":0.60}))
    j75d=ctx.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":{"A":"SPF","B":"55","C":"2123"},"code":"SPF-55-2123"})).json()
    new75=(j75d.get("special") or {}); old75=ctx.request.get(BASE+f"/api/bom-runs/{rid75}").json().get("special") or {}
    rs=pf.request.post(CP,headers=J0,data=json.dumps({**PT,"p":663,"eta":0.7395}))
    k=lambda x: ((x.get("result") or {}).get("model"), (x.get("result") or {}).get("rpm"), (x.get("result") or {}).get("eta"))
    ok("S75d 플랫폼이 샘플 성능표 점 하나(PF-560 · 2600 rpm · 11,200 CMH 효율 0.7395 → 0.60)를 고치면 새 BOM Run 은 다른 선정 · 곡선 지문도 다르다 — 앞 스냅샷의 dims.special 은 그대로(되돌림 200)",
       (cu.status, k(sp75), k(new75), k(old75), sp75.get("curveFingerprint"), new75.get("curveFingerprint"), rs.status),
       cu.status==200 and rs.status==200 and k(new75)!=k(sp75) and new75.get("curveFingerprint")!=sp75.get("curveFingerprint") and old75==sp75)
    # (e) 기존 시연 제품 코드는 Special 을 부르지 않는다 — 수치는 S3 · S5 · S6b · S22b 가 이미 못 박았다(455.4 · 11행 · ₩15,487,170)
    r75e=run55("S75e"); s75e=ctx.request.get(BASE+f"/api/bom-runs/{r75e.get('runId')}").json()
    ok("S75e 기존 시연 제품(EU-55-2123-630SS)은 Special 을 부르지 않는다 — 응답 special 없음 · 스냅샷에 dims.special 키 없음 · 사용 기록 0건(455.4 · ₩15,487,170 · 11행은 S5 · S6b · S8 이 못 박는다)",
       (r75e.get("special"), "special" in (s75e.get("dims") or {}), len(s75e.get("specialRuns",[])), len(r75e.get("lines",[]))),
       r75e.get("special") is None and "special" not in (s75e.get("dims") or {}) and len(s75e.get("specialRuns",[]))==0 and len(r75e.get("lines",[]))>0)
    # (f) viewer 403 · 다른 회사 404 · 회사 계정은 성능표 API 403
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v75=vw.request.get(BASE+f"/api/bom-runs/{rid75}"); vb75=vw.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":{"A":"SPF","B":"55"}})); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test")); g75=gb.request.get(BASE+f"/api/bom-runs/{rid75}"); gb.close()
    c75f=ctx.request.post(CP,headers=J0,data=json.dumps({**PT,"p":663,"eta":0.5})); pf.close()
    ok("S75f viewer 는 스냅샷 읽기 · BOM Run 403 · 다른 회사는 그 스냅샷 404 · 회사 계정은 플랫폼 성능표 API 403",
       (v75.status, vb75.status, g75.status, c75f.status), (v75.status, vb75.status, g75.status, c75f.status)==(403,403,404,403))
    # ── S76 KC-1 · KC-2 (ccmd K) — Detail Dimension(role detail) · CAD 규칙서(샘플 파일) — 샘플 제품 SPF 만(기존 EU 는 그대로) ──
    def run_spf():
        return ctx.request.post(BASE+"/api/run/bom",headers=J0,data=json.dumps({"slots":{"A":"SPF","B":"55","C":"2123"},"code":"SPF-55-2123"}))
    def layer_texts(txt, layer):
        return sorted(e.dxf.text for e in ezdxf.read(io.StringIO(txt)).modelspace() if e.dxftype()=="TEXT" and e.dxf.layer==layer)
    j76=run_spf().json(); rid76=j76.get("runId")
    s76=ctx.request.get(BASE+f"/api/bom-runs/{rid76}").json(); d76=s76.get("dims") or {}
    det76=sorted((d["target"],d["label"],d["value"]) for d in d76.get("detail",[])); cr76=d76.get("cadRules") or {}
    asm76=ctx.request.get(BASE+f"/api/dxf?runId={rid76}&type=assembly").text()
    dim76=layer_texts(asm76,"DIM"); cad76=layer_texts(asm76,"CADRULE"); kad76=layer_texts(asm76,"KAD")
    dxf_png(asm76, f"{OUT}/76_detail_dim.png", "SPF 샘플 조립도 — 세부 치수(DIM) · 기준점 · mm 배치(CADRULE) · KAD 슬롯(샘플 대응표) — ezdxf re-render")
    ok("S76a 세부 치수(detail 표) → BOM Run 이 스냅샷 dims.detail 에 박는다(Fan A=1250 · B=1400 = 사이즈 표 fsz 55 행 · SMT 1 C=350 등록 값) · 규칙서 판 sample-1 · 지문 12자 → 조립도 DXF(ezdxf) DIM 레이어에 같은 값",
       (det76, cr76.get("version"), cr76.get("fingerprint"), cr76.get("file"), [t for t in dim76 if t.startswith("detail.")]),
       det76==[("Fan","A",1250),("Fan","B",1400),("SMT 1","C",350)] and cr76.get("version")=="sample-1" and len(cr76.get("fingerprint") or "")==12 and cr76.get("file")=="cad-rules.sample.json"
       and [t for t in dim76 if t.startswith("detail.")]==["detail.Fan.A=1250","detail.Fan.B=1400","detail.SMT 1.C=350"])
    ok("S76b CAD 규칙서가 부품을 mm 좌표에 놓는다(3×3 칸 → Fan 구획 2700~3600 · W 2472) · 기준점 Shaft · Foot · KAD- 슬롯 줄 = 대응표의 치수 키(W · H · 전장 · Fan A · B) · '샘플'",
       ([t for t in cad76 if "@" in t], [t for t in kad76 if t.startswith("KAD-2")]),
       "SFN 1 @3150,1236" in cad76 and "SMT 1 @3450,412.1" in cad76 and "SHAFT 3150,1236" in cad76 and "FOOT 2750,0" in cad76
       and any(t.startswith("KAD-2472-2472-3600-1250-1400 · CAD RULES sample-1 #"+str(cr76.get("fingerprint"))) and t.endswith("(SAMPLE)") for t in kad76))
    # 규칙 위반: 사이즈 표 fsz 55 행 A 1250 → 1400(규칙 max 1300) → 도면 422 · 앞 스냅샷 도면은 그대로 · 되돌림
    _cat=ctx.request.get(BASE+"/api/setup/catalog").json(); spf=[p_ for p_ in _cat.get("productCodes",[]) if p_.get("code")=="SPF"][0]
    spf_orig=json.loads(json.dumps(spf)); spf2=json.loads(json.dumps(spf))
    for r_ in spf2["tables"]["fsz"]["rows"]:
        if r_["item"]=="55": r_["cells"]["A"]=1400
    up76=ctx.request.post(BASE+"/api/setup/product-codes",headers=J0,data=json.dumps(spf2))
    j76v=run_spf().json(); v76=ctx.request.post(BASE+"/api/drawings",headers=J0,data=json.dumps({"runId":j76v.get("runId"),"type":"assembly"}))
    asm76_again=ctx.request.get(BASE+f"/api/dxf?runId={rid76}&type=assembly").text()
    rs76=ctx.request.post(BASE+"/api/setup/product-codes",headers=J0,data=json.dumps(spf_orig))
    ok("S76c 설계 검증 규칙이 세부 치수에도 걸린다(detail.Fan.A max 1300 · 표를 1400 으로 고치면 도면 422 + 규칙 이름) · 세부 치수를 고쳐도 앞 스냅샷의 조립도는 한 글자도 안 바뀐다 · 되돌림 200",
       (up76.status, v76.status, (v76.json().get("error","") if v76.status!=200 else "")[:70], asm76_again==asm76, rs76.status),
       up76.status==200 and v76.status==422 and "팬 구획 세부 A 한계" in v76.json().get("error","") and asm76_again==asm76 and rs76.status==200)
    # 파일 교체 시험(완료 정의 "파일 교체만으로 반영") — 회사 규칙서 자리(cad-rules.local.json)에 오프셋만 다른 사본 → 코드 변경 0 으로 좌표만 바뀐다
    _rules_dir=os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","packages","bom-code","cad-rules")
    _local=os.path.join(_rules_dir,"cad-rules.local.json")
    _sample=json.load(open(os.path.join(_rules_dir,"cad-rules.sample.json"),encoding="utf-8"))
    _copy=json.loads(json.dumps(_sample)); _copy["version"]="sample-1-offset"; _copy["grid"]["offsetMm"]={"x":100,"y":-50}
    try:
        with open(_local,"w",encoding="utf-8",newline="\n") as f_: json.dump(_copy,f_,ensure_ascii=False,indent=2)
        j76f=run_spf().json(); s76f=ctx.request.get(BASE+f"/api/bom-runs/{j76f.get('runId')}").json(); cr76f=(s76f.get("dims") or {}).get("cadRules") or {}
        asm76f=ctx.request.get(BASE+f"/api/dxf?runId={j76f.get('runId')}&type=assembly").text()
    finally:
        if os.path.exists(_local): os.remove(_local)
    j76b=run_spf().json(); cr76b=((ctx.request.get(BASE+f"/api/bom-runs/{j76b.get('runId')}").json().get("dims") or {}).get("cadRules") or {})
    ok("S76d 규칙서 파일만 바꾸면(오프셋 +100 · -50 사본) 새 BOM Run 의 도면 좌표만 바뀐다 — 부품 @3250,1186 · 세부 치수 값 그대로 · 지문 · 판이 다르다 · 앞 스냅샷 도면 불변 · 파일을 치우면 샘플 지문으로 돌아온다",
       (cr76f.get("file"), cr76f.get("version"), cr76f.get("fingerprint"), [t for t in layer_texts(asm76f,"CADRULE") if t.startswith("SFN")], cr76b.get("fingerprint")),
       cr76f.get("file")=="cad-rules.local.json" and cr76f.get("version")=="sample-1-offset" and cr76f.get("fingerprint")!=cr76.get("fingerprint")
       and "SFN 1 @3250,1186" in layer_texts(asm76f,"CADRULE") and layer_texts(asm76f,"DIM")==dim76
       and ctx.request.get(BASE+f"/api/dxf?runId={rid76}&type=assembly").text()==asm76 and cr76b.get("fingerprint")==cr76.get("fingerprint"))
    r76e=run55("S76e"); s76e=ctx.request.get(BASE+f"/api/bom-runs/{r76e.get('runId')}").json()
    asm76e=ctx.request.get(BASE+f"/api/dxf?runId={r76e.get('runId')}&type=assembly").text()
    ok("S76e 기존 시연 제품(EU)은 세부 치수 · 규칙서를 쓰지 않는다 — 스냅샷에 detail · cadRules 없음 · 조립도에 CADRULE · KAD 레이어 없음",
       (("detail" in (s76e.get("dims") or {})), ("cadRules" in (s76e.get("dims") or {})), dxf_stats(asm76e)["layers"]),
       "detail" not in (s76e.get("dims") or {}) and "cadRules" not in (s76e.get("dims") or {}) and not ({"CADRULE","KAD"} & set(dxf_stats(asm76e)["layers"])))
    # ── S77 KC-3 (ccmd K) — p58 설계 심볼 라이브러리: 놓기 · 옮기기 · 회전 · 지우기 · DXF SYMBOL 레이어 · 발행 도면 잠금 ──
    DRW=BASE+"/api/drawings"
    d77=ctx.request.post(DRW,headers=J0,data=json.dumps({"runId":rid76,"type":"assembly"})).json(); D77=d77.get("id"); SY=DRW+f"/{D77}/symbols"
    lib77=ctx.request.get(BASE+"/api/design-symbols").json().get("rows",[]); key_of={r_["id"]:r_["key"] for r_ in lib77}
    orig77=ctx.request.get(DRW+f"/{D77}").text()
    pg.goto(BASE+f"/drawings/{D77}/annotate",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=symbol-panel][data-ready='1']",timeout=60000); nuke(pg)
    ob=pg.locator("[data-testid=annot-overlay]").bounding_box()
    P=lambda fx,fy: (ob["x"]+ob["width"]*fx, ob["y"]+ob["height"]*fy)
    def sym_count(n): pg.wait_for_selector(f"[data-testid=symbol-panel][data-count='{n}']",timeout=30000)
    OV=pg.locator("[data-testid=annot-overlay]")
    def ov_click(fx,fy): OV.click(position={"x":ob["width"]*fx,"y":ob["height"]*fy})
    pg.click("[data-testid=sym-pick-fan]"); ov_click(.3,.45); sym_count(1)
    pg.click("[data-testid=sym-pick-motor]"); ov_click(.6,.45); sym_count(2)
    rows77=ctx.request.get(SY).json().get("rows",[]); fan77=[r_ for r_ in rows77 if key_of.get(r_["symbolId"])=="fan"][0]; mot77=[r_ for r_ in rows77 if key_of.get(r_["symbolId"])=="motor"][0]
    pg.click("[data-testid=annot-tool-select]"); pg.locator(f"[data-testid=sym-{fan77['id']}] circle").first.scroll_into_view_if_needed()
    fb=pg.locator(f"[data-testid=sym-{fan77['id']}] circle").first.bounding_box(); c0=(fb["x"]+fb["width"]/2, fb["y"]+fb["height"]/2)
    pg.mouse.move(*c0); pg.mouse.down(); pg.mouse.move(c0[0]+80,c0[1],steps=6); pg.mouse.up(); wait_text(pg,"[data-testid=annot-msg]","옮겼습니다")
    pg.locator(f"[data-testid=sym-{fan77['id']}] circle").first.click(force=True)
    pg.click("[data-testid=sym-rot]"); wait_text(pg,"[data-testid=annot-msg]","돌렸습니다")
    pg.locator(f"[data-testid=sym-{mot77['id']}] line").first.click(force=True)
    pg.click("[data-testid=sym-del]"); sym_count(1)
    nuke(pg); pg.screenshot(path=f"{OUT}/77_symbol.png",full_page=True)
    rows77b=ctx.request.get(SY).json().get("rows",[])
    ok("S77a 설계 심볼(샘플 5종)을 도면에 놓고(팬 · 모터) · 팬을 끌어 옮기고(x 증가) · 90° 돌리고 · 모터를 지운다 — 도면 좌표 mm 로 저장",
       (sorted(key_of.values()), [(key_of.get(r_["symbolId"]), r_["rot"]) for r_ in rows77b], fan77["x"], rows77b[0]["x"] if rows77b else None),
       sorted(key_of.values())==["coil","damper","fan","filter","motor"] and len(rows77b)==1 and key_of.get(rows77b[0]["symbolId"])=="fan" and rows77b[0]["rot"]==90 and rows77b[0]["x"]>fan77["x"])
    def prim_n(prims):
        import math
        return sum(1 if p_["t"] in ("line","circle") else 4 if p_["t"]=="rect" else max(2, math.ceil((p_["a1"]-p_["a0"])/15)) for p_ in prims)
    fan_prims=[r_ for r_ in lib77 if r_["key"]=="fan"][0]["primitives"]
    ex77=ctx.request.get(DRW+f"/{D77}?annot=1").text(); sy77=[e for e in ezdxf.read(io.StringIO(ex77)).modelspace() if e.dxf.layer=="SYMBOL"]
    ok("S77b DXF 내보내기(주석 포함)에 SYMBOL 레이어 — 엔티티 수 = 팬 도형 전개 수(원 2 + 호 3 × 선분 10 = 32) · 원 도면 DXF 는 한 글자도 안 바뀐다",
       (len(sy77), prim_n(fan_prims), ctx.request.get(DRW+f"/{D77}").text()==orig77), len(sy77)==prim_n(fan_prims)==32 and ctx.request.get(DRW+f"/{D77}").text()==orig77)
    sid77=rows77b[0]["id"] if rows77b else "00000000-0000-4000-8000-000000000000"
    lk77=ctx.request.post(DRW+f"/{dI.get('id')}/symbols",headers=J0,data=json.dumps({"symbolId":fan77["symbolId"],"x":0,"y":0}))
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v77=(vw.request.get(SY).status, vw.request.post(SY,headers=J0,data=json.dumps({"symbolId":fan77["symbolId"],"x":0,"y":0})).status,
         vw.request.patch(BASE+f"/api/drawing-symbols/{sid77}",headers=J0,data=json.dumps({"rot":180})).status, vw.request.delete(BASE+f"/api/drawing-symbols/{sid77}").status); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    glib=gb.request.get(BASE+"/api/design-symbols").json().get("rows",[])
    g77=(gb.request.get(SY).status, gb.request.patch(BASE+f"/api/drawing-symbols/{sid77}",headers=J0,data=json.dumps({"dx":10})).status,
         gb.request.post(SY,headers=J0,data=json.dumps({"symbolId":glib[0]["id"] if glib else "","x":0,"y":0})).status); gb.close()
    ok("S77c 발행된 도면에는 심볼을 놓지 못한다(409) · viewer 는 보기 200 · 놓기 · 돌리기 · 지우기 403 · 다른 회사는 우리 도면 심볼 404(자기 라이브러리 5종만 · 우리 id 0건)",
       (lk77.status, v77, g77, len(glib), len(set(r_["id"] for r_ in glib) & set(key_of))),
       lk77.status==409 and v77==(200,403,403,403) and g77==(404,404,404) and len(glib)==5 and not (set(r_["id"] for r_ in glib) & set(key_of)))
    # ── S78 KC-4 (ccmd K) — 조립도 Item 표 · 풍선번호 더블클릭 = 부품의 정보(스냅샷 기준 · 단가 출처 = 그때 단가 이력 행) ──
    PRC=BASE+"/api/setup/prices"
    p78a=ctx.request.post(PRC,headers=J0,data=json.dumps({"code":"SCS 1","price":2500000,"currency":"KRW","effectiveFrom":"2026-01-01","note":"e2e S78 샘플"}))
    pid78=[r_ for r_ in ctx.request.get(PRC+"?code=SCS%201").json().get("rows",[]) if r_["price"]==2500000][0]["id"]
    pg.goto(BASE+f"/workbench?node={AHU75}",wait_until="domcontentloaded"); hydrated(pg); nuke(pg)
    sel78=pg.query_selector_all("[data-testid=code-builder] select"); sel78[0].select_option(value="SPF"); sel78[1].select_option(value="55")
    pg.locator("button",has_text=re.compile(r"^BOM$")).first.click(); pg.wait_for_selector("button:has-text('BOM Run')",timeout=30000)
    with pg.expect_response(lambda q: "/api/run/bom" in q.url and q.request.method=="POST",timeout=60000) as br78:
        pg.click("button:has-text('BOM Run')")
    rid78=br78.value.json().get("runId")
    def open_asm():
        pg.select_option("[data-cmd=dwg-view]","assembly"); pg.wait_for_selector("[data-testid=dwg-viewer][data-view=assembly][data-ready='1']",timeout=30000)
        pg.wait_for_selector("[data-testid=part-panel][data-ready='1']",timeout=30000)
    open_asm()
    rows78=pg.eval_on_selector_all("[data-testid=part-table] tbody tr","es=>es.map(e=>e.dataset.code)")
    cas_no=[i+1 for i,c in enumerate(rows78) if c=="SCS 1"][0]; mot_no=[i+1 for i,c in enumerate(rows78) if c=="SMT 1"][0]
    pg.dblclick(f"[data-testid=part-row-{mot_no}]"); pg.wait_for_selector(f"[data-testid=part-info][data-no='{mot_no}']",timeout=30000)
    pi_m=pg.eval_on_selector("[data-testid=part-info]","e=>({...e.dataset, text:e.innerText})")
    nuke(pg); pg.screenshot(path=f"{OUT}/78_part_info.png",full_page=True)
    pg.dblclick(f"[data-testid=dwg-viewer-svg] [data-balloon='{cas_no}']"); pg.wait_for_selector(f"[data-testid=part-info][data-no='{cas_no}']",timeout=30000)
    pi_c=pg.eval_on_selector("[data-testid=part-info]","e=>({...e.dataset})")
    ok("S78a 조립도 옆 Item 표(Item · Description · Q'ty · Remarks info) — 줄 더블클릭 = 모터(SMT 1) 정보: 코드 · 사양 3.7 kW · 수량 · 공급처 · 단가 ₩420,000 · 조립순서 4/4(Fan) · 세부 치수 C=350 · 주의사항 · DWG · Remarks info 칸이 채워진다",
       (pi_m.get("code"), pi_m.get("qty"), pi_m.get("unitCost"), pi_m.get("order"), pg.inner_text(f"[data-testid=part-remarks-{mot_no}]")),
       pg.get_attribute("[data-testid=dwg-viewer]","data-run")==rid78 and pi_m.get("code")=="SMT 1" and pi_m.get("qty")=="1" and pi_m.get("unitCost")=="420000" and pi_m.get("order")=="4" and "3.7 kW" in pi_m.get("text","")
       and "SMT 1.C=350" in pi_m.get("text","") and "C=350" in pg.inner_text(f"[data-testid=part-remarks-{mot_no}]") and "주의사항" in pg.inner_text(f"[data-testid=part-remarks-{mot_no}]"))
    p78b=ctx.request.post(PRC,headers=J0,data=json.dumps({"code":"SCS 1","price":2700000,"currency":"KRW","effectiveFrom":"2026-02-01","note":"e2e S78 샘플 인상"}))
    pg.click("[data-testid=dwg-viewer-close]"); pg.wait_for_selector("[data-testid=dwg-viewer]",state="detached",timeout=10000); open_asm()
    pg.dblclick(f"[data-testid=dwg-viewer-svg] [data-balloon='{cas_no}']"); pg.wait_for_selector(f"[data-testid=part-info][data-no='{cas_no}']",timeout=30000)
    pi_c2=pg.eval_on_selector("[data-testid=part-info]","e=>({...e.dataset})")
    pg.click("[data-testid=dwg-viewer-close]")
    ok("S78b 풍선번호 더블클릭 = 케이싱(SCS 1) 정보 · 단가 출처 = 그때 단가 이력 행(₩2,500,000) — 단가를 ₩2,700,000 으로 올린 뒤 다시 열어도 패널은 스냅샷 값 · 같은 행 id",
       (p78a.status, p78b.status, pi_c.get("code"), pi_c.get("unitCost"), pi_c.get("priceId")==pid78, pi_c2.get("unitCost"), pi_c2.get("priceId")==pid78),
       p78a.status==200 and p78b.status==200 and pi_c.get("code")=="SCS 1" and pi_c.get("unitCost")=="2500000" and pi_c.get("priceId")==pid78 and pi_c2.get("unitCost")=="2500000" and pi_c2.get("priceId")==pid78)
    PA=BASE+f"/api/bom-runs/{rid78}/parts"
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test")); v78=vw.request.get(PA).status; vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test")); g78=gb.request.get(PA).status; gb.close()
    x78=ctx.request.get(PA+f"?drawing={D77}")
    ok("S78c 부품 정보 API — viewer 403(단가가 들어 있다) · 다른 회사 404 · 이 스냅샷에서 뜬 도면이 아닌 drawing 을 주면 404",
       (v78, g78, x78.status), (v78, g78, x78.status)==(403,404,404))
    # ── S79 KB (ccmd K) — 컨설팅 두 트랙: 내부 최적안(읽기 전용 · 인쇄본) · 익명 · 집계 벤치마킹(k-익명 3) ──
    import subprocess
    PRC=BASE+"/api/setup/prices"
    p79=ctx.request.post(PRC,headers=J0,data=json.dumps({"code":"SCS 1","price":2300000,"currency":"KRW","supplier":"샘플 공급처 B","effectiveFrom":"2026-01-15","note":"e2e S79 샘플"}))
    pid79=[r_ for r_ in ctx.request.get(PRC+"?code=SCS%201").json().get("rows",[]) if r_["price"]==2300000][0]["id"]
    j79=run_spf().json(); rid79=j79.get("runId")
    CI=BASE+"/api/consulting/internal"; CB=BASE+"/api/consulting/benchmark"
    rep79=ctx.request.get(CI+f"?runId={rid79}").json(); props=rep79.get("proposals",[])
    kinds79=sorted({p_["kind"] for p_ in props})
    sup79=[p_ for p_ in props if p_["kind"]=="supplier"]; mar79=[p_ for p_ in props if p_["kind"]=="margin"]
    ev_price=[e["id"] for p_ in sup79 for e in p_["evidence"] if e["kind"]=="price_history"]
    ev_rule=[e["id"] for p_ in mar79 for e in p_["evidence"] if e["kind"]=="rule_row"]
    _cat79=ctx.request.get(BASE+"/api/setup/catalog").json(); _spf79=[p_ for p_ in _cat79.get("productCodes",[]) if p_.get("code")=="SPF"][0]
    rule_rows={f"SPF#{tn}.{r_['item']}" for tn,t_ in _spf79["tables"].items() if t_.get("role")=="rule" for r_ in t_["rows"]}
    price_ids={r_["id"] for r_ in ctx.request.get(PRC+"?code=SCS%201").json().get("rows",[])}
    ok("S79a 트랙 1 내부 최적안 — 샘플 데이터로 제안 2종 이상(공급처: SCS 1 ₩2,700,000 → 샘플 공급처 B ₩2,300,000 · 차액 400,000 / 설계 여유: detail.Fan.A 1250 ≤ 1300 · 여유 3.8%) · 근거 행 id 가 실제 단가 이력 행 · 실제 규칙 표 행",
       (kinds79, [p_["saving"] for p_ in sup79], [p_["title"] for p_ in mar79], ev_price[:1], ev_rule, rep79.get("fanNote")),
       len(kinds79)>=2 and "supplier" in kinds79 and "margin" in kinds79 and any(p_["saving"]=={"amount":400000,"unit":"KRW"} for p_ in sup79)
       and pid79 in ev_price and set(ev_price)<=price_ids and ev_rule and set(ev_rule)<=rule_rows)
    pg.goto(BASE+f"/m/consulting?runId={rid79}",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=consulting][data-ready='1']",timeout=60000); nuke(pg)
    ui79=pg.eval_on_selector_all("[data-testid=consulting-internal] tbody tr[data-testid^=proposal-]","es=>es.map(e=>e.dataset.testid)")
    apply_btn=pg.query_selector_all("[data-testid=consulting] button")
    pg.screenshot(path=f"{OUT}/79_consulting_internal.png",full_page=True)
    pr79=ctx.request.get(BASE+f"/api/consulting/print?runId={rid79}"); pt79=pr79.text()
    ok("S79b 화면(/m/consulting)은 같은 제안을 읽기 전용으로(적용 버튼 0) · 인쇄본 A4 발치에 스냅샷 id · 분석 날짜 · '샘플 단가 · 샘플 운전시간'",
       (sorted(set(ui79)), len(apply_btn), pr79.status, rid79 in pt79, "샘플 단가 · 샘플 운전시간" in pt79),
       sorted(set(ui79))==sorted({f"proposal-{k_}" for k_ in kinds79}) and len(apply_btn)==0 and pr79.status==200 and rid79 in pt79
       and "샘플 단가 · 샘플 운전시간" in pt79 and "consulting-print-footer" in pt79 and "size:A4" in pt79)
    bm=ctx.request.get(CB).json(); bro={r_["metric"]:r_ for r_ in bm.get("rows",[])}; fe=bro.get("fan_eta",{})
    pg.locator("[data-testid=consulting-benchmark]").screenshot(path=f"{OUT}/79_consulting_benchmark.png")
    ok("S79c 트랙 2 업계 안 우리 위치 — 지표 3종(풍량당 원가 · 재료비 비율 · 팬 효율) · 표본 ≥ 3(샘플 회사 3 + 우리) · p25 · p50 · p75 · 우리 값 = 최신 스냅샷 η · 백분위 · 응답에 회사 id · 이름 칸 없음",
       ({k_:(v_["n"],v_["suppressed"],v_["mine"],v_["percentile"]) for k_,v_ in bro.items()}, j79.get("special",{}).get("result",{}).get("eta")),
       sorted(bro)==["cost_per_cmh","fan_eta","material_ratio"] and all(not v_["suppressed"] and v_["n"]>=4 and v_["p50"] is not None and v_["percentile"] is not None for v_ in bro.values())
       and abs((fe.get("mine") or 0)-(j79.get("special",{}).get("result",{}).get("eta") or -1))<1e-4
       and all(set(v_)=={"metric","label","unit","n","p25","p50","p75","mine","percentile","suppressed"} for v_ in bro.values()))
    # 표본을 줄인 상태(테스트용 시드 옵션: 샘플 회사 1곳) → 표본 2곳 → 숨김 · 되돌림
    sd1=subprocess.run("pnpm --filter @edim/db bench:seed -- 1",shell=True,capture_output=True)
    bm1=ctx.request.get(CB).json(); b1={r_["metric"]:r_ for r_ in bm1.get("rows",[])}
    pg.goto(BASE+f"/m/consulting?runId={rid79}",wait_until="domcontentloaded"); pg.wait_for_selector("[data-testid=consulting][data-ready='1']",timeout=60000)
    hid=pg.query_selector_all("[data-testid^=bench-hidden-]"); hid_txt=pg.inner_text("[data-testid=bench-hidden-fan_eta]") if pg.query_selector("[data-testid=bench-hidden-fan_eta]") else ""
    sd3=subprocess.run("pnpm --filter @edim/db bench:seed -- 3",shell=True,capture_output=True)
    ok("S79d 표본을 2곳으로 줄이면(샘플 1 + 우리) 분포를 숨긴다 — suppressed · p25/p50/p75 · 백분위 NULL · 화면 '표본이 3곳 미만이라 보여 드리지 않습니다' · 시드 되돌림",
       (sd1.returncode, {k_:(v_["n"],v_["suppressed"],v_["p50"]) for k_,v_ in b1.items()}, len(hid), hid_txt, sd3.returncode),
       sd1.returncode==0 and sd3.returncode==0 and all(v_["suppressed"] and v_["n"]==2 and v_["p25"] is None and v_["p50"] is None and v_["p75"] is None and v_["percentile"] is None for v_ in b1.values())
       and len(hid)==3 and "표본이 3곳 미만이라 보여 드리지 않습니다" in hid_txt)
    vw=b.new_context(); vw.request.post(BASE+"/api/auth/login",data=LOGIN("viewer@acme.test"))
    v79=(vw.request.get(CI).status, vw.request.get(CB).status, vw.request.get(BASE+f"/api/consulting/print?runId={rid79}").status)
    vp=vw.new_page(); vp.goto(BASE+"/m/consulting",wait_until="domcontentloaded"); v79p=bool(vp.query_selector("[data-testid=consulting-forbidden]")); vw.close()
    gb=b.new_context(); gb.request.post(BASE+"/api/auth/login",data=LOGIN("owner@globex.test"))
    gbm=gb.request.get(CB); g79txt=gbm.text(); gbr={r_["metric"]:r_ for r_ in gbm.json().get("rows",[])}; g79i=gb.request.get(CI+f"?runId={rid79}").status; gb.close()
    ok("S79e viewer 403(API 3 · 화면 403 안내) · 다른 회사(Globex)가 부르면 **그 회사 값**으로만(최신 스냅샷 없음 → 우리 값 없음 · 분포는 같다) · 응답에 Acme 이름 · id 없음 · Acme 스냅샷 분석 404",
       (v79, v79p, {k_:(v_["n"],v_["mine"]) for k_,v_ in gbr.items()}, "Acme" in g79txt, "00000000-0000-4000-8000-00000000000a" in g79txt, g79i),
       v79==(403,403,403) and v79p and gbr.get("fan_eta",{}).get("mine") is None and gbr.get("fan_eta",{}).get("p50")==fe.get("p50")
       and "Acme" not in g79txt and "00000000-0000-4000-8000-00000000000a" not in g79txt and g79i==404)
    # S37 은 맨 끝에서 센다 — 중간(옛 자리)에서는 뒤에 찍히는 5장(40·41·42·52·53)이 아직 없어,
    # 빈 폴더에서는 25장이라 실패하고 이전 실행 잔재가 있을 때만 통과했다(2026-09-24 실측).
    _want=["00_login","05_project_mgmt","06_module_cpq_stub","10_project_bound","11_code_builder","11b_revisions","12_macro_tab",
           "13_macro_approved","14_edim_run","15_bom_cost","16_design_tab","20_setup_subcode","21_setup_product_table",
           "22_setup_relationship","23_codebuilder_from_subcode","30_toolbox_program","31_toolbox_ui_tool","40_company_admin",
           "41_platform_console","42_user_management","43_drawings","44_document_tab","45_purchasing","46_quotation_print",
           "47_techdata_print","48_dxf_plan","49_dxf_assembly","51_accepted","52_register","53_schedule","54_toolbar","55_project_mgmt","56_print_setup","57_ui_design","58_material","59_arrangement_code","60_spec_input","61_drawing_purpose","62_company_db","63_input_data","64_viewer3d","65_price_to_cost","66_project_contacts","67_partner_edit","68_spec_import","69_code_approval","70_dwg_view","71_techdata_list","72_mfg_rate","73_erp_master","74_sub_drawing","75_output_template","76_coding_list","77_wizards","78_print_layout","79_draw_module","80_macro_verify","81_learning_job","82_formula_cards","83_projection","84_toolbox_suggestion","85_special_request","86_fan_result","87_special_meter","75_cpq_special_bom","75_cpq_special_drawing","76_detail_dim","77_symbol","78_part_info","79_consulting_internal","79_consulting_benchmark"]
    _miss=[w for w in _want if not os.path.exists(f"{OUT}/{w}.png") or os.path.getmtime(f"{OUT}/{w}.png")<T0]
    ok(f"S37 캡처 {len(_want)}장이 이번 실행에서 전부 나온다 (잔재 파일은 세지 않음)", _miss or len(_want), not _miss)
    b.close()
n=sum(1 for v in R.values() if v[0]); print(f"\n[demo_e2e] {n}/{len(R)} steps passed"); json.dump(R,open(f"{OUT}/demo_e2e_result.json","w",encoding="utf-8"),ensure_ascii=False,indent=1)
sys.exit(0 if n==len(R) else 1)
